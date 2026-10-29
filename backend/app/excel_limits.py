"""Bound XLSX decompression and XML work before openpyxl allocates objects."""

from __future__ import annotations

import io
import posixpath
import re
import zlib
from dataclasses import dataclass, field
from xml.parsers import expat
from zipfile import ZIP_DEFLATED, ZIP_STORED, BadZipFile, ZipFile

from openpyxl.formula.tokenizer import Token, Tokenizer, TokenizerError

from app.limits import (
    EXCEL_MAX_CELLS,
    EXCEL_MAX_COLUMNS,
    EXCEL_MAX_DECODED_CHARACTERS,
    EXCEL_MAX_EXPANDED_BYTES,
    EXCEL_MAX_MEMBER_BYTES,
    EXCEL_MAX_MEMBERS,
    EXCEL_MAX_ROWS,
    EXCEL_MAX_SHEETS,
    TEXT_MAX_LENGTH,
)

_CHUNK_BYTES = 16 * 1024
_COORDINATE = re.compile(r"([A-Za-z]{1,3})([1-9][0-9]{0,6})\Z")
_SHEET_NAMESPACE = "http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_RELATIONSHIP_NAMESPACE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _check(value: int, limit: int, label: str) -> None:
    if value > limit:
        raise ValueError(f"Excel {label} must be at most {limit}")


def _coordinate(value: str) -> tuple[int, int]:
    match = _COORDINATE.fullmatch(value)
    if match is None:
        raise ValueError("Invalid Excel cell coordinate")
    column = 0
    for letter in match[1].upper():
        column = column * 26 + ord(letter) - ord("A") + 1
    return int(match[2]), column


@dataclass
class _Budget:
    rows: int = 0
    cells: int = 0
    sheets: int = 0
    sheet_declarations: int = 0
    elements: int = 0
    sheet_targets: set[str] = field(default_factory=set)
    sheet_relationships: dict[str, set[str]] = field(default_factory=dict)
    decoded_characters: int = 0
    shared_strings: dict[int, int] = field(default_factory=dict)
    shared_string_uses: dict[int, int] = field(default_factory=dict)

    def add_text(self, length: int) -> None:
        self.decoded_characters += length
        _check(self.decoded_characters, EXCEL_MAX_DECODED_CHARACTERS, "decoded characters")

    def check_shared_strings(self) -> None:
        for index, count in self.shared_string_uses.items():
            length = self.shared_strings.get(index)
            if length is None:
                raise ValueError("Invalid Excel shared-string index")
            self.add_text(length * count)


def _validate_sheet_relationships(archive: ZipFile, path: str, ids: set[str]) -> None:
    found = {}
    depth = 0

    def start(name: str, attrs: dict[str, str]) -> None:
        nonlocal depth
        depth += 1
        if depth == 2 and name.rsplit("}", 1)[-1] == "Relationship" and attrs.get("Id") in ids:
            found[attrs["Id"]] = (attrs.get("Type"), attrs.get("TargetMode"))

    def end(_name: str) -> None:
        nonlocal depth
        depth -= 1

    parser = expat.ParserCreate(namespace_separator="}")
    parser.StartElementHandler = start
    parser.EndElementHandler = end
    if path not in archive.namelist():
        raise ValueError("Excel worksheet relationships are missing")
    # These immutable members have already passed the archive and XML checks.
    with archive.open(path) as stream:
        while chunk := stream.read(_CHUNK_BYTES):
            parser.Parse(chunk, False)
    parser.Parse(b"", True)
    for relationship_id in ids:
        kind, mode = found.get(relationship_id, (None, None))
        if kind != f"{_RELATIONSHIP_NAMESPACE}/worksheet" or mode == "External":
            raise ValueError("Excel sheets must use internal worksheet relationships")


def _shared_formula_bound(formula: str) -> int:
    try:
        tokens = Tokenizer("=" + formula).items
    except TokenizerError as exc:
        raise ValueError("Invalid Excel shared formula") from exc
    # A translated coordinate component can grow by at most seven characters.
    growth = sum(
        7 * sum(1 for _ in re.finditer(r"[A-Za-z]+|[0-9]+", token.value))
        for token in tokens
        if token.type == Token.OPERAND and token.subtype == Token.RANGE
    )
    return len(formula) + 1 + growth


class _XmlCheck:
    def __init__(self, budget: _Budget, part_name: str):
        self.budget = budget
        self.depth = 0
        self.elements: list[str] = []
        self.worksheet = False
        self.max_row = self.max_column = 0
        self.row_count = self.cell_count = 0
        self.row_has_extension = False
        self.row_extensions = 0
        self.row = self.column = 0
        self.text_depth = self.text_length = 0
        self.value_depth = 0
        self.formula_depth = 0
        self.cell_depth = 0
        self.cell_has_formula = False
        self.cell_shared_string = False
        self.cell_has_value = False
        self.shared_value_depth = 0
        self.shared_value_parts: list[str] = []
        self.shared_string_depth = self.shared_string_length = self.shared_string_count = 0
        self.formula_parts: list[str] = []
        self.shared_formula_id: str | None = None
        self.shared_formulae: dict[str | None, int] = {}
        self.sheet_ids: set[str] = set()
        self.sheet_names: set[str] = set()
        self.relationship_base = "/" + posixpath.dirname(posixpath.dirname(part_name))
        folder, filename = posixpath.split(part_name)
        self.relationship_path = posixpath.join(folder, "_rels", filename + ".rels")
        self.parser = expat.ParserCreate(namespace_separator="}")
        self.parser.StartElementHandler = self.start
        self.parser.EndElementHandler = self.end
        self.parser.CharacterDataHandler = self.text
        self.parser.StartDoctypeDeclHandler = self.reject_declaration
        self.parser.EntityDeclHandler = self.reject_declaration
        self.parser.ExternalEntityRefHandler = self.reject_declaration

    @staticmethod
    def reject_declaration(*_args):
        raise ValueError("Excel XML must not contain DTD or entity declarations")

    def check_grid(self) -> None:
        _check(self.max_column, EXCEL_MAX_COLUMNS, "columns")
        rows = max(self.max_row, self.row_count)
        _check(self.budget.rows + rows, EXCEL_MAX_ROWS, "total row slots")
        cells = max(rows * self.max_column, self.cell_count) + self.row_extensions
        _check(self.budget.cells + cells, EXCEL_MAX_CELLS, "total cell slots")

    def start(self, name: str, attrs: dict[str, str]) -> None:
        parent = self.elements[-1] if self.elements else None
        if self.shared_value_depth:
            raise ValueError("Invalid Excel shared-string index")
        if parent == _SHEET_NAMESPACE + "row":
            if name == _SHEET_NAMESPACE + "extLst":
                if self.row_has_extension:
                    raise ValueError("Excel rows must have at most one extension list")
                self.row_has_extension = True
                # openpyxl also processes the extension as an empty cell record.
                self.row_extensions += 1
                self.check_grid()
            elif name != _SHEET_NAMESPACE + "c":
                raise ValueError("Invalid Excel worksheet row element")
            elif self.row_has_extension:
                raise ValueError("Excel row extensions must follow their cells")
        if name == _SHEET_NAMESPACE + "row":
            if name in self.elements:
                raise ValueError("Excel worksheet rows must not be nested")
            self.row_has_extension = False
        if name == _SHEET_NAMESPACE + "c" and parent == _SHEET_NAMESPACE + "row":
            self.cell_depth = self.depth + 1
            self.cell_has_formula = False
            self.cell_shared_string = attrs.get("t") == "s"
            self.cell_has_value = False
            self.shared_value_parts = []
        if name == _SHEET_NAMESPACE + "si":
            if self.shared_string_depth:
                raise ValueError("Excel shared strings must not be nested")
            self.shared_string_depth = self.depth + 1
            self.shared_string_length = 0
        if (
            name == _SHEET_NAMESPACE + "v"
            and self.cell_shared_string
            and self.cell_depth
            and self.depth == self.cell_depth
        ):
            if self.cell_has_value:
                raise ValueError("Excel cells must have at most one shared-string index")
            self.cell_has_value = True
            self.shared_value_depth = self.depth + 1
        self.elements.append(name)
        spreadsheet_element = name.startswith(_SHEET_NAMESPACE)
        name = name.rsplit("}", 1)[-1]
        self.depth += 1
        _check(self.depth, 128, "XML nesting depth")
        self.budget.elements += 1
        _check(self.budget.elements, EXCEL_MAX_CELLS * 8, "XML elements")
        # openpyxl reads worksheet records even under a nonstandard root.
        if (
            not self.worksheet
            and spreadsheet_element
            and name in {"worksheet", "sheetData", "dimension", "row", "c"}
        ):
            self.worksheet = True
            self.budget.sheets += 1
            _check(self.budget.sheets, EXCEL_MAX_SHEETS, "worksheets")
        if self.depth == 3 and parent is not None and parent.rsplit("}", 1)[-1] == "sheets":
            if name != "sheet" or "id" in attrs:
                raise ValueError("Invalid Excel worksheet declaration")
            self.budget.sheet_declarations += 1
            _check(self.budget.sheet_declarations, EXCEL_MAX_SHEETS, "worksheets")
            if "name" in attrs:
                title = attrs["name"]
                if title in self.sheet_names:
                    raise ValueError("Excel worksheets must have distinct names")
                self.sheet_names.add(title)
            relationship_id = attrs.get(
                "http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
            )
            if relationship_id is not None:
                if relationship_id in self.sheet_ids:
                    raise ValueError("Excel worksheets must use distinct relationships")
                self.sheet_ids.add(relationship_id)
                self.budget.sheet_relationships.setdefault(self.relationship_path, set()).add(
                    relationship_id
                )
        if (
            name == "Relationship"
            and self.depth == 2
            and attrs.get("Type", "").endswith("/worksheet")
        ):
            target = posixpath.normpath(
                posixpath.join(self.relationship_base, attrs.get("Target", ""))
            )
            target = "/" + target.lstrip("/")
            if target in self.budget.sheet_targets:
                raise ValueError("Excel worksheets must use distinct targets")
            self.budget.sheet_targets.add(target)
        if self.worksheet and spreadsheet_element:
            if name == "dimension":
                for point in attrs.get("ref", "A1").split(":"):
                    row, column = _coordinate(point)
                    self.max_row = max(self.max_row, row)
                    self.max_column = max(self.max_column, column)
            elif name == "row":
                self.row_count += 1
                self.row = int(attrs.get("r", self.row + 1))
                if self.row < 1:
                    raise ValueError("Invalid Excel row coordinate")
                self.max_row = max(self.max_row, self.row)
                self.column = 0
            elif name == "c":
                self.cell_count += 1
                if "r" in attrs:
                    row, self.column = _coordinate(attrs["r"])
                    self.max_row = max(self.max_row, row)
                else:
                    self.column += 1
                self.max_column = max(self.max_column, self.column)
            if name in {"dimension", "row", "c"}:
                self.check_grid()
        # Count all rich-text fragments together without retaining their text.
        if spreadsheet_element and name in {"c", "si"} and not self.text_depth:
            self.text_depth = self.depth
            self.text_length = 0
        if self.text_depth and not self.value_depth and name in {"t", "v", "f"}:
            self.value_depth = self.depth
        if (
            self.cell_depth
            and spreadsheet_element
            and name == "f"
            and self.depth == self.cell_depth + 1
        ):
            if self.cell_has_formula:
                raise ValueError("Excel cells must have at most one formula")
            self.cell_has_formula = True
            if attrs.get("t") == "shared":
                self.formula_depth = self.depth
                self.shared_formula_id = attrs.get("si")
                self.formula_parts = []

    def text(self, value: str) -> None:
        if self.value_depth:
            self.text_length += len(value)
            _check(self.text_length, TEXT_MAX_LENGTH, "cell/shared-string characters")
            self.budget.add_text(len(value))
            if self.shared_string_depth:
                self.shared_string_length += len(value)
            if self.shared_value_depth:
                self.shared_value_parts.append(value)
            if self.formula_depth:
                self.formula_parts.append(value)

    def end(self, _name: str) -> None:
        if self.depth == self.shared_string_depth:
            index = self.shared_string_count
            self.shared_string_count += 1
            _check(self.shared_string_count, EXCEL_MAX_CELLS, "shared strings per part")
            # Multiple tables cannot lower the bound for the table openpyxl selects.
            self.budget.shared_strings[index] = max(
                self.budget.shared_strings.get(index, 0), self.shared_string_length
            )
            self.shared_string_depth = 0
        if self.depth == self.shared_value_depth:
            self.shared_value_depth = 0
        if self.depth == self.formula_depth:
            formula = "".join(self.formula_parts)
            index = self.shared_formula_id
            if index not in self.shared_formulae and formula:
                self.shared_formulae[index] = _shared_formula_bound(formula)
            bound = self.shared_formulae.get(index, 1)
            self.budget.add_text(max(bound - len(formula), 0))
            self.formula_depth = 0
            self.formula_parts = []
        if self.depth == self.value_depth:
            self.value_depth = 0
        if self.depth == self.text_depth:
            self.text_depth = 0
        if self.depth == self.cell_depth:
            if self.cell_shared_string and not self.cell_has_formula and self.shared_value_parts:
                try:
                    index = int("".join(self.shared_value_parts))
                except ValueError as exc:
                    raise ValueError("Invalid Excel shared-string index") from exc
                if index < 0:
                    raise ValueError("Invalid Excel shared-string index")
                self.budget.shared_string_uses[index] = (
                    self.budget.shared_string_uses.get(index, 0) + 1
                )
            self.shared_value_parts = []
            self.cell_depth = 0
        self.elements.pop()
        self.depth -= 1

    def finish(self) -> None:
        self.parser.Parse(b"", True)
        if self.worksheet:
            self.check_grid()
            rows = max(self.max_row, self.row_count)
            self.budget.rows += rows
            self.budget.cells += max(rows * self.max_column, self.cell_count) + self.row_extensions


def validate_workbook(raw: bytes) -> None:
    """Validate every member, including metadata and ignored worksheets."""
    try:
        with ZipFile(io.BytesIO(raw)) as archive:
            members = archive.infolist()
            _check(len(members), EXCEL_MAX_MEMBERS, "archive entries")
            names = [member.filename for member in members]
            if len(names) != len(set(names)):
                raise ValueError("Excel archive contains duplicate member names")
            _check(
                sum(member.file_size for member in members),
                EXCEL_MAX_EXPANDED_BYTES,
                "expanded archive bytes",
            )
            for member in members:
                _check(member.file_size, EXCEL_MAX_MEMBER_BYTES, "expanded member bytes")
                # zipfile only bounds decompressor output for stored/deflate members.
                if member.compress_type not in {ZIP_STORED, ZIP_DEFLATED}:
                    raise ValueError("Excel archive must use stored or deflate compression")
            budget = _Budget()
            total = 0
            for member in members:
                xml = (
                    _XmlCheck(budget, member.filename)
                    if member.filename.lower().endswith((".xml", ".rels"))
                    else None
                )
                size = 0
                sniffed = xml is not None
                with archive.open(member) as stream:
                    while chunk := stream.read(_CHUNK_BYTES):
                        size += len(chunk)
                        total += len(chunk)
                        _check(size, EXCEL_MAX_MEMBER_BYTES, "expanded member bytes")
                        _check(total, EXCEL_MAX_EXPANDED_BYTES, "expanded archive bytes")
                        if not sniffed:
                            # Relationship targets may use arbitrary file extensions.
                            prefix = chunk.lstrip()
                            if prefix:
                                sniffed = True
                                if prefix.startswith(
                                    (
                                        b"<",
                                        b"\xef\xbb\xbf",
                                        b"\xff\xfe",
                                        b"\xfe\xff",
                                        b"\x00<",
                                        b"\x00\x00\xfe\xff",
                                    )
                                ):
                                    xml = _XmlCheck(budget, member.filename)
                        if xml is not None:
                            xml.parser.Parse(chunk, False)
                if xml is not None:
                    xml.finish()
            for path, ids in budget.sheet_relationships.items():
                _validate_sheet_relationships(archive, path, ids)
            budget.check_shared_strings()
    except (
        BadZipFile,
        expat.ExpatError,
        EOFError,
        RuntimeError,
        NotImplementedError,
        zlib.error,
    ) as exc:
        raise ValueError("Invalid or unsupported Excel archive/XML") from exc
