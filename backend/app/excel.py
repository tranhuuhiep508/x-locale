"""Excel import/export via openpyxl."""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from app.excel_limits import validate_workbook
from app.helpers import (
    ensure_max_length,
    key_preview,
    reject_nul,
    validate_locale_code,
    validate_module_slug,
)
from app.limits import KEY_MAX_LENGTH, TAG_NAME_MAX_LENGTH, TEXT_MAX_LENGTH
from app.models import (
    Project,
    ProjectLayout,
    StringEntry,
    TranslationStatus,
)
from app.schemas import ImportResult
from app.services.catalog import owned_module
from app.services.import_bulk import ImportRow, import_rows
from app.services.sync import check_export_keys

UNASSIGNED_SHEET = "_unassigned"
FLAT_SHEET = "strings"
MODULE_SLUG_PREFIX = "module_slug:"

DEMO_IMPORT_ROWS: dict[str, list[tuple[str, str]]] = {
    "auth": [("auth.email", "Địa chỉ email"), ("auth.sign_in", "Đăng nhập")],
    "common": [("common.save", "Lưu"), ("common.cancel", "Hủy")],
    UNASSIGNED_SHEET: [("hello", "Xin chào")],
}


def _layout_value(project: Project) -> str:
    layout = project.layout
    return layout.value if hasattr(layout, "value") else str(layout)


def _is_modular(project: Project) -> bool:
    return _layout_value(project) == ProjectLayout.modular.value


def _export_locales(project: Project, locale: str | None) -> list[str]:
    locales = [project.base_language, *project.target_languages]
    if not locale:
        return locales
    if locale not in locales:
        raise ValueError(f"Unknown locale '{locale}'")
    return [locale]


def _module_sheet_names(slugs: list[str]) -> dict[str, str]:
    used = {"_meta", UNASSIGNED_SHEET}
    names = {}
    # Reserve complete short slugs before shortened names can claim them.
    for slug in sorted(slugs):
        if len(slug) <= 31 and slug.casefold() not in used:
            names[slug] = slug
            used.add(slug.casefold())
    for slug in sorted(slugs):
        if slug in names:
            continue
        name, number = slug[:31], 0
        while name.casefold() in used:
            number += 1
            suffix = f"_{number}"
            name = slug[: 31 - len(suffix)] + suffix
        names[slug] = name
        used.add(name.casefold())
    return names


def build_workbook(
    project: Project,
    entries: list[StringEntry],
    stage: str,
    locale: str | None = None,
) -> bytes:
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    locales = _export_locales(project, locale)
    module_sheets = {}
    check_export_keys(project, entries, stage, modular=_is_modular(project))
    if _is_modular(project):
        by_module: dict[str | None, list[StringEntry]] = {}
        for e in entries:
            module = e.published_module if stage == "public" else e.module
            owned = owned_module(project.id, module)
            slug = owned.slug if owned else None
            by_module.setdefault(slug, []).append(e)
        names = _module_sheet_names([k for k in by_module if k is not None])
        for slug in sorted(names):
            name = names[slug]
            _write_sheet(wb, name, by_module[slug], locales, project, stage)
            module_sheets[name] = slug
        if None in by_module or not by_module:
            _write_sheet(wb, UNASSIGNED_SHEET, by_module.get(None, []), locales, project, stage)
            module_sheets[UNASSIGNED_SHEET] = UNASSIGNED_SHEET
    else:
        _write_sheet(wb, FLAT_SHEET, entries, locales, project, stage)

    _write_meta_sheet(wb, project, stage, module_sheets)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_template_workbook(project: Project) -> bytes:
    """Demo workbook matching the project's layout so users can see the import format."""
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    locales = [project.base_language, *project.target_languages]
    module_sheets = {}
    if _is_modular(project):
        names = _module_sheet_names([slug for slug in DEMO_IMPORT_ROWS if slug != UNASSIGNED_SHEET])
        for slug, rows in DEMO_IMPORT_ROWS.items():
            name = UNASSIGNED_SHEET if slug == UNASSIGNED_SHEET else names[slug]
            _write_template_sheet(wb, name, rows, locales, project.base_language)
            module_sheets[name] = slug
    else:
        flat_rows = [row for rows in DEMO_IMPORT_ROWS.values() for row in rows]
        _write_template_sheet(wb, FLAT_SHEET, flat_rows, locales, project.base_language)
    _write_meta_sheet(wb, project, "draft", module_sheets)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _write_meta_sheet(
    wb: Workbook, project: Project, stage: str, module_sheets: dict[str, str]
) -> None:
    meta = wb.create_sheet("_meta")
    _append_text_row(meta, ["key", "value"])
    _append_text_row(meta, ["project_id", str(project.id)])
    _append_text_row(meta, ["project_slug", project.slug])
    _append_text_row(meta, ["base_language", project.base_language])
    _append_text_row(meta, ["stage", stage])
    _append_text_row(meta, ["target_languages", ",".join(project.target_languages)])
    _append_text_row(meta, ["layout", _layout_value(project)])
    for name, slug in module_sheets.items():
        _append_text_row(meta, [f"{MODULE_SLUG_PREFIX}{name}", slug])
    meta.sheet_state = "hidden"


def _write_template_sheet(
    wb: Workbook,
    name: str,
    rows: list[tuple[str, str]],
    locales: list[str],
    base_language: str,
) -> None:
    ws = wb.create_sheet(name)
    _append_text_row(ws, ["key", "description", "tags", *locales])
    for key, source in rows:
        row: list[Any] = [key, "", ""]
        for locale in locales:
            row.append(source if locale == base_language else "")
        _append_text_row(ws, row)


def _excel_flag(value: Any) -> bool:
    return value is True or (isinstance(value, str) and value.lower() in {"1", "true"})


def _data_table_text(formula: DataTableFormula) -> str:
    # The file stores input cells, not a formula string. Excel shows TABLE().
    # shortcut: deleted inputs (del1/del2) stay as their cell ref, turn them into #REF! if those files show up.
    first = formula.r1 or ""
    second = formula.r2 or ""
    if _excel_flag(formula.dt2D):
        return f"=TABLE({first},{second})"
    if _excel_flag(formula.dtr):
        return f"=TABLE({first},)"
    return f"=TABLE(,{first})"


def _text_value(val: Any) -> str:
    if isinstance(val, ArrayFormula):
        return val.text or ""
    if isinstance(val, DataTableFormula):
        return _data_table_text(val)
    return "" if val is None else str(val)


def _append_text_row(ws: Worksheet, values: list[Any]) -> None:
    cells = []
    for value in values:
        cell = WriteOnlyCell(ws, value=_text_value(value))
        # Explicit string cells preserve literal text without creating Excel formulas.
        cell.data_type = "s"
        cells.append(cell)
    ws.append(cells)


def _write_sheet(
    wb: Workbook,
    name: str,
    entries: list[StringEntry],
    locales: list[str],
    project: Project,
    stage: str,
) -> None:
    ws = wb.create_sheet(name)
    headers = ["key", "description", "tags", *locales]
    _append_text_row(ws, headers)

    def _row_key(entry: StringEntry) -> str:
        if stage == "public":
            return entry.published_key or entry.key
        return entry.key

    for entry in sorted(entries, key=_row_key):
        if entry.deleted_at is not None:
            continue
        if stage == "public":
            if entry.status != TranslationStatus.public:
                continue
            row_key = entry.published_key or entry.key
            source = (
                entry.published_source_text
                if entry.published_source_text is not None
                else entry.source_text
            )
        else:
            if entry.pending_delete:
                continue
            row_key = entry.key
            source = entry.source_text
        tags = ",".join(
            tag.name for tag in sorted(entry.tags or [], key=lambda tag: (tag.name, str(tag.id)))
        )
        row: list[Any] = [
            row_key,
            entry.description,
            tags,
        ]
        for locale in locales:
            if locale == project.base_language:
                row.append(source)
            else:
                t = next((x for x in entry.translations if x.locale == locale), None)
                if stage == "public":
                    if t is not None and t.published_value is not None:
                        row.append(t.published_value)
                    else:
                        row.append("")
                else:
                    row.append(t.value if t else "")
        _append_text_row(ws, row)


def import_workbook(
    db: Session,
    project: Project,
    raw: bytes,
    *,
    dry_run: bool = False,
    status: TranslationStatus = TranslationStatus.draft,
) -> ImportResult:
    validate_workbook(raw)
    # data_only would replace a formula cell with its cached result, or blank.
    wb = load_workbook(io.BytesIO(raw), read_only=True, keep_links=False)
    rows, groups = [], []
    try:
        modular = _is_modular(project)
        module_sheets = {}
        if "_meta" in wb.sheetnames:
            meta = {}
            identities = set()
            for row in wb["_meta"].iter_rows(min_row=2, values_only=True):
                if not row or not row[0]:
                    continue
                key, value = row[0], row[1] if len(row) > 1 else None
                if modular and isinstance(key, str) and key.startswith(MODULE_SLUG_PREFIX):
                    name = key[len(MODULE_SLUG_PREFIX) :]
                    if name in module_sheets:
                        raise ValueError(f"Duplicate module mapping for sheet {name!r}")
                    if name == "_meta" or name not in wb.sheetnames:
                        raise ValueError(f"Module mapping references unknown data sheet {name!r}")
                    if not isinstance(value, str):
                        raise ValueError(f"Invalid module slug mapping for sheet {name!r}")
                    if (name == UNASSIGNED_SHEET) != (value == UNASSIGNED_SHEET):
                        raise ValueError(f"Invalid unassigned module mapping for sheet {name!r}")
                    if value != UNASSIGNED_SHEET:
                        reject_nul(value, field="module slug", context=f"Sheet {name!r}: ")
                        value = validate_module_slug(value)
                    if value in identities:
                        raise ValueError(f"Conflicting module mappings for slug {value!r}")
                    identities.add(value)
                    module_sheets[name] = value
                meta[key] = value
            if meta.get("project_id") and str(meta["project_id"]) != str(project.id):
                if meta.get("project_slug") and meta["project_slug"] != project.slug:
                    raise ValueError(
                        f"Workbook belongs to project {meta.get('project_slug')}, not {project.slug}"
                    )
        for sheet_name in wb.sheetnames:
            if sheet_name == "_meta":
                continue
            iterator = wb[sheet_name].iter_rows(values_only=True)
            first = next(iterator, None)
            if first is None:
                continue
            headers = [str(value).strip() if value else "" for value in first]
            if "key" not in headers:
                continue
            key_idx = headers.index("key")
            desc_idx = headers.index("description") if "description" in headers else None
            tags_idx = headers.index("tags") if "tags" in headers else None
            module = None
            if modular:
                if module_sheets and sheet_name not in module_sheets:
                    raise ValueError(f"Missing module mapping for data sheet {sheet_name!r}")
                slug = module_sheets.get(sheet_name, sheet_name)
                if slug != UNASSIGNED_SHEET:
                    reject_nul(slug, field="module slug", context=f"Sheet {sheet_name!r}: ")
                    try:
                        module = validate_module_slug(slug)
                    except ValueError as exc:
                        raise ValueError(
                            f"Invalid module slug for sheet {sheet_name!r}: {exc}"
                        ) from exc
            groups.append((module, None, modular))
            locales = {}
            for index, header in enumerate(headers):
                if not header or header in {"key", "description", "tags"}:
                    continue
                reject_nul(header, field="locale", context=f"Sheet {sheet_name!r}: ")
                try:
                    locales[validate_locale_code(header)] = index
                except ValueError as exc:
                    raise ValueError(f"Invalid locale column {header!r}: {exc}") from exc
            for number, row in enumerate(iterator, start=2):
                if not row or row[key_idx] is None:
                    continue
                key = _text_value(row[key_idx]).strip()
                if not key:
                    continue
                context = f"Sheet {sheet_name!r} row {number}: "
                reject_nul(key, field="key", key=key, context=context)
                if len(key) > KEY_MAX_LENGTH:
                    raise ValueError(
                        f"{context}key {key_preview(key)} must be at most {KEY_MAX_LENGTH} characters"
                    )
                description = (
                    ensure_max_length(
                        _text_value(row[desc_idx]).strip(),
                        field="description",
                        limit=TEXT_MAX_LENGTH,
                    )
                    if desc_idx is not None and row[desc_idx]
                    else None
                ) or None
                names = ()
                if tags_idx is not None and row[tags_idx]:
                    names = tuple(
                        ensure_max_length(
                            value.strip(),
                            field="tag name",
                            limit=TAG_NAME_MAX_LENGTH,
                        )
                        for value in _text_value(row[tags_idx]).split(",")
                        if value.strip()
                    )
                values = {locale: _text_value(row[index]) for locale, index in locales.items()}
                label = f"{module}/{key}" if module is not None else key
                rows.append(
                    ImportRow(
                        key=key,
                        values=values,
                        module=module,
                        scoped=modular,
                        label=label,
                        excel=True,
                        description=description,
                        tag_names=names,
                        context=context,
                    )
                )
        return import_rows(
            db, project, rows, groups=groups, dry_run=dry_run, status=status, report_orphans=False
        )
    finally:
        wb.close()
