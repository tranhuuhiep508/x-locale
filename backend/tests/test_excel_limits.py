"""Resource preflight tests; these also run offline with pytest --noconftest."""

import io
from types import SimpleNamespace
from xml.sax.saxutils import escape
from zipfile import ZIP_BZIP2, ZIP_DEFLATED, ZIP_LZMA, ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.formula.translate import Translator

from app import excel_limits
from app.excel_limits import validate_workbook
from app.limits import TEXT_MAX_LENGTH


def archive_bytes(parts):
    raw = io.BytesIO()
    with ZipFile(raw, "w", ZIP_DEFLATED) as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    return raw.getvalue()


def worksheet(body):
    return f'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">{body}</worksheet>'


def shared_formula_sheet(formula, count, *, index="0"):
    rows = []
    for number in range(1, count + 1):
        text = escape(formula) if number == 1 else ""
        rows.append(
            f'<row r="{number}"><c r="A{number}"><f t="shared" si="{index}">{text}</f><v>0</v></c></row>'
        )
    return worksheet(f"<sheetData>{''.join(rows)}</sheetData>")


def shared_string_table(values, *, rich=False):
    entries = []
    for value in values:
        if rich:
            middle = len(value) // 2
            content = (
                f"<r><t>{escape(value[:middle])}</t></r><r><t>{escape(value[middle:])}</t></r>"
            )
        else:
            content = f'<t xml:space="preserve">{escape(value)}</t>'
        entries.append(f"<si>{content}</si>")
    return f'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">{"".join(entries)}</sst>'


def shared_string_sheet(indices):
    rows = []
    for number, index in enumerate(indices, start=1):
        value = f"<v>{escape(str(index))}</v>" if index is not None else ""
        rows.append(f'<row r="{number}"><c r="A{number}" t="s">{value}</c></row>')
    return worksheet(f"<sheetData>{''.join(rows)}</sheetData>")


@pytest.mark.parametrize(
    "setting,parts,message",
    [
        ("EXCEL_MAX_MEMBERS", {"a": b"", "b": b"", "c": b""}, "archive entries"),
        ("EXCEL_MAX_MEMBER_BYTES", {"a": b"xxx"}, "expanded member bytes"),
        ("EXCEL_MAX_EXPANDED_BYTES", {"a": b"xx", "b": b"x"}, "expanded archive bytes"),
    ],
)
def test_archive_metadata_limits(monkeypatch, setting, parts, message):
    raw = archive_bytes(parts)
    monkeypatch.setattr(excel_limits, setting, 2)

    def fail_open(*_args, **_kwargs):
        raise AssertionError("Archive member opened before metadata validation")

    monkeypatch.setattr(ZipFile, "open", fail_open)
    with pytest.raises(ValueError, match=message):
        validate_workbook(raw)


@pytest.mark.parametrize("setting", ["EXCEL_MAX_MEMBER_BYTES", "EXCEL_MAX_EXPANDED_BYTES"])
def test_actual_stream_bytes_are_checked(monkeypatch, setting):
    raw = archive_bytes({"binary": b"x"})
    original = ZipFile.infolist

    def understated(archive):
        members = original(archive)
        for member in members:
            member.file_size = 1
        return members

    # The stream deliberately disagrees with the directory metadata.
    monkeypatch.setattr(ZipFile, "infolist", understated)
    monkeypatch.setattr(ZipFile, "open", lambda *_args: io.BytesIO(b"xxx"))
    monkeypatch.setattr(excel_limits, setting, 2)
    with pytest.raises(ValueError, match="expanded .* bytes"):
        validate_workbook(raw)


def workbook_parts(sheets=None, shared=None):
    """Build real package metadata around the supplied worksheet XML."""
    from xml.etree import ElementTree as ET

    sheets = sheets or {"xl/worksheets/sheet1.xml": worksheet("")}
    workbook = Workbook()
    for i in range(1, len(sheets)):
        workbook.create_sheet(f"sheet{i}")
    raw = io.BytesIO()
    workbook.save(raw)
    with ZipFile(raw) as archive:
        parts = {n: archive.read(n) for n in archive.namelist()}
    rel_path = "xl/_rels/workbook.xml.rels"
    relationships = ET.fromstring(parts[rel_path])
    sheet_rels = [r for r in relationships if r.get("Type").endswith("/worksheet")]
    for relationship, path in zip(sheet_rels, sheets, strict=True):
        old = relationship.get("Target").lstrip("/")
        del parts[old]
        relationship.set("Target", "/" + path)
    parts[rel_path] = ET.tostring(relationships)
    parts.update(sheets)
    if shared is not None:
        parts["xl/sharedStrings.xml"] = shared
        parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(
            b"</Types>",
            b'<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>',
        )
    return parts


@pytest.mark.parametrize("kind", ["inline", "shared", "rich", "utf16", "alternate-target"])
def test_oversized_cell_text_is_rejected_before_loading(kind):
    text = "x" * (TEXT_MAX_LENGTH + 1)
    shared = (
        shared_string_table([text], rich=kind == "rich") if kind in {"shared", "rich"} else None
    )
    xml = (
        shared_string_sheet([0])
        if shared
        else worksheet(
            f'<sheetData><row><c t="inlineStr"><is><t>{text}</t></is></c></row></sheetData>'
        )
    )
    if kind == "utf16":
        xml = xml.encode("utf-16")
    path = "xl/worksheet.data" if kind == "alternate-target" else "xl/worksheets/sheet1.xml"
    with pytest.raises(ValueError, match="cell/shared-string characters"):
        validate_workbook(archive_bytes(workbook_parts({path: xml}, shared)))


@pytest.mark.parametrize(
    "body,message",
    [
        ('<dimension ref="A1:A100001"/>', "row slots"),
        ('<dimension ref="A1:IW1"/>', "columns"),
        ('<dimension ref="A1:K100000"/>', "cell slots"),
        ('<sheetData><row r="100001"/></sheetData>', "row slots"),
        ('<sheetData><row><c r="IW1"/></row></sheetData>', "columns"),
        ('<sheetData><row><c r="A100001"/></row></sheetData>', "row slots"),
        ('<sheetData><row r="0"/></sheetData>', "row coordinate"),
        ('<dimension ref="invalid"/>', "cell coordinate"),
    ],
)
def test_sparse_coordinates_and_dimensions_are_bounded(body, message):
    with pytest.raises(ValueError, match=message):
        validate_workbook(archive_bytes(workbook_parts({"sheet.xml": worksheet(body)})))


@pytest.mark.parametrize(
    "body",
    [
        "<row><other><v>1</v></other></row>",
        "<row><c><row><c/></row></c></row>",
        '<row><c><f>1</f><f t="shared" si="0">2</f></c></row>',
        '<row><c t="s"><v>0</v><v>1</v></c></row>',
        '<row><c t="s"><v><nested>0</nested></v></c></row>',
        '<row><c><unused><f t="shared" si="0">1</f></unused></c></row>',
        '<row><c><fake:f xmlns:fake="urn:fake">1</fake:f></c></row>',
        "<row><c/><extLst/><extLst/></row>",
        "<row><extLst/><c/></row>",
    ],
)
def test_malformed_worksheet_records_are_rejected(body):
    with pytest.raises(ValueError):
        validate_workbook(
            archive_bytes(
                workbook_parts({"sheet.xml": worksheet(f"<sheetData>{body}</sheetData>")})
            )
        )


@pytest.mark.parametrize("root", ["wrapper", "externalLink"])
def test_worksheet_root_is_required(root):
    xml = worksheet("<sheetData><row/></sheetData>").replace("worksheet", root)
    with pytest.raises(ValueError, match="worksheet root"):
        validate_workbook(archive_bytes(workbook_parts({"sheet.rels": xml})))


@pytest.mark.parametrize("resource", ["rows", "cells", "sheets"])
def test_budgets_are_cumulative(monkeypatch, resource):
    setting, limit, body = {
        "rows": ("EXCEL_MAX_ROWS", 3, '<dimension ref="A1:A2"/>'),
        "cells": ("EXCEL_MAX_CELLS", 100, '<dimension ref="A1:BH1"/>'),
        "sheets": ("EXCEL_MAX_SHEETS", 1, ""),
    }[resource]
    monkeypatch.setattr(excel_limits, setting, limit)
    parts = workbook_parts({f"sheet{i}.xml": worksheet(body) for i in range(2)})
    with pytest.raises(ValueError, match="row slots|cell slots|worksheets"):
        validate_workbook(archive_bytes(parts))


@pytest.mark.parametrize("body", ['<row r="1"/>' * 1001, "<row>" + '<c r="A1"/>' * 1001 + "</row>"])
def test_repeated_coordinates_count_as_work(monkeypatch, body):
    monkeypatch.setattr(excel_limits, "EXCEL_MAX_ROWS", 1000)
    monkeypatch.setattr(excel_limits, "EXCEL_MAX_CELLS", 1000)
    parts = workbook_parts({"sheet.xml": worksheet(f"<sheetData>{body}</sheetData>")})
    with pytest.raises(ValueError, match="total .* slots"):
        validate_workbook(archive_bytes(parts))


@pytest.mark.parametrize("index", [-1, 1, "bad", "0.0"])
def test_invalid_shared_string_indices_are_rejected(index):
    parts = workbook_parts(
        {"sheet.xml": shared_string_sheet([index])}, shared_string_table(["valid"])
    )
    with pytest.raises(ValueError, match="shared-string index"):
        validate_workbook(archive_bytes(parts))


@pytest.mark.parametrize("table_first", [False, True])
@pytest.mark.parametrize("rich", [False, True])
def test_shared_string_expansion_is_bounded(monkeypatch, table_first, rich):
    monkeypatch.setattr(excel_limits, "EXCEL_MAX_DECODED_CHARACTERS", 100)
    parts = workbook_parts(
        {"first.xml": shared_string_sheet([0, 0]), "second.xml": shared_string_sheet([0, 0])},
        shared_string_table([" " + "x" * 19], rich=rich),
    )
    if table_first:
        parts = dict(reversed(list(parts.items())))
    with pytest.raises(ValueError, match="decoded characters"):
        validate_workbook(archive_bytes(parts))


@pytest.mark.parametrize(
    "formula",
    ['"literal"', "SUM(A1:B2)", "$A1+A$1+$A$1", "SUM(1:2,A:B)", "'Sheet 1'!A1", "A1", "999"],
)
def test_shared_formulas_remain_supported(formula):
    assert len(
        Translator("=" + formula, "A1").translate_formula("IV100000")
    ) <= excel_limits._shared_formula_bound(formula)
    validate_workbook(
        archive_bytes(workbook_parts({"sheet.xml": shared_formula_sheet(formula, 3)}))
    )


@pytest.mark.parametrize("redefine", [False, True])
def test_shared_formula_expansion_is_bounded(monkeypatch, redefine):
    monkeypatch.setattr(excel_limits, "EXCEL_MAX_DECODED_CHARACTERS", 100)
    xml = shared_formula_sheet('"' + "x" * 40 + '"', 3)
    if redefine:
        xml = xml.replace('<f t="shared" si="0"></f>', '<f t="shared" si="0">1</f>')
    with pytest.raises(ValueError, match="decoded characters"):
        validate_workbook(archive_bytes(workbook_parts({"sheet.xml": xml})))


@pytest.mark.parametrize("rich", [False, True])
def test_maximum_length_text_and_formatting_whitespace_are_supported(rich):
    table = shared_string_table(["x" * TEXT_MAX_LENGTH], rich=rich).replace("><", ">\n<")
    validate_workbook(archive_bytes(workbook_parts(shared=table)))


@pytest.mark.parametrize(
    "xml",
    [
        '<!DOCTYPE a [<!ENTITY x "expanded">]><a>&x;</a>',
        '<!DOCTYPE a SYSTEM "file:///etc/passwd"><a/>',
        "<a>",
    ],
)
def test_unsafe_xml_is_rejected_even_in_ignored_parts(xml):
    parts = workbook_parts()
    parts["ignored.xml"] = xml
    with pytest.raises(ValueError, match="DTD|archive/XML"):
        validate_workbook(archive_bytes(parts))


@pytest.mark.parametrize(
    "path", ["/sheet.xml", "xl/../sheet.xml", "xl/./sheet.xml", "../sheet.xml", "xl\\sheet.xml"]
)
def test_noncanonical_archive_paths_are_rejected(path):
    with pytest.raises(ValueError, match="member path"):
        validate_workbook(archive_bytes({path: worksheet("")}))


@pytest.mark.parametrize(
    "kind",
    [
        "name",
        "relationship",
        "target",
        "renamed-rel",
        "duplicate-rel",
        "renamed-sheet",
        "unqualified-id",
        "nested-rel",
        "type",
        "external",
        "missing-target",
    ],
)
def test_ambiguous_or_unsupported_packages_are_rejected(kind):
    from xml.etree import ElementTree as ET

    parts = workbook_parts({"first.xml": worksheet(""), "second.xml": worksheet("")})
    root = ET.fromstring(parts["xl/workbook.xml"])
    sheets = root.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheets")
    rel_path = "xl/_rels/workbook.xml.rels"
    rels = ET.fromstring(parts[rel_path])
    sheet_rels = [r for r in rels if r.get("Type").endswith("/worksheet")]
    rid = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    if kind == "name":
        sheets[1].set("name", sheets[0].get("name"))
    elif kind == "relationship":
        sheets[1].set(rid, sheets[0].get(rid))
    elif kind == "target":
        sheet_rels[1].set("Target", sheet_rels[0].get("Target"))
    elif kind == "renamed-rel":
        sheet_rels[0].tag = "renamed"
    elif kind == "duplicate-rel":
        rels.append(ET.fromstring(ET.tostring(sheet_rels[0])))
    elif kind == "renamed-sheet":
        sheets[0].tag = "renamed"
    elif kind == "unqualified-id":
        sheets[0].set("id", sheets[0].attrib.pop(rid))
    elif kind == "nested-rel":
        wrapper = ET.SubElement(rels, "unused")
        wrapper.append(sheet_rels[0])
    elif kind == "type":
        sheet_rels[0].set("Type", "example/styles")
    elif kind == "external":
        sheet_rels[0].set("TargetMode", "External")
    else:
        del parts["first.xml"]
    parts["xl/workbook.xml"], parts[rel_path] = ET.tostring(root), ET.tostring(rels)
    with pytest.raises(ValueError):
        validate_workbook(archive_bytes(parts))


@pytest.mark.parametrize("layout", ["flat", "modular"])
def test_twenty_thousand_key_workbooks_fit(layout):
    workbook = Workbook(write_only=True)
    sheets = [
        workbook.create_sheet(name)
        for name in (["strings"] if layout == "flat" else ["auth", "home", "common"])
    ]
    for sheet in sheets:
        sheet.append(["key", "vi", "en", "fr"])
    for number in range(20000):
        sheets[number % len(sheets)].append(
            [f"key.{number}", f"Source {number}", f"English {number}", f"French {number}"]
        )
    raw = io.BytesIO()
    workbook.save(raw)
    validate_workbook(raw.getvalue())


def workbook_with_external_cache():
    relationship = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    workbook = Workbook()
    workbook.active.append(["key", "vi"])
    workbook.active.append(["hello", "Hello"])
    raw = io.BytesIO()
    workbook.save(raw)
    with ZipFile(raw) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["xl/workbook.xml"] = parts["xl/workbook.xml"].replace(
        b"</workbook>",
        f'<externalReferences><externalReference xmlns:r="{relationship}" r:id="external"/></externalReferences></workbook>'.encode(),
    )
    parts["xl/_rels/workbook.xml.rels"] = parts["xl/_rels/workbook.xml.rels"].replace(
        b"</Relationships>",
        f'<Relationship Id="external" Type="{relationship}/externalLink" Target="/xl/externalLinks/externalLink1.xml"/></Relationships>'.encode(),
    )
    parts["xl/externalLinks/externalLink1.xml"] = (
        f'<externalLink xmlns="{namespace}" xmlns:r="{relationship}">'
        '<externalBook r:id="rId1"><sheetNames><sheetName val="Sheet1"/></sheetNames>'
        '<sheetDataSet><sheetData sheetId="0"><row r="1"><cell r="A1" t="str">'
        "<v>cached text</v></cell></row></sheetData></sheetDataSet></externalBook></externalLink>"
    ).encode()
    parts["xl/externalLinks/_rels/externalLink1.xml.rels"] = (
        f'<Relationships><Relationship Id="rId1" Type="{relationship}/externalLinkPath" '
        'Target="file:///example.xlsx" TargetMode="External"/></Relationships>'
    ).encode()
    return parts


@pytest.mark.parametrize("cell_type", ["str", "s"])
def test_external_workbook_cache_is_supported(monkeypatch, cell_type):
    from app import excel

    parts = workbook_with_external_cache()
    path = "xl/externalLinks/externalLink1.xml"
    parts[path] = parts[path].replace(b't="str"', f't="{cell_type}"'.encode())
    raw = archive_bytes(parts)
    validate_workbook(raw)
    workbook = load_workbook(io.BytesIO(raw), read_only=True)
    try:
        assert list(workbook.active.values) == [("key", "vi"), ("hello", "Hello")]
        assert len(workbook._external_links) == 1
    finally:
        workbook.close()

    def fail_external_link(*_args, **_kwargs):
        raise AssertionError("Import must not load external workbook caches")

    monkeypatch.setattr("openpyxl.reader.workbook.read_external_link", fail_external_link)
    monkeypatch.setattr(excel, "import_rows", lambda _db, _project, rows, **_kwargs: rows)
    rows = excel.import_workbook(None, SimpleNamespace(layout="flat"), raw)
    assert [(row.key, row.values) for row in rows] == [("hello", {"vi": "Hello"})]


@pytest.mark.parametrize("resource", ["elements", "depth", "bytes"])
def test_external_caches_still_obey_global_limits(monkeypatch, resource):
    parts = workbook_with_external_cache()
    path = "xl/externalLinks/externalLink1.xml"
    if resource == "elements":
        monkeypatch.setattr(excel_limits, "EXCEL_MAX_CELLS", 1)
        message = "XML elements"
    elif resource == "depth":
        parts[path] = parts[path].replace(b"cached text", b"<a>" * 129 + b"</a>" * 129)
        message = "nesting depth"
    else:
        monkeypatch.setattr(excel_limits, "EXCEL_MAX_MEMBER_BYTES", 4096)
        parts[path] = parts[path].replace(b"cached text", b"x" * 8192)
        message = "expanded member bytes"
    with pytest.raises(ValueError, match=message):
        validate_workbook(archive_bytes({path: parts[path]}))


@pytest.mark.parametrize("relationships_first", [False, True])
def test_external_link_root_cannot_hide_a_worksheet_target(relationships_first):
    parts = workbook_with_external_cache()
    parts["xl/_rels/workbook.xml.rels"] = parts["xl/_rels/workbook.xml.rels"].replace(
        b"/xl/worksheets/sheet1.xml", b"/xl/externalLinks/externalLink1.xml"
    )
    if relationships_first:
        parts = dict(reversed(list(parts.items())))
    with pytest.raises(ValueError, match="worksheet root"):
        validate_workbook(archive_bytes(parts))


@pytest.mark.parametrize("suffix", ["xml", "rels"])
@pytest.mark.parametrize("worksheet_first", [False, True])
def test_external_cache_classification_cannot_hide_shared_string_expansion(suffix, worksheet_first):
    parts = workbook_with_external_cache()
    path = f"xl/worksheets/disguised.{suffix}"
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rows = "".join(
        f'<row r="{number}"><cell r="A{number}" t="s"><v>0</v></cell></row>'
        for number in range(1, 5)
    )
    parts[path] = (
        f'<externalLink xmlns="{namespace}"><sheetData>{rows}</sheetData></externalLink>'.encode()
    )
    parts["xl/_rels/workbook.xml.rels"] = parts["xl/_rels/workbook.xml.rels"].replace(
        b"/xl/worksheets/sheet1.xml", f"/{path}".encode()
    )
    parts["xl/sharedStrings.xml"] = shared_string_table(["x" * 40]).encode()
    parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(
        b"</Types>",
        b'<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>',
    )
    worksheet = parts.pop(path)
    parts = {path: worksheet, **parts} if worksheet_first else {**parts, path: worksheet}
    raw = archive_bytes(parts)
    workbook = load_workbook(io.BytesIO(raw), read_only=True, keep_links=False)
    try:
        assert sum(len(row[0]) for row in workbook.active.values) == 160
    finally:
        workbook.close()
    with pytest.raises(ValueError, match="worksheet root"):
        validate_workbook(raw)


def test_invalid_zip_is_a_validation_error():
    with pytest.raises(ValueError, match="archive/XML"):
        validate_workbook(b"not a zip")


@pytest.mark.parametrize("compression", [ZIP_BZIP2, ZIP_LZMA])
def test_compression_with_unbounded_decoder_output_is_rejected(monkeypatch, compression):
    raw = io.BytesIO()
    with ZipFile(raw, "w", compression) as archive:
        archive.writestr("part.xml", "<a/>")

    def fail_open(*_args, **_kwargs):
        raise AssertionError("Unsafe decompressor started")

    monkeypatch.setattr(ZipFile, "open", fail_open)
    with pytest.raises(ValueError, match="stored or deflate"):
        validate_workbook(raw.getvalue())


def test_duplicate_archive_member_names_are_rejected():
    raw = io.BytesIO()
    with ZipFile(raw, "w") as archive:
        archive.writestr("part.xml", "<a/>")
        with pytest.warns(UserWarning, match="Duplicate name"):
            archive.writestr("part.xml", "<b/>")
    with pytest.raises(ValueError, match="duplicate member"):
        validate_workbook(raw.getvalue())


def test_xml_nesting_is_bounded():
    with pytest.raises(ValueError, match="nesting depth"):
        validate_workbook(archive_bytes({"deep.xml": "<a>" * 129 + "</a>" * 129}))


def test_xml_element_count_is_bounded_across_parts(monkeypatch):
    monkeypatch.setattr(excel_limits, "EXCEL_MAX_CELLS", 1)
    with pytest.raises(ValueError, match="XML elements"):
        validate_workbook(
            archive_bytes({"one.xml": "<a><b/><b/><b/><b/></a>", "two.xml": "<a><b/><b/><b/></a>"})
        )


@pytest.mark.parametrize(
    "entry",
    [
        "<si><si><t>nested</t></si></si>",
        '<foreign:si xmlns:foreign="urn:test"><t>text</t></foreign:si>',
    ],
)
def test_malformed_shared_string_entries_are_rejected(entry):
    table = shared_string_table([]).replace("</sst>", entry + "</sst>")
    with pytest.raises(ValueError, match="shared-string entry structure"):
        validate_workbook(archive_bytes(workbook_parts(shared=table)))


def test_empty_shared_values_and_ignored_formula_caches_are_supported():
    sheet = shared_string_sheet([0, "", None]).replace(
        "</sheetData>", '<row><c t="s"><f>1+1</f><v>invalid-cache</v></c></row></sheetData>'
    )
    validate_workbook(
        archive_bytes(workbook_parts({"sheet.xml": sheet}, shared_string_table([""])))
    )


def test_ignored_binary_parts_do_not_receive_xml_checks():
    parts = workbook_parts()
    parts["embedding.bin"] = b"\x00\x01binary data"
    validate_workbook(archive_bytes(parts))


@pytest.mark.parametrize("wrapper", ["is", "extLst"])
def test_nested_values_cannot_replace_shared_string_indices(monkeypatch, wrapper):
    monkeypatch.setattr(excel_limits, "EXCEL_MAX_DECODED_CHARACTERS", 100)
    rows = "".join(
        f'<row r="{number}"><c r="A{number}" t="s"><v>0</v>'
        f"<{wrapper}><v>1</v></{wrapper}></c></row>"
        for number in range(1, 5)
    )
    raw = archive_bytes(
        workbook_parts(
            {"sheet.xml": worksheet(f"<sheetData>{rows}</sheetData>")},
            shared_string_table(["x" * 40, "y"]),
        )
    )
    workbook = load_workbook(io.BytesIO(raw), read_only=True, keep_links=False)
    try:
        assert sum(len(row[0]) for row in workbook.active.values) == 160
    finally:
        workbook.close()
    with pytest.raises(ValueError, match="cell value/formula structure"):
        validate_workbook(raw)


def test_unqualified_formula_cannot_hide_shared_string_expansion(monkeypatch):
    monkeypatch.setattr(excel_limits, "EXCEL_MAX_DECODED_CHARACTERS", 100)
    rows = "".join(
        f'<row r="{number}"><c r="A{number}" t="s"><f xmlns="">1</f><v>0</v></c></row>'
        for number in range(1, 5)
    )
    raw = archive_bytes(
        workbook_parts(
            {"sheet.xml": worksheet(f"<sheetData>{rows}</sheetData>")},
            shared_string_table([" " + "x" * 39]),
        )
    )
    workbook = load_workbook(io.BytesIO(raw), read_only=True, keep_links=False)
    try:
        assert sum(len(row[0]) for row in workbook.active.values) == 160
    finally:
        workbook.close()
    with pytest.raises(ValueError, match="spreadsheet element namespace"):
        validate_workbook(raw)


@pytest.mark.parametrize("element", ["row", "c", "v", "si"])
def test_unqualified_structural_elements_are_rejected(element):
    xml = worksheet('<sheetData><row><c t="s"><v>0</v></c></row></sheetData>')
    table = shared_string_table(["text"])
    if element == "si":
        table = table.replace("<si>", '<si xmlns="">')
    else:
        xml = xml.replace(f"<{element}", f'<{element} xmlns=""', 1)
    with pytest.raises(ValueError, match="spreadsheet element namespace"):
        validate_workbook(archive_bytes(workbook_parts({"sheet.xml": xml}, table)))
