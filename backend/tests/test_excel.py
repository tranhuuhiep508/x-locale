"""Comprehensive unit and integration tests for app.excel."""

from __future__ import annotations

import io
import uuid
from datetime import datetime
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import event

from app.excel import (
    FLAT_SHEET,
    UNASSIGNED_SHEET,
    _module_sheet_names,
    build_template_workbook,
    build_workbook,
)
from app.models import (
    Module,
    Project,
    ProjectLayout,
    StringEntry,
    Tag,
    Translation,
    TranslationStatus,
)
from tests.helpers import json_upload, publish_strings
from tests.test_import_jobs import poll


def test_module_sheet_names_reserve_short_names_and_bound_collision_suffixes():
    prefix = "a" * 31
    slugs = [prefix + str(number) for number in range(15)] + [prefix, "a" * 29 + "_1"]
    names = _module_sheet_names(slugs)
    assert names[prefix] == prefix
    assert names["a" * 29 + "_1"] == "a" * 29 + "_1"
    assert len({name.casefold() for name in names.values()}) == len(slugs)
    assert all(len(name) <= 31 for name in names.values())
    assert names == _module_sheet_names(list(reversed(slugs)))
    case_names = _module_sheet_names(["auth", "Auth", "_meta", "_unassigned"])
    assert len({name.casefold() for name in case_names.values()}) == 4
    assert not {"_meta", "_unassigned"} & {name.casefold() for name in case_names.values()}


@pytest.mark.parametrize("stage", ["draft", "public"])
def test_module_mapping_uses_the_selected_stage(client, stage):
    project = _make_project(client, layout="modular")
    root = f"/api/projects/{project['id']}"
    module_ids = []
    for slug in ("a" * 128, "b" * 128):
        response = client.post(f"{root}/modules", json={"slug": slug, "name": slug})
        assert response.status_code == 201, response.text
        module_ids.append(response.json()["id"])
    entry = client.post(
        f"{root}/strings",
        json={"key": "title", "source_text": "Source", "module_id": module_ids[0]},
    )
    assert entry.status_code == 201, entry.text
    sid = entry.json()["id"]
    publish_strings(client, project["id"], [sid])
    moved = client.patch(f"{root}/strings/{sid}", json={"module_id": module_ids[1]})
    assert moved.status_code == 200, moved.text
    exported = client.get(f"{root}/export", params={"format": "xlsx", "stage": stage})
    assert exported.status_code == 200, exported.text
    wb = load_workbook(io.BytesIO(exported.content))
    mappings = [value for key, value in wb["_meta"].values if key.startswith("module_slug:")]
    assert mappings == [("a" if stage == "public" else "b") * 128]
    wb.close()


@pytest.mark.parametrize("stage", ["draft", "public"])
@pytest.mark.parametrize("allow_duplicates", [False, True])
@pytest.mark.parametrize("as_job", [False, True])
def test_long_module_slugs_roundtrip_preserves_identity(client, stage, allow_duplicates, as_job):
    project = _make_project(client, layout="modular")
    root = f"/api/projects/{project['id']}"
    assert client.patch(root, json={"allow_duplicate_keys": allow_duplicates}).status_code == 200
    slugs = ["a" * 31, "a" * 32, "a" * 31 + "b", "b" * 128, "a" * 29 + "_1"]
    payload = {"modules": {}, "unassigned": {"vi": {"loose": "Unassigned"}}}
    for number, slug in enumerate(slugs):
        key = "title" if allow_duplicates else f"title{number}"
        payload["modules"][slug] = {"vi": {key: "Source"}, "en": {key: "Original"}}
    seeded = client.post(f"{root}/import", files=json_upload(payload))
    assert seeded.status_code == 200, seeded.text
    before = client.get(f"{root}/strings").json()["items"]
    identities = {(row["id"], row["module_id"], row["key"]) for row in before}
    modules_before = client.get(f"{root}/modules").json()
    if stage == "public":
        publish_strings(client, project["id"], [row["id"] for row in before])
    exported = client.get(f"{root}/export", params={"format": "xlsx", "stage": stage})
    assert exported.status_code == 200, exported.text
    wb = load_workbook(io.BytesIO(exported.content))
    assert wb["_meta"].sheet_state == "hidden"
    mapping = {
        key.removeprefix("module_slug:"): value
        for key, value in wb["_meta"].iter_rows(min_row=2, values_only=True)
        if key.startswith("module_slug:")
    }
    assert set(mapping.values()) == {*slugs, UNASSIGNED_SHEET}
    assert mapping["a" * 31] == "a" * 31
    assert mapping["a" * 29 + "_1"] == "a" * 29 + "_1"
    assert all(len(name) <= 31 for name in wb.sheetnames)

    def upload(raw):
        response = client.post(
            f"{root}/import",
            files={"file": ("roundtrip.xlsx", raw)},
            headers={"Prefer": "respond-async"} if as_job else {},
        )
        if as_job:
            assert response.status_code == 202, response.text
            job = poll(client, project, "xlsx", response.json()["job_id"])
            assert job["status"] == "completed", job
            return job["result"]
        assert response.status_code == 200, response.text
        return response.json()

    assert upload(exported.content)["created"] == 0
    edited_name = next(name for name, slug in mapping.items() if slug == "b" * 128)
    wb[edited_name].cell(2, 5, "Edited translation")
    wb[edited_name].cell(2, 4, "Edited source")
    edited = io.BytesIO()
    wb.save(edited)
    preview = client.post(
        f"{root}/import",
        params={"dry_run": True},
        files={"file": ("preview.xlsx", edited.getvalue())},
    )
    assert preview.status_code == 200, preview.text
    edited_key = wb[edited_name].cell(2, 1).value
    assert preview.json()["diff"]["update"][0]["key"] == f"{'b' * 128}/{edited_key}"
    wb[edited_name].cell(2, 4, "Source")
    edited = io.BytesIO()
    wb.save(edited)
    wb.close()
    assert upload(edited.getvalue())["created"] == 0
    after = client.get(f"{root}/strings").json()["items"]
    assert {(row["id"], row["module_id"], row["key"]) for row in after} == identities
    assert client.get(f"{root}/modules").json() == modules_before
    changed = next(row for row in after if row["module_slug"] == "b" * 128)
    assert (
        next(t["value"] for t in changed["translations"] if t["locale"] == "en")
        == "Edited translation"
    )
    if stage == "public":
        public = client.get(f"{root}/export", params={"stage": "public"}).json()
        assert public["modules"]["b" * 128]["en"][changed["key"]] == "Original"


@pytest.mark.parametrize("legacy", [False, True])
def test_modular_template_mapping_and_legacy_workbooks(client, legacy):
    project = _make_project(client, layout="modular")
    root = f"/api/projects/{project['id']}"
    template = client.get(f"{root}/import-template", params={"format": "xlsx"})
    assert template.status_code == 200, template.text
    wb = load_workbook(io.BytesIO(template.content))
    mapping = {
        key: value
        for key, value in wb["_meta"].iter_rows(min_row=2, values_only=True)
        if key.startswith("module_slug:")
    }
    assert mapping == {
        "module_slug:auth": "auth",
        "module_slug:common": "common",
        "module_slug:_unassigned": UNASSIGNED_SHEET,
    }
    wb.create_sheet("Notes").append(["Read me"])
    if legacy:
        for number in range(wb["_meta"].max_row, 1, -1):
            if str(wb["_meta"].cell(number, 1).value).startswith("module_slug:"):
                wb["_meta"].delete_rows(number)
    raw = io.BytesIO()
    wb.save(raw)
    wb.close()
    imported = client.post(f"{root}/import", files={"file": ("template.xlsx", raw.getvalue())})
    assert imported.status_code == 200, imported.text
    assert imported.json()["created"] == 5
    assert {m["slug"] for m in client.get(f"{root}/modules").json()} == {"auth", "common"}
    rows = client.get(f"{root}/strings").json()["items"]
    assert next(row for row in rows if row["key"] == "hello")["module_id"] is None


@pytest.mark.parametrize(
    "mappings,error",
    [
        ([("auth", "new-module")], "Missing module mapping"),
        (
            [("auth", "new-module"), ("auth", "new-module"), ("home", "home")],
            "Duplicate module mapping",
        ),
        ([("auth", "new-module"), ("home", "new-module")], "Conflicting module mappings"),
        ([("auth", "new-module"), ("home", "INVALID")], "Invalid module slug"),
        ([("auth", "new-module"), ("home", "a" * 129)], "at most 128"),
        ([("auth", "new-module"), ("home", None)], "Invalid module slug mapping"),
        ([("auth", "new-module"), ("home", 123)], "Invalid module slug mapping"),
        ([("auth", "new-module"), ("renamed", "home")], "unknown data sheet"),
        ([("auth", "new-module"), ("home", UNASSIGNED_SHEET)], "Invalid unassigned"),
        ([("auth", "new-module"), ("home", "home"), ("_meta", "meta")], "unknown data sheet"),
        (
            [("auth", "new-module"), ("home", "home"), (UNASSIGNED_SHEET, "assigned")],
            "Invalid unassigned",
        ),
    ],
)
@pytest.mark.parametrize("dry_run", [False, True])
def test_invalid_module_mappings_rejected_before_writes(client, mappings, error, dry_run):
    project = _make_project(client, layout="modular")
    root = f"/api/projects/{project['id']}"
    wb = Workbook()
    wb.remove(wb.active)
    for name in ("auth", "home", UNASSIGNED_SHEET):
        ws = wb.create_sheet(name)
        ws.append(["key", "vi"])
        ws.append([name, "Source"])
    meta = wb.create_sheet("_meta")
    meta.append(["key", "value"])
    for name, slug in mappings:
        meta.append([f"module_slug:{name}", slug])
    raw = io.BytesIO()
    wb.save(raw)
    wb.close()
    response = client.post(
        f"{root}/import",
        params={"dry_run": dry_run},
        files={"file": ("invalid.xlsx", raw.getvalue())},
    )
    assert response.status_code == 400, response.text
    assert error in response.json()["detail"]
    assert client.get(f"{root}/strings").json()["total"] == 0
    assert client.get(f"{root}/modules").json() == []


def _make_project(client, name="Excel Test Project", layout="flat", base="vi", targets=None):
    if targets is None:
        targets = ["en"]
    r = client.post(
        "/api/projects",
        json={
            "name": name,
            "layout": layout,
            "base_language": base,
            "target_languages": targets,
        },
    )
    assert r.status_code == 201, r.text
    return r.json()


@pytest.mark.parametrize(("count", "expected_queries"), [(4, 1), (501, 2)])
@pytest.mark.parametrize("tagged", [False, True])
def test_excel_export_batches_tag_queries(
    client, session_factory, database_engine, count, expected_queries, tagged
):
    project = _make_project(client, targets=["en", "ja"])
    pid = uuid.UUID(project["id"])
    with session_factory() as db:
        tags = [Tag(project_id=pid, name=name) for name in ("z-last", "a-first")] if tagged else []
        db.add_all(
            StringEntry(
                project_id=pid,
                key=f"key{i:04}",
                key_policy_allows_duplicates=False,
                source_text=f"Source {i}",
                description=f"Description {i}",
                status=TranslationStatus.draft,
                tags=tags if i % 2 == 0 else [],
                translations=[
                    Translation(locale="en", value=f"English {i}"),
                    Translation(locale="ja", value=f"Japanese {i}"),
                ],
            )
            for i in range(count)
        )
        db.commit()

    tag_queries = []

    def observe(_conn, _cursor, statement, _parameters, _context, _many):
        if statement.lstrip().upper().startswith("SELECT") and "string_tags" in statement:
            tag_queries.append(statement)

    event.listen(database_engine, "before_cursor_execute", observe)
    try:
        response = client.get(
            f"/api/projects/{pid}/export", params={"format": "xlsx", "stage": "draft"}
        )
    finally:
        event.remove(database_engine, "before_cursor_execute", observe)
    assert response.status_code == 200, response.text
    assert len(tag_queries) == expected_queries
    wb = load_workbook(io.BytesIO(response.content))
    try:
        assert list(wb[FLAT_SHEET].values) == [
            ("key", "description", "tags", "vi", "en", "ja"),
            *[
                (
                    f"key{i:04}",
                    f"Description {i}",
                    "a-first,z-last" if tagged and i % 2 == 0 else None,
                    f"Source {i}",
                    f"English {i}",
                    f"Japanese {i}",
                )
                for i in range(count)
            ],
        ]
    finally:
        wb.close()


def test_build_workbook_flat_structure():
    project = Project(
        name="Flat App",
        slug="flat-app",
        base_language="vi",
        target_languages=["en", "ja"],
        layout=ProjectLayout.flat,
    )
    tag = Tag(name="v1")
    t_en = Translation(locale="en", value="Save")
    entry = StringEntry(
        key="btn.save",
        source_text="Lưu",
        description="Save button label",
        status=TranslationStatus.draft,
        tags=[tag],
        translations=[t_en],
    )

    wb_bytes = build_workbook(project, [entry], stage="draft")
    wb = load_workbook(io.BytesIO(wb_bytes))

    assert FLAT_SHEET in wb.sheetnames
    assert "_meta" in wb.sheetnames

    ws = wb[FLAT_SHEET]
    headers = [cell for cell in next(ws.iter_rows(values_only=True))]
    assert headers == ["key", "description", "tags", "vi", "en", "ja"]

    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) == 2  # header + 1 row
    row = rows[1]
    assert row[0] == "btn.save"
    assert row[1] == "Save button label"
    assert row[2] == "v1"
    assert row[3] == "Lưu"
    assert row[4] == "Save"
    assert row[5] in ("", None)  # missing ja is empty


def test_build_workbook_modular_structure():
    project = Project(
        name="Modular App",
        slug="modular-app",
        base_language="vi",
        target_languages=["en"],
        layout=ProjectLayout.modular,
    )
    mod_auth = Module(slug="auth", name="Authentication")
    mod_home = Module(slug="home", name="Home Page")

    entry_auth = StringEntry(
        key="login",
        source_text="Đăng nhập",
        module=mod_auth,
        status=TranslationStatus.draft,
        translations=[],
    )
    entry_home = StringEntry(
        key="welcome",
        source_text="Chào mừng",
        module=mod_home,
        status=TranslationStatus.draft,
        translations=[],
    )
    entry_unassigned = StringEntry(
        key="common.ok",
        source_text="Đồng ý",
        module=None,
        status=TranslationStatus.draft,
        translations=[],
    )

    wb_bytes = build_workbook(project, [entry_auth, entry_home, entry_unassigned], stage="draft")
    wb = load_workbook(io.BytesIO(wb_bytes))

    assert "auth" in wb.sheetnames
    assert "home" in wb.sheetnames
    assert UNASSIGNED_SHEET in wb.sheetnames

    auth_rows = list(wb["auth"].iter_rows(values_only=True))
    assert auth_rows[1][0] == "login"

    home_rows = list(wb["home"].iter_rows(values_only=True))
    assert home_rows[1][0] == "welcome"

    unassigned_rows = list(wb[UNASSIGNED_SHEET].iter_rows(values_only=True))
    assert unassigned_rows[1][0] == "common.ok"


def test_build_workbook_public_stage_snapshot_isolation():
    project = Project(
        name="Public App",
        slug="public-app",
        base_language="vi",
        target_languages=["en"],
        layout=ProjectLayout.flat,
    )

    # 1. Draft string -> must be omitted
    draft_entry = StringEntry(
        key="draft_key",
        source_text="Bản nháp",
        status=TranslationStatus.draft,
        translations=[],
    )

    # 2. Public string with edits -> must use published snapshot
    t_live = Translation(
        locale="en",
        value="Working Copy Save",
        published_value="Published Save",
    )
    public_edited = StringEntry(
        key="save",
        source_text="Lưu (Working)",
        status=TranslationStatus.public,
        published_key="save",
        published_source_text="Lưu (Published)",
        translations=[t_live],
    )

    # 3. Pending delete string -> still public until published delete
    pending_delete = StringEntry(
        key="remove_me",
        source_text="Sắp xóa",
        status=TranslationStatus.public,
        published_key="remove_me",
        published_source_text="Vẫn còn trên prod",
        pending_delete=True,
        translations=[],
    )

    wb_bytes = build_workbook(project, [draft_entry, public_edited, pending_delete], stage="public")
    wb = load_workbook(io.BytesIO(wb_bytes))
    rows = list(wb[FLAT_SHEET].iter_rows(values_only=True))

    keys = [r[0] for r in rows[1:]]
    assert "draft_key" not in keys
    assert "save" in keys
    assert "remove_me" in keys

    save_row = next(r for r in rows[1:] if r[0] == "save")
    assert save_row[3] == "Lưu (Published)"  # vi base language
    assert save_row[4] == "Published Save"  # en translation


def test_build_template_workbook_flat_and_modular():
    flat_proj = Project(
        name="Flat Template",
        slug="flat-tpl",
        base_language="vi",
        target_languages=["en"],
        layout=ProjectLayout.flat,
    )
    flat_bytes = build_template_workbook(flat_proj)
    flat_wb = load_workbook(io.BytesIO(flat_bytes))
    assert FLAT_SHEET in flat_wb.sheetnames

    modular_proj = Project(
        name="Modular Template",
        slug="mod-tpl",
        base_language="vi",
        target_languages=["en"],
        layout=ProjectLayout.modular,
    )
    mod_bytes = build_template_workbook(modular_proj)
    mod_wb = load_workbook(io.BytesIO(mod_bytes))
    assert "auth" in mod_wb.sheetnames
    assert "common" in mod_wb.sheetnames
    assert UNASSIGNED_SHEET in mod_wb.sheetnames


def test_import_workbook_dry_run_via_api(client):
    project = _make_project(client, name="Import Dry Run App")
    pid = project["id"]

    wb = Workbook()
    ws = wb.active
    ws.title = FLAT_SHEET
    ws.append(["key", "description", "tags", "vi", "en"])
    ws.append(["new_key_1", "Test desc", "tag1", "Nguồn 1", "Source 1"])
    ws.append(["new_key_2", "", "", "Nguồn 2", "Source 2"])

    buf = io.BytesIO()
    wb.save(buf)
    content = buf.getvalue()

    # Dry-run = True
    r = client.post(
        f"/api/projects/{pid}/import?dry_run=true",
        files={
            "file": (
                "test.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["dry_run"] is True
    assert data["created"] == 0
    assert data["diff"]["create_count"] == 2

    # Ensure nothing was actually written to catalog
    catalog = client.get(f"/api/projects/{pid}/strings").json()
    assert catalog["total"] == 0

    # Now apply (dry_run = False)
    r_apply = client.post(
        f"/api/projects/{pid}/import?dry_run=false",
        files={
            "file": (
                "test.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert r_apply.status_code == 200, r_apply.text
    catalog_applied = client.get(f"/api/projects/{pid}/strings").json()
    assert catalog_applied["total"] == 2
    keys = {s["key"] for s in catalog_applied["items"]}
    assert keys == {"new_key_1", "new_key_2"}


def test_import_workbook_mismatched_project_slug_rejected(client):
    project = _make_project(client, name="Real Project")
    pid = project["id"]

    wb = Workbook()
    ws = wb.active
    ws.title = FLAT_SHEET
    ws.append(["key", "vi", "en"])
    ws.append(["k1", "v1", "e1"])

    meta = wb.create_sheet("_meta")
    meta.append(["key", "value"])
    meta.append(["project_id", "00000000-0000-0000-0000-000000000000"])
    meta.append(["project_slug", "completely-different-project"])

    buf = io.BytesIO()
    wb.save(buf)

    r = client.post(
        f"/api/projects/{pid}/import",
        files={
            "file": (
                "mismatch.xlsx",
                buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert r.status_code == 400
    assert "belongs to project" in r.text.lower()


def test_import_workbook_skips_sheet_without_key_column(client):
    project = _make_project(client, name="Skip Sheet App")
    pid = project["id"]

    wb = Workbook()
    # Sheet 1: Notes (no key column)
    ws_notes = wb.active
    ws_notes.title = "Notes"
    ws_notes.append(["Title", "Description", "Author"])
    ws_notes.append(["Read me", "Some translation notes", "John"])

    # Sheet 2: strings (has key column)
    ws_strings = wb.create_sheet(FLAT_SHEET)
    ws_strings.append(["key", "vi", "en"])
    ws_strings.append(["valid_key", "Giá trị", "Value"])

    buf = io.BytesIO()
    wb.save(buf)

    r = client.post(
        f"/api/projects/{pid}/import",
        files={
            "file": (
                "multi.xlsx",
                buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert r.status_code == 200, r.text
    catalog = client.get(f"/api/projects/{pid}/strings").json()
    assert catalog["total"] == 1
    assert catalog["items"][0]["key"] == "valid_key"


def test_workbook_tag_order_matches_catalog_sort():
    low = uuid.UUID("00000000-0000-0000-0000-000000000001")
    mid = uuid.UUID("00000000-0000-0000-0000-000000000002")
    high = uuid.UUID("00000000-0000-0000-0000-000000000003")
    project = Project(
        name="Tag Order",
        slug="tag-order",
        base_language="vi",
        target_languages=["en"],
        layout=ProjectLayout.flat,
    )

    def entry(tags: list[Tag]) -> StringEntry:
        return StringEntry(
            key="greet",
            source_text="Xin chào",
            description="",
            status=TranslationStatus.draft,
            tags=tags,
            translations=[],
        )

    forward = [
        Tag(id=high, name="beta"),
        Tag(id=low, name="alpha"),
        Tag(id=mid, name="beta"),
    ]
    reverse = list(reversed(forward))
    expected = "alpha,beta,beta"
    assert [tag.name for tag in sorted(forward, key=lambda tag: (tag.name, str(tag.id)))] == [
        "alpha",
        "beta",
        "beta",
    ]

    def tags_cell(tags: list[Tag]) -> str:
        workbook = load_workbook(io.BytesIO(build_workbook(project, [entry(tags)], stage="draft")))
        rows = list(workbook[FLAT_SHEET].iter_rows(values_only=True))
        return rows[1][2]

    assert tags_cell(forward) == expected
    assert tags_cell(reverse) == expected


def test_build_workbook_serializes_formulas_as_text():
    project = Project(
        name="Security App",
        slug="security-app",
        base_language="vi",
        target_languages=["en"],
        layout=ProjectLayout.flat,
    )
    dangerous_entry = StringEntry(
        key='=HYPERLINK("http://evil")',
        source_text="=SUM(A1:A10)",
        description="+cmd|' /C calc'",
        status=TranslationStatus.draft,
        tags=[Tag(name="@malicious"), Tag(name="safe"), Tag(name="-1day")],
        translations=[Translation(locale="en", value="@malicious_link")],
    )
    normal_entry = StringEntry(
        key="btn.save",
        source_text="Lưu",
        description="Save button",
        status=TranslationStatus.draft,
        tags=[Tag(name="v1")],
        translations=[Translation(locale="en", value="Save")],
    )
    empty_tags_entry = StringEntry(
        key="plain.key",
        source_text="Xin chào",
        description="",
        status=TranslationStatus.draft,
        tags=[],
        translations=[],
    )
    wb_bytes = build_workbook(
        project, [dangerous_entry, normal_entry, empty_tags_entry], stage="draft"
    )
    _assert_workbook_is_text(wb_bytes)
    wb = load_workbook(io.BytesIO(wb_bytes))
    rows = list(wb[FLAT_SHEET].iter_rows(values_only=True))
    by_key = {r[0]: r for r in rows[1:]}

    dangerous = by_key['=HYPERLINK("http://evil")']
    assert dangerous[0] == '=HYPERLINK("http://evil")'
    assert dangerous[1] == "+cmd|' /C calc'"
    assert dangerous[2] == "-1day,@malicious,safe"
    assert dangerous[3] == "=SUM(A1:A10)"
    assert dangerous[4] == "@malicious_link"

    normal = by_key["btn.save"]
    assert normal[0] == "btn.save"
    assert normal[1] == "Save button"
    assert normal[2] == "v1"
    assert normal[3] == "Lưu"
    assert normal[4] == "Save"

    empty_tags = by_key["plain.key"]
    assert empty_tags[0] == "plain.key"
    assert empty_tags[2] in ("", None)


def test_import_workbook_literal_key_and_tags_round_trip(client):
    project = _make_project(client, name="Formula Key Tags Round Trip")
    pid = project["id"]

    export_project = Project(
        name=project["name"],
        slug=project["slug"],
        base_language=project["base_language"],
        target_languages=project["target_languages"],
        layout=ProjectLayout.flat,
    )
    export_project.id = uuid.UUID(pid)

    dangerous_entry = StringEntry(
        key='=HYPERLINK("http://evil")',
        source_text="=SUM(A1:A10)",
        description="+cmd|' /C calc'",
        status=TranslationStatus.draft,
        tags=[Tag(name="@malicious"), Tag(name="safe"), Tag(name="-1day")],
        translations=[Translation(locale="en", value="@malicious_link")],
    )
    plus_key_entry = StringEntry(
        key="+plus.key",
        source_text="Nguồn",
        status=TranslationStatus.draft,
        tags=[Tag(name="v2"), Tag(name="=cmd")],
        translations=[],
    )
    normal_entry = StringEntry(
        key="btn.save",
        source_text="Lưu",
        description="Save button",
        status=TranslationStatus.draft,
        tags=[Tag(name="v1")],
        translations=[Translation(locale="en", value="Save")],
    )
    empty_tags_entry = StringEntry(
        key="plain.key",
        source_text="Xin chào",
        status=TranslationStatus.draft,
        tags=[],
        translations=[],
    )
    apostrophe_key_entry = StringEntry(
        key="'quoted.key",
        source_text="Giữ apostrophe",
        status=TranslationStatus.draft,
        tags=[Tag(name="'keep")],
        translations=[],
    )

    content = build_workbook(
        export_project,
        [
            dangerous_entry,
            plus_key_entry,
            normal_entry,
            empty_tags_entry,
            apostrophe_key_entry,
        ],
        stage="draft",
    )

    r = client.post(
        f"/api/projects/{pid}/import?dry_run=false",
        files={
            "file": (
                "formula.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert r.status_code == 200, r.text
    result = r.json()
    assert result["created"] == 5
    assert result["dry_run"] is False

    catalog = client.get(f"/api/projects/{pid}/strings").json()
    assert catalog["total"] == 5
    by_key = {item["key"]: item for item in catalog["items"]}

    dangerous = by_key['=HYPERLINK("http://evil")']
    assert {t["name"] for t in dangerous["tags"]} == {"@malicious", "safe", "-1day"}
    assert dangerous["source_text"] == "=SUM(A1:A10)"
    assert dangerous["description"] == "+cmd|' /C calc'"

    plus_key = by_key["+plus.key"]
    assert {t["name"] for t in plus_key["tags"]} == {"v2", "=cmd"}

    normal = by_key["btn.save"]
    assert {t["name"] for t in normal["tags"]} == {"v1"}
    assert normal["source_text"] == "Lưu"
    assert normal["description"] == "Save button"

    empty_tags = by_key["plain.key"]
    assert empty_tags["tags"] == []

    apostrophe = by_key["'quoted.key"]
    assert {t["name"] for t in apostrophe["tags"]} == {"'keep"}

    # Re-importing literal keys updates in place instead of creating duplicates.
    r_again = client.post(
        f"/api/projects/{pid}/import?dry_run=false",
        files={
            "file": (
                "formula.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert r_again.status_code == 200, r_again.text
    catalog_again = client.get(f"/api/projects/{pid}/strings").json()
    assert catalog_again["total"] == 5
    assert {item["key"] for item in catalog_again["items"]} == {
        '=HYPERLINK("http://evil")',
        "+plus.key",
        "btn.save",
        "plain.key",
        "'quoted.key",
    }


def _assert_workbook_is_text(raw):
    with ZipFile(io.BytesIO(raw)) as archive:
        for name in archive.namelist():
            if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                root = ElementTree.fromstring(archive.read(name))
                assert root.findall(".//{*}f") == []
    wb = load_workbook(io.BytesIO(raw))
    try:
        for ws in wb:
            for row in ws:
                for cell in row:
                    if cell.value is not None:
                        assert cell.data_type == "s", (ws.title, cell.coordinate, cell.value)
    finally:
        wb.close()


@pytest.mark.parametrize("template", [False, True])
def test_all_workbook_text_including_templates_and_metadata_is_literal(monkeypatch, template):
    from app import excel

    project = Project(
        id=uuid.uuid4(),
        name="Text cells",
        slug='=HYPERLINK("https://example.com")',
        base_language="vi",
        target_languages=["en"],
        layout=ProjectLayout.flat,
    )
    value = '=HYPERLINK("https://example.com")'
    monkeypatch.setattr(excel, "DEMO_IMPORT_ROWS", {"auth": [(value, value)]})
    if template:
        raw = build_template_workbook(project)
    else:
        raw = build_workbook(
            project,
            [StringEntry(key=value, source_text=value, status=TranslationStatus.draft)],
            stage="draft",
        )
    _assert_workbook_is_text(raw)
    wb = load_workbook(io.BytesIO(raw), data_only=True)
    try:
        assert wb[FLAT_SHEET].cell(2, 1).value == value
        assert wb[FLAT_SHEET].cell(2, 4).value == value
        assert dict(wb["_meta"].values)["project_slug"] == project.slug
    finally:
        wb.close()


def _cache_formula_result(raw: bytes, ref: str, formula: str, cached: str) -> bytes:
    needle = f'<c r="{ref}"><f>{formula}</f><v /></c>'.encode()
    replacement = f'<c r="{ref}" t="str"><f>{formula}</f><v>{cached}</v></c>'.encode()
    out = io.BytesIO()
    with ZipFile(io.BytesIO(raw)) as src, ZipFile(out, "w") as dst:
        for name in src.namelist():
            data = src.read(name)
            if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                assert needle in data
                data = data.replace(needle, replacement, 1)
            dst.writestr(name, data)
    return out.getvalue()


def test_import_stores_formula_cells_as_typed_text(client):
    project = _make_project(client, name="Formula Cell Text")
    root = f"/api/projects/{project['id']}"
    wb = Workbook()
    ws = wb.active
    ws.title = FLAT_SHEET
    ws.append(["key", "vi", "en"])
    ws.append(["plain", "Xin chào", "Hello"])
    ws["A3"] = "number"
    ws["B3"] = 42
    ws["C3"] = 7.5
    ws["A4"] = "dated"
    ws["B4"] = datetime(2024, 1, 2)
    ws["C4"] = datetime(2024, 3, 4, 5, 6, 7)
    ws["A5"] = "uncached"
    ws["B5"] = "=B2"
    ws["C5"] = "=SUM(C2)"
    ws["A6"] = "cached"
    ws["B6"] = "=B2"
    ws["C6"] = "=C2"
    raw = io.BytesIO()
    wb.save(raw)
    content = _cache_formula_result(raw.getvalue(), "B6", "B2", "Cached source")
    content = _cache_formula_result(content, "C6", "C2", "Cached translation")

    response = client.post(f"{root}/import", files={"file": ("formulas.xlsx", content)})
    assert response.status_code == 200, response.text
    by_key = {item["key"]: item for item in client.get(f"{root}/strings").json()["items"]}

    def en(key: str) -> str:
        return next(t["value"] for t in by_key[key]["translations"] if t["locale"] == "en")

    assert by_key["plain"]["source_text"] == "Xin chào"
    assert en("plain") == "Hello"
    assert by_key["number"]["source_text"] == "42"
    assert en("number") == "7.5"
    assert by_key["dated"]["source_text"] == "2024-01-02 00:00:00"
    assert en("dated") == "2024-03-04 05:06:07"
    assert by_key["uncached"]["source_text"] == "=B2"
    assert en("uncached") == "=SUM(C2)"
    assert by_key["cached"]["source_text"] == "=B2"
    assert en("cached") == "=C2"


@pytest.mark.parametrize("stage", ["draft", "public"])
@pytest.mark.parametrize("layout", ["flat", "modular"])
@pytest.mark.parametrize("as_job", [False, True])
def test_formula_prefixes_and_literal_apostrophes_roundtrip(
    client, session_factory, stage, layout, as_job
):
    project = _make_project(client, name="Literal Excel Text", layout=layout)
    pid = uuid.UUID(project["id"])
    root = f"/api/projects/{pid}"
    values = ["'" * quotes + prefix + "literal" for prefix in "=+-@" for quotes in range(3)]
    tag_names = ["!first", *values, "z-last"]
    with session_factory() as db:
        tags = {name: Tag(project_id=pid, name=name) for name in tag_names}
        module = Module(project_id=pid, slug="auth", name="Auth") if layout == "modular" else None
        entries = [
            StringEntry(
                project_id=pid,
                key=value,
                key_policy_allows_duplicates=False,
                source_text=value,
                description=value,
                module=module,
                status=TranslationStatus.draft,
                tags=[tags[value]],
                translations=[Translation(locale="en", value=value)],
            )
            for value in values
        ]
        # Each special tag also occurs after a safe first tag in a combined cell.
        entries.append(
            StringEntry(
                project_id=pid,
                key="all-tags",
                key_policy_allows_duplicates=False,
                source_text="All tags",
                module=module,
                status=TranslationStatus.draft,
                tags=list(tags.values()),
                translations=[Translation(locale="en", value="All tags")],
            )
        )
        db.add_all(entries)
        db.commit()
        ids = [str(entry.id) for entry in entries]
    if stage == "public":
        publish_strings(client, str(pid), ids)
    before = client.get(f"{root}/strings").json()
    exported = client.get(f"{root}/export", params={"format": "xlsx", "stage": stage})
    assert exported.status_code == 200, exported.text
    _assert_workbook_is_text(exported.content)
    for _ in range(2):
        response = client.post(
            f"{root}/import",
            files={"file": ("literal.xlsx", exported.content)},
            headers={"Prefer": "respond-async"} if as_job else {},
        )
        if as_job:
            assert response.status_code == 202, response.text
            job = poll(client, project, "xlsx", response.json()["job_id"])
            assert job["status"] == "completed", job
            result = job["result"]
        else:
            assert response.status_code == 200, response.text
            result = response.json()
        assert result["created"] == 0
        assert result["updated"] == 0
        assert client.get(f"{root}/strings").json() == before
