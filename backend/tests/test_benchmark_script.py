"""Smoke-check scale data generation before the PostgreSQL nightly run."""

from __future__ import annotations

import io
import json

from openpyxl import load_workbook
from sqlalchemy import create_engine, func, select

from app.database import Base
from app.models import StringEntry, Translation
from scripts import benchmark_scale


def test_benchmark_generates_ten_locale_catalog_and_import_files(tmp_path, monkeypatch):
    monkeypatch.setattr(benchmark_scale, "ROWS", 5)
    engine = create_engine(f"sqlite:///{tmp_path / 'benchmark.db'}")
    Base.metadata.create_all(engine)
    project_id, key = benchmark_scale._project(engine, "benchmark-test")
    assert project_id and key
    benchmark_scale.seed(engine, project_id)
    with engine.connect() as connection:
        assert connection.scalar(select(func.count()).select_from(StringEntry)) == 5
        assert connection.scalar(select(func.count()).select_from(Translation)) == 50
    payload = json.loads(benchmark_scale.json_file())
    assert len(payload) == 11
    assert all(len(strings) == 5 for strings in payload.values())
    workbook = load_workbook(io.BytesIO(benchmark_scale.excel_file()), read_only=True)
    assert sum(1 for _ in workbook["strings"].iter_rows(values_only=True)) == 6
    workbook.close()
    engine.dispose()
