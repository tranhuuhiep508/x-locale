"""Optional module-scoped keys preserve import, history, and export identities."""

import io
import time
import uuid
from concurrent.futures import ThreadPoolExecutor
from threading import Event

import pytest
from alembic.config import Config
from openpyxl import load_workbook
from sqlalchemy import select, text
from sqlalchemy.exc import IntegrityError

from alembic import command
from app.config import settings
from app.models import Activity, Project, StringEntry
from app.services import import_bulk, jobs
from tests.helpers import json_upload, make_project, publish_strings
from tests.postgres_support import (
    BACKEND_DIR,
    _client_for,
    alembic_check,
    migration_database_url,
    prepare_schema,
    set_sqlalchemy_url,
)
from tests.test_import_concurrency import BUSY, seed_import_project
from tests.test_project_members import _add_user, _as
from tests.test_unique_keys import _upload_xlsx, _xlsx
from tests.test_write_transactions import database_error


def enabled_project(client, **kwargs):
    project = make_project(client, **kwargs)
    root = f"/api/projects/{project['id']}"
    result = client.patch(root, json={"allow_duplicate_keys": True})
    assert result.status_code == 200, result.text
    return result.json(), root


def module(client, root, slug):
    result = client.post(f"{root}/modules", json={"slug": slug, "name": slug})
    assert result.status_code == 201, result.text
    return result.json()["id"]


def create(client, root, key="title", mid=None, value="Title"):
    return client.post(f"{root}/strings", json={"key": key, "module_id": mid, "source_text": value})


def rows(client, root):
    return client.get(f"{root}/strings").json()["items"]


def modular_payload(a="Auth", b="Home"):
    return {"modules": {"auth": {"vi": {"title": a}}, "home": {"vi": {"title": b}}}}


def test_defaults_create_read_and_disable_without_duplicates(client, session_factory):
    project = make_project(client)
    root = f"/api/projects/{project['id']}"
    assert project["allow_duplicate_keys"] is False
    created = create(client, root).json()
    assert client.patch(root, json={"allow_duplicate_keys": True}).status_code == 200
    assert client.get(root).json()["allow_duplicate_keys"] is True
    assert client.get("/api/projects").json()[0]["allow_duplicate_keys"] is True
    key = client.post(f"{root}/api-keys", json={"name": "test"}).json()["key"]
    assert (
        client.get("/api/bootstrap", headers={"X-API-Key": key}).json()["allow_duplicate_keys"]
        is True
    )
    with session_factory() as db:
        assert db.get(StringEntry, uuid.UUID(created["id"])).key_policy_allows_duplicates is True
    assert client.patch(root, json={"allow_duplicate_keys": False}).status_code == 200
    with session_factory() as db:
        assert db.get(StringEntry, uuid.UUID(created["id"])).key_policy_allows_duplicates is False
    response = client.post("/api/projects", json={"name": "Opt in", "allow_duplicate_keys": True})
    assert response.status_code == 201 and response.json()["allow_duplicate_keys"] is True


@pytest.mark.parametrize("allow_duplicates", [False, True])
@pytest.mark.parametrize("layout", ["flat", "modular"])
def test_sync_state_reports_current_key_policy(client, allow_duplicates, layout):
    project = make_project(client)
    root = f"/api/projects/{project['id']}"
    assert client.patch(root, json={"allow_duplicate_keys": allow_duplicates}).status_code == 200
    response = client.get(f"{root}/sync-state", params={"layout": layout})
    assert response.status_code == 200, response.text
    assert response.json()["allow_duplicate_keys"] is allow_duplicates


def test_only_admin_can_change_policy(client):
    project = make_project(client)
    root = f"/api/projects/{project['id']}"
    editor = _add_user("editor-policy@example.com")
    assert (
        client.post(
            f"{root}/members", json={"email": editor["email"], "role": "editor"}
        ).status_code
        == 201
    )
    response = client.patch(root, json={"allow_duplicate_keys": True}, cookies=_as(editor))
    assert response.status_code == 403
    assert client.get(root).json()["allow_duplicate_keys"] is False


def test_enabled_keys_are_unique_per_module_and_unassigned(client):
    _, root = enabled_project(client)
    auth, home = module(client, root, "auth"), module(client, root, "home")
    for mid in (auth, home, None):
        assert create(client, root, mid=mid).status_code == 201
        assert create(client, root, mid=mid).status_code == 409
    assert len(rows(client, root)) == 3
    response = client.patch(root, json={"allow_duplicate_keys": False, "name": "Must roll back"})
    assert response.status_code == 409
    assert response.json()["detail"]["conflicts"] == [
        {"key": "title", "modules": ["_unassigned", "auth", "home"]}
    ]
    project = client.get(root).json()
    assert project["allow_duplicate_keys"] is True and project["name"] != "Must roll back"
    for row in rows(client, root)[1:]:
        assert client.delete(f"{root}/strings/{row['id']}").status_code == 204
    assert client.patch(root, json={"allow_duplicate_keys": False}).status_code == 200


def test_move_rename_and_module_delete_validate_final_scope(client):
    _, root = enabled_project(client)
    auth, home = module(client, root, "auth"), module(client, root, "home")
    first = create(client, root, mid=auth).json()
    create(client, root, mid=home)
    assert (
        client.patch(f"{root}/strings/{first['id']}", json={"module_id": home}).status_code == 409
    )
    moved = client.patch(f"{root}/strings/{first['id']}", json={"module_id": home, "key": "other"})
    assert moved.status_code == 200, moved.text
    assert client.patch(f"{root}/strings/{first['id']}", json={"key": "title"}).status_code == 409
    assert client.delete(f"{root}/modules/{home}").status_code == 204
    assert all(row["module_id"] is None for row in rows(client, root))
    auth_title = create(client, root, mid=auth)
    assert auth_title.status_code == 201
    deleted = client.delete(f"{root}/modules/{auth}")
    assert deleted.status_code == 409
    assert client.get(f"{root}/strings/{auth_title.json()['id']}").json()["module_id"] == auth
    assert any(m["id"] == auth for m in client.get(f"{root}/modules").json())


@pytest.mark.parametrize("allow_duplicates", [False, True])
@pytest.mark.parametrize("pending_delete", [False, True])
def test_module_delete_rejects_published_unassigned_collision(
    client, allow_duplicates, pending_delete
):
    project = make_project(client)
    root = f"/api/projects/{project['id']}"
    assert client.patch(root, json={"allow_duplicate_keys": allow_duplicates}).status_code == 200
    auth, home = module(client, root, "auth"), module(client, root, "home")
    first = create(client, root, mid=auth, value="Auth title").json()
    publish_strings(client, project["id"], [first["id"]])
    # Only the published snapshot still belongs to the module being deleted.
    assert client.patch(
        f"{root}/strings/{first['id']}", json={"key": "renamed", "module_id": home}
    ).status_code == 200
    if pending_delete:
        assert client.delete(f"{root}/strings/{first['id']}").status_code == 204
    second = create(client, root, value="Unassigned title").json()
    publish_strings(client, project["id"], [second["id"]])
    before = [client.get(f"{root}/strings/{sid}").json() for sid in (first["id"], second["id"])]
    for fmt in ("json", "xlsx"):
        assert client.get(
            f"{root}/export", params={"stage": "public", "format": fmt}
        ).status_code == 200

    response = client.delete(f"{root}/modules/{auth}")
    assert response.status_code == 409, response.text
    assert "Published key 'title'" in response.json()["detail"]
    assert any(m["id"] == auth for m in client.get(f"{root}/modules").json())
    after = [client.get(f"{root}/strings/{sid}").json() for sid in (first["id"], second["id"])]
    assert after == before
    for fmt in ("json", "xlsx"):
        assert client.get(
            f"{root}/export", params={"stage": "public", "format": fmt}
        ).status_code == 200

    publish_strings(client, project["id"], [first["id"]])
    assert client.delete(f"{root}/modules/{auth}").status_code == 204
    assert client.get(f"{root}/export", params={"stage": "public"}).status_code == 200


@pytest.mark.parametrize("excluded_state", ["unpublished", "tombstone"])
def test_module_delete_ignores_nonexported_published_snapshots(client, excluded_state):
    project, root = enabled_project(client)
    auth = module(client, root, "auth")
    first = create(client, root, mid=auth).json()
    publish_strings(client, project["id"], [first["id"]])
    assert client.patch(f"{root}/strings/{first['id']}", json={"key": "renamed"}).status_code == 200
    second = create(client, root).json()
    publish_strings(client, project["id"], [second["id"]])
    if excluded_state == "unpublished":
        response = client.post(
            f"{root}/strings/batch", json={"action": "unpublish", "string_ids": [second["id"]]}
        )
        assert response.status_code == 200, response.text
    else:
        assert client.delete(f"{root}/strings/{second['id']}").status_code == 204
        publish_strings(client, project["id"], [second["id"]])

    response = client.delete(f"{root}/modules/{auth}")
    assert response.status_code == 204, response.text
    assert client.get(f"{root}/export", params={"stage": "public"}).status_code == 200


@pytest.mark.parametrize("kind", ["cli", "json", "excel"])
@pytest.mark.parametrize("dry_run", [False, True])
def test_modular_import_preserves_each_module(client, kind, dry_run):
    project, root = enabled_project(client)
    params = {"dry_run": dry_run}
    if kind == "cli":
        response = client.post(f"{root}/strings/import", json=modular_payload(), params=params)
    elif kind == "json":
        response = client.post(
            f"{root}/import", files=json_upload(modular_payload()), params=params
        )
    else:
        response = _upload_xlsx(
            client,
            project["id"],
            _xlsx({"auth": [("title", "Auth")], "home": [("title", "Home")]}),
            dry_run=dry_run,
        )
    assert response.status_code == 200, response.text
    assert response.json()["diff"]["create_count"] == 2
    if dry_run:
        assert rows(client, root) == []
        assert client.get(f"{root}/modules").json() == []
    else:
        assert {(row["module_slug"], row["source_text"]) for row in rows(client, root)} == {
            ("auth", "Auth"),
            ("home", "Home"),
        }
        response = client.post(f"{root}/strings/import", json=modular_payload("Updated", "Home"))
        assert response.status_code == 200 and response.json()["updated"] == 1
        assert {(row["module_slug"], row["source_text"]) for row in rows(client, root)} == {
            ("auth", "Updated"),
            ("home", "Home"),
        }


def test_flat_import_ambiguity_and_explicit_unassigned(client):
    _, root = enabled_project(client)
    auth = module(client, root, "auth")
    create(client, root, mid=auth, value="Auth")
    # A flat file with one candidate preserves its existing module.
    assert client.get(f"{root}/export", params={"layout": "flat"}).status_code == 200
    assert client.get(f"{root}/sync-state", params={"layout": "flat"}).status_code == 200
    response = client.post(f"{root}/import", files=json_upload({"title": "Unique update"}))
    assert response.status_code == 200 and rows(client, root)[0]["module_id"] == auth
    payload = {
        "modules": {"home": {"vi": {"title": "Home"}}},
        "unassigned": {"vi": {"title": "Unassigned"}},
    }
    assert client.post(f"{root}/import", files=json_upload(payload)).status_code == 200
    for dry_run in (True, False):
        response = client.post(
            f"{root}/import",
            files=json_upload({"title": "Ambiguous", "new": "Must not insert"}),
            params={"dry_run": dry_run},
        )
        assert response.status_code == 409 and response.json()["detail"]["code"] == "ambiguous_key"
    assert len(rows(client, root)) == 3
    # Explicit Unassigned is never resolved to an assigned module.
    response = client.post(
        f"{root}/import",
        files=json_upload(
            {
                "unassigned": {"vi": {"title": "New unassigned"}},
                "modules": {"auth": {"vi": {"title": "Auth"}}},
            }
        ),
    )
    assert response.status_code == 200, response.text
    assert (
        next(row for row in rows(client, root) if row["module_id"] is None)["source_text"]
        == "New unassigned"
    )


def test_modular_deleted_protection_sync_state_and_orphans(client):
    _, root = enabled_project(client)
    response = client.post(f"{root}/strings/import", json=modular_payload())
    assert response.status_code == 200
    auth = next(row for row in rows(client, root) if row["module_slug"] == "auth")
    assert client.delete(f"{root}/strings/{auth['id']}").status_code == 204
    state = client.get(f"{root}/sync-state", params={"layout": "modular"}).json()
    assert state["tombstones"] == ["auth/title"]
    # A flat import with one live candidate updates that module, not another module's tombstone.
    assert (
        client.post(
            f"{root}/strings/import", json={"strings": {"title": "Flat home update"}}
        ).status_code
        == 200
    )
    assert (
        client.post(
            f"{root}/strings/import", json={"modules": {"home": {"vi": {"title": "Updated home"}}}}
        ).status_code
        == 200
    )
    response = client.post(f"{root}/strings/import", json=modular_payload())
    assert response.status_code == 409 and response.json()["detail"]["tombstones"] == ["auth/title"]
    # UI imports can restore tombstones in exactly the selected module.
    assert client.post(f"{root}/import", files=json_upload(modular_payload())).status_code == 200
    response = client.post(
        f"{root}/import",
        files=json_upload({"modules": {"auth": {"vi": {}}, "home": {"vi": {"title": "Home"}}}}),
        params={"dry_run": True},
    )
    assert response.status_code == 200 and response.json()["diff"]["orphan_count"] == 1
    assert response.json()["diff"]["orphan"][0]["key"] == "auth/title"


@pytest.mark.parametrize("stage", ["draft", "public"])
def test_modular_roundtrip_and_flat_export_conflicts(client, stage):
    project, root = enabled_project(client)
    assert client.post(f"{root}/strings/import", json=modular_payload()).status_code == 200
    if stage == "public":
        publish_strings(client, project["id"], [row["id"] for row in rows(client, root)])
    response = client.get(f"{root}/export", params={"stage": stage, "layout": "modular"})
    assert response.status_code == 200
    data = response.json()
    assert data["modules"]["auth"]["vi"]["title"] == "Auth"
    assert data["modules"]["home"]["vi"]["title"] == "Home"
    workbook = client.get(f"{root}/export", params={"stage": stage, "format": "xlsx"})
    assert workbook.status_code == 200
    wb = load_workbook(io.BytesIO(workbook.content))
    assert wb["auth"].cell(2, 1).value == wb["home"].cell(2, 1).value == "title"
    assert (
        client.post(
            f"{root}/import", files={"file": ("roundtrip.xlsx", workbook.content)}
        ).status_code
        == 200
    )
    for suffix in ("export", "translations.json", "sync-state"):
        response = client.get(f"{root}/{suffix}", params={"stage": stage, "layout": "flat"})
        assert response.status_code == 409, response.text
        assert response.json()["detail"]["code"] == "ambiguous_key"
    # Layout remains independent; only ambiguous operations are blocked.
    assert client.patch(root, json={"layout": "flat"}).status_code == 200
    assert (
        client.get(f"{root}/export", params={"stage": stage, "format": "xlsx"}).status_code == 409
    )


def test_public_export_checks_snapshot_identity(client):
    project, root = enabled_project(client)
    auth = module(client, root, "auth")
    first = create(client, root, mid=auth).json()
    publish_strings(client, project["id"], [first["id"]])
    assert client.patch(f"{root}/strings/{first['id']}", json={"key": "renamed"}).status_code == 200
    second = create(client, root, mid=auth).json()
    preview = client.post(
        f"{root}/strings/publish-preview", json={"string_ids": [second["id"]]}
    ).json()
    rejected = client.post(
        f"{root}/strings/batch",
        json={
            "action": "publish",
            "string_ids": [second["id"]],
            "fingerprint": preview["fingerprint"],
        },
    )
    assert rejected.status_code == 409, rejected.text
    assert rejected.json()["detail"]["code"] == "published_key_conflict"
    assert client.get(f"{root}/export", params={"stage": "draft"}).status_code == 200
    public = client.get(f"{root}/export", params={"stage": "public"})
    assert public.status_code == 200, public.text
    assert public.json()["modules"]["auth"]["vi"]["title"] == "Title"
    assert (
        client.get(f"{root}/export", params={"stage": "public", "format": "xlsx"}).status_code
        == 200
    )


def test_restore_history_and_batch_undo_use_module_identity(client):
    _, root = enabled_project(client)
    response = client.post(f"{root}/strings/import", json=modular_payload())
    batch_id = response.json()["batch_id"]
    auth = next(row for row in rows(client, root) if row["module_slug"] == "auth")
    original = client.get(f"{root}/strings/{auth['id']}/activities").json()["items"][0]["id"]
    assert client.patch(f"{root}/strings/{auth['id']}", json={"key": "renamed"}).status_code == 200
    preview = client.get(f"{root}/strings/{auth['id']}/activities/{original}/restore/preview")
    assert preview.status_code == 200 and preview.json()["can_restore"]
    response = client.post(f"{root}/strings/{auth['id']}/activities/{original}/restore")
    assert response.status_code == 200, response.text
    assert client.delete(f"{root}/strings/{auth['id']}").status_code == 204
    restored = client.post(
        f"{root}/strings/batch", json={"action": "restore", "string_ids": [auth["id"]]}
    )
    assert restored.status_code == 200, restored.text
    preview = client.get(f"{root}/activities/batch/{batch_id}/revert/preview")
    assert preview.status_code == 200 and preview.json()["can_revert"], preview.text
    undone = client.post(f"{root}/activities/batch/{batch_id}/revert", params={"force": True})
    assert undone.status_code == 200, undone.text
    assert rows(client, root) == []


def test_async_modular_import(client, monkeypatch):
    _, root = enabled_project(client)
    monkeypatch.setattr(settings, "import_async_threshold", 1)
    response = client.post(
        f"{root}/strings/import", json=modular_payload(), headers={"Prefer": "respond-async"}
    )
    assert response.status_code == 202, response.text
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline:
        result = client.get(f"/api/jobs/{response.json()['job_id']}").json()
        if result["status"] in ("completed", "failed"):
            break
        time.sleep(0.02)
    assert result["status"] == "completed", result
    assert len(rows(client, root)) == 2


@pytest.mark.parametrize("mode", ["alembic", "create_all"])
def test_database_constraints_and_toggle_race(throwaway_database, monkeypatch, mode):
    prepare_schema(throwaway_database, mode)
    monkeypatch.setattr(settings, "database_url", throwaway_database)
    for client, factory in _client_for(throwaway_database):
        project, root = enabled_project(client)
        auth, home = module(client, root, "auth"), module(client, root, "home")
        first = create(client, root, mid=auth)
        assert first.status_code == 201
        pid = uuid.UUID(project["id"])
        # The internal flag cannot bypass the project's policy.
        with factory() as db:
            with pytest.raises(IntegrityError):
                db.add(StringEntry(project_id=pid, key="new", source_text="New"))
                db.commit()
        # NULL modules also participate in uniqueness.
        for mid in (auth, None):
            with factory() as db:
                if mid is None:
                    db.add(
                        StringEntry(
                            project_id=pid,
                            key="title",
                            source_text="U",
                            key_policy_allows_duplicates=True,
                        )
                    )
                    db.commit()
                with pytest.raises(IntegrityError):
                    db.add(
                        StringEntry(
                            project_id=pid,
                            module_id=uuid.UUID(mid) if mid else None,
                            key="title",
                            source_text="Duplicate",
                            key_policy_allows_duplicates=True,
                        )
                    )
                    db.commit()
        # A concurrent toggle waits for an inserting transaction, then detects its duplicates.
        with factory() as db:
            db.add(
                StringEntry(
                    project_id=pid,
                    module_id=uuid.UUID(home),
                    key="title",
                    source_text="Home",
                    key_policy_allows_duplicates=True,
                )
            )
            db.flush()
            started = Event()

            def disable(started=started, client=client, root=root):
                started.set()
                return client.patch(root, json={"allow_duplicate_keys": False})

            with ThreadPoolExecutor(max_workers=1) as pool:
                future = pool.submit(disable)
                assert started.wait(2)
                blocker = db.scalar(text("SELECT pg_backend_pid()"))
                blocked = False
                deadline = time.monotonic() + 5
                with factory() as observer:
                    while time.monotonic() < deadline:
                        blocked = observer.scalar(
                            text(
                                "SELECT EXISTS (SELECT 1 FROM pg_stat_activity "
                                "WHERE :blocker = ANY(pg_blocking_pids(pid)))"
                            ),
                            {"blocker": blocker},
                        )
                        if blocked:
                            break
                        observer.rollback()  # Refresh pg_stat_activity for newly connected waiters.
                        time.sleep(0.02)
                db.commit()
                assert future.result(timeout=10).status_code == 409
                assert blocked, "policy toggle must wait for the inserting transaction"
        if mode == "alembic":
            alembic_check(throwaway_database)
            cfg = Config(str(BACKEND_DIR / "alembic.ini"))
            set_sqlalchemy_url(cfg, throwaway_database)
            with migration_database_url(throwaway_database):
                with pytest.raises(RuntimeError, match="duplicate keys"):
                    command.downgrade(cfg, "r7d05b6c8901")
        with factory() as db:
            assert db.scalar(select(Project.allow_duplicate_keys).where(Project.id == pid)) is True
            db.execute(
                text("UPDATE strings SET deleted_at = now() WHERE project_id = :pid"), {"pid": pid}
            )
            db.commit()
        assert client.patch(root, json={"allow_duplicate_keys": False}).status_code == 200
        with factory() as db:
            with pytest.raises(IntegrityError):
                db.add(
                    StringEntry(
                        project_id=pid,
                        key="bypass",
                        source_text="No bypass",
                        key_policy_allows_duplicates=True,
                    )
                )
                db.commit()
            db.rollback()
            db.add(
                StringEntry(
                    project_id=pid, module_id=uuid.UUID(auth), key="strict", source_text="First"
                )
            )
            db.flush()
            with pytest.raises(IntegrityError):
                db.add(
                    StringEntry(
                        project_id=pid,
                        module_id=uuid.UUID(home),
                        key="strict",
                        source_text="Conflict",
                    )
                )
                db.commit()
        if mode == "alembic":
            with migration_database_url(throwaway_database):
                command.downgrade(cfg, "r7d05b6c8901")
                command.upgrade(cfg, "head")
            alembic_check(throwaway_database)


def test_pending_deletes_restore_and_bulk_moves_keep_scope(client):
    project, root = enabled_project(client)
    auth, home = module(client, root, "auth"), module(client, root, "home")
    first = create(client, root, mid=auth).json()
    publish_strings(client, project["id"], [first["id"]])
    assert client.delete(f"{root}/strings/{first['id']}").status_code == 204
    assert create(client, root, mid=auth).status_code == 409
    second = create(client, root, mid=home).json()
    move = client.post(
        f"{root}/strings/batch",
        json={
            "action": "move_module",
            "string_ids": [second["id"]],
            "payload": {"module_id": auth},
        },
    )
    assert move.status_code == 409
    assert client.get(f"{root}/strings/{second['id']}").json()["module_id"] == home
    state = client.get(f"{root}/sync-state", params={"layout": "modular"}).json()
    assert state["pending_remove"] == ["auth/title"]
    publish_strings(client, project["id"], [first["id"]])
    replacement = create(client, root, mid=auth)
    assert replacement.status_code == 201
    restore = client.post(
        f"{root}/strings/batch", json={"action": "restore", "string_ids": [first["id"]]}
    )
    assert restore.status_code == 409
    assert client.get(f"{root}/strings/{first['id']}").json()["deleted_at"] is not None


def test_history_restore_and_undo_preview_block_same_module(client):
    _, root = enabled_project(client)
    auth, home = module(client, root, "auth"), module(client, root, "home")
    first = create(client, root, mid=auth).json()
    version = client.get(f"{root}/strings/{first['id']}/activities").json()["items"][0]["id"]
    assert client.patch(f"{root}/strings/{first['id']}", json={"key": "renamed"}).status_code == 200
    rename = client.get(f"{root}/strings/{first['id']}/activities").json()["items"][0]["id"]
    for mid in (auth, home):
        assert create(client, root, mid=mid).status_code == 201
    preview = client.get(f"{root}/strings/{first['id']}/activities/{version}/restore/preview")
    assert not preview.json()["can_restore"]
    assert (
        client.post(f"{root}/strings/{first['id']}/activities/{version}/restore").status_code == 409
    )
    preview = client.get(f"{root}/activities/{rename}/revert/preview")
    assert preview.status_code == 200 and not preview.json()["can_revert"], preview.text
    assert (
        client.post(f"{root}/activities/{rename}/revert", params={"force": True}).status_code == 409
    )


def test_waiting_async_import_blocks_policy_change(client, monkeypatch):
    project = make_project(client)
    root = f"/api/projects/{project['id']}"
    entered, release = Event(), Event()
    original = jobs.wait_for_import_lock

    def pause(db, project_id):
        original(db, project_id)
        entered.set()
        assert release.wait(10)

    monkeypatch.setattr(jobs, "wait_for_import_lock", pause)
    monkeypatch.setattr(settings, "import_async_threshold", 1)
    response = client.post(
        f"{root}/strings/import",
        json={"strings": {"a": "A", "b": "B"}},
        headers={"Prefer": "respond-async"},
    )
    assert response.status_code == 202
    try:
        assert entered.wait(5)
        changed = client.patch(root, json={"allow_duplicate_keys": True})
        assert changed.status_code == 409 and changed.json() == BUSY
    finally:
        release.set()
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline:
        result = client.get(f"/api/jobs/{response.json()['job_id']}").json()
        if result["status"] in ("completed", "failed"):
            break
        time.sleep(0.02)
    assert result["status"] == "completed", result
    assert len(rows(client, root)) == 2
    assert client.patch(root, json={"allow_duplicate_keys": True}).status_code == 200


@pytest.mark.parametrize("initial_policy", [False, True])
@pytest.mark.parametrize("asynchronous", [False, True])
def test_database_policy_change_during_import_returns_retryable_conflict(
    client, session_factory, monkeypatch, initial_policy, asynchronous
):
    project = make_project(client)
    root = f"/api/projects/{project['id']}"
    assert client.patch(root, json={"allow_duplicate_keys": initial_policy}).status_code == 200
    existing = create(client, root, key="existing", value="Original").json()
    pid = uuid.UUID(project["id"])
    with session_factory() as db:
        activity_count = db.query(Activity).filter(Activity.project_id == pid).count()
    entered, release = Event(), Event()
    original = import_bulk.insert_rows
    attempts = []

    def pause_before_insert(db, table, new_rows, **kwargs):
        if table is StringEntry.__table__:
            attempts.append(True)
            entered.set()
            assert release.wait(10)
        return original(db, table, new_rows, **kwargs)

    monkeypatch.setattr(import_bulk, "insert_rows", pause_before_insert)
    monkeypatch.setattr(settings, "import_async_threshold", 1)
    with ThreadPoolExecutor(max_workers=1) as pool:
        future = pool.submit(
            client.post,
            f"{root}/strings/import",
            json={"strings": {"existing": "Updated", "a": "A", "b": "B"}},
            headers={"Prefer": "respond-async"} if asynchronous else {},
        )
        try:
            if asynchronous:
                assert future.result(timeout=5).status_code == 202
            assert entered.wait(5)
            # Bypass the API guard to exercise the database safety net for external writes.
            with session_factory() as db:
                db.get(Project, pid).allow_duplicate_keys = not initial_policy
                db.commit()
        finally:
            release.set()
        response = future.result(timeout=5)

    if asynchronous:
        deadline = time.monotonic() + 10
        while time.monotonic() < deadline:
            result = client.get(f"/api/jobs/{response.json()['job_id']}").json()
            if result["status"] in ("completed", "failed"):
                break
            time.sleep(0.02)
        assert result["status"] == "failed", result
        error = result["result"]["error"]
        assert result["error"] == "The project key setting changed. Please retry."
    else:
        error = {"status_code": response.status_code, "detail": response.json()["detail"]}
    assert error == {
        "status_code": 409,
        "detail": {
            "code": "concurrent_modification",
            "message": "The project key setting changed. Please retry.",
            "retryable": True,
        },
    }
    assert attempts == [True]
    assert [(row["id"], row["source_text"]) for row in rows(client, root)] == [
        (existing["id"], "Original")
    ]
    with session_factory() as db:
        assert db.query(Activity).filter(Activity.project_id == pid).count() == activity_count


@pytest.mark.parametrize("initial_policy", [False, True])
@pytest.mark.parametrize("asynchronous", [False, True])
def test_policy_toggle_does_not_deadlock_with_existing_string_import(
    client, session_factory, monkeypatch, initial_policy, asynchronous
):
    project = seed_import_project(client, session_factory)
    root = f"/api/projects/{project['id']}"
    assert client.patch(root, json={"allow_duplicate_keys": initial_policy}).status_code == 200
    entered, release = Event(), Event()
    original = import_bulk.insert_rows

    def pause_before_audit(db, table, new_rows, **kwargs):
        if table is Activity.__table__ and new_rows:
            entered.set()
            assert release.wait(10)
        return original(db, table, new_rows, **kwargs)

    monkeypatch.setattr(import_bulk, "insert_rows", pause_before_audit)
    monkeypatch.setattr(settings, "import_async_threshold", 0)
    with ThreadPoolExecutor(max_workers=2) as pool:
        importing = pool.submit(
            client.post,
            f"{root}/strings/import",
            json={"strings": {"changed": "Updated"}},
            headers={"Prefer": "respond-async"} if asynchronous else {},
        )
        try:
            if asynchronous:
                assert importing.result(timeout=5).status_code == 202
            assert entered.wait(5)
            changing = pool.submit(
                client.patch,
                root,
                json={"allow_duplicate_keys": not initial_policy, "name": "Must roll back"},
            )
            rejected = changing.result(timeout=5)
            assert rejected.status_code == 409 and rejected.json() == BUSY
            unchanged = client.get(root).json()
            assert unchanged["allow_duplicate_keys"] is initial_policy
            assert unchanged["name"] == project["name"]
            saved = client.patch(
                root,
                json={
                    "name": "Saved during import",
                    "base_language": project["base_language"],
                    "target_languages": project["target_languages"],
                    "layout": project["layout"],
                    "allow_duplicate_keys": initial_policy,
                    "translation_context": "Saved during import",
                },
            )
            assert saved.status_code == 200, saved.text
            assert saved.json()["name"] == "Saved during import"
            assert saved.json()["translation_context"] == "Saved during import"
            assert saved.json()["allow_duplicate_keys"] is initial_policy
        finally:
            release.set()
        response = importing.result(timeout=5)

    if asynchronous:
        deadline = time.monotonic() + 10
        while time.monotonic() < deadline:
            result = client.get(f"/api/jobs/{response.json()['job_id']}").json()
            if result["status"] in ("completed", "failed"):
                break
            time.sleep(0.02)
        assert result["status"] == "completed", result
    else:
        assert response.status_code == 200, response.text
    assert (
        next(row for row in rows(client, root) if row["key"] == "changed")["source_text"]
        == "Updated"
    )
    assert client.patch(root, json={"allow_duplicate_keys": not initial_policy}).status_code == 200


def test_policy_toggle_rejects_while_batch_undo_is_active(client, monkeypatch):
    project = make_project(client)
    root = f"/api/projects/{project['id']}"
    imported = client.post(f"{root}/strings/import", json={"strings": {"a": "A", "b": "B"}})
    assert imported.status_code == 200, imported.text
    from app.services import activities

    entered, release = Event(), Event()
    original = activities._persist_revert_plan

    def pause_before_write(*args, **kwargs):
        entered.set()
        assert release.wait(10)
        return original(*args, **kwargs)

    monkeypatch.setattr(activities, "_persist_revert_plan", pause_before_write)
    with ThreadPoolExecutor(max_workers=1) as pool:
        undoing = pool.submit(
            client.post, f"{root}/activities/batch/{imported.json()['batch_id']}/revert"
        )
        try:
            assert entered.wait(5)
            rejected = client.patch(root, json={"allow_duplicate_keys": True})
            assert rejected.status_code == 409 and rejected.json() == BUSY
            saved = client.patch(
                root, json={"allow_duplicate_keys": False, "name": "Saved during undo"}
            )
            assert saved.status_code == 200, saved.text
            assert saved.json()["name"] == "Saved during undo"
        finally:
            release.set()
        response = undoing.result(timeout=5)
    assert response.status_code == 200, response.text
    assert response.json()["reverted"] == 2
    assert client.patch(root, json={"allow_duplicate_keys": True}).status_code == 200


def test_async_import_keeps_unrelated_database_errors_generic(client, monkeypatch):
    project = make_project(client)
    root = f"/api/projects/{project['id']}"

    def fail_insert(db, table, new_rows, **kwargs):
        raise database_error("23503", "other_fk", integrity=True)

    monkeypatch.setattr(import_bulk, "insert_rows", fail_insert)
    monkeypatch.setattr(settings, "import_async_threshold", 1)
    response = client.post(
        f"{root}/strings/import",
        json={"strings": {"a": "A", "b": "B"}},
        headers={"Prefer": "respond-async"},
    )
    assert response.status_code == 202, response.text
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline:
        result = client.get(f"/api/jobs/{response.json()['job_id']}").json()
        if result["status"] in ("completed", "failed"):
            break
        time.sleep(0.02)
    assert result["status"] == "failed", result
    assert result["result"]["error"] == {"status_code": 500, "detail": "Job failed unexpectedly"}
    assert rows(client, root) == []


@pytest.mark.parametrize("dry_run", [False, True])
@pytest.mark.parametrize(
    "payload,expected_orphans",
    [
        ({"modules": {}, "unassigned": {"vi": {}}}, ["welcome"]),
        ({"modules": {"auth": {"vi": {}}}}, ["auth/login"]),
        (
            {"modules": {"auth": {"vi": {}}, "home": {"vi": {}}}, "unassigned": {"vi": {}}},
            ["auth/login", "home/title", "welcome"],
        ),
        ({}, ["login", "title", "welcome"]),
        ({"modules": {}, "unassigned": {"vi": {"welcome": "Welcome"}}}, []),
        (
            {"modules": {"home": {"vi": {"title": "Home"}}}, "unassigned": {"vi": {}}},
            ["welcome"],
        ),
    ],
)
def test_import_orphans_use_explicit_group_scope(client, dry_run, payload, expected_orphans):
    _, root = enabled_project(client)
    auth, home = module(client, root, "auth"), module(client, root, "home")
    for key, mid, value in (
        ("login", auth, "Auth"),
        ("title", home, "Home"),
        ("welcome", None, "Welcome"),
    ):
        assert create(client, root, key=key, mid=mid, value=value).status_code == 201
    original = rows(client, root)
    response = client.post(
        f"{root}/import", files=json_upload(payload), params={"dry_run": dry_run}
    )
    assert response.status_code == 200, response.text
    diff = response.json()["diff"]
    assert diff["orphan_count"] == len(expected_orphans)
    assert [item["key"] for item in diff["orphan"]] == expected_orphans
    assert diff["create_count"] == diff["update_count"] == 0
    assert rows(client, root) == original
