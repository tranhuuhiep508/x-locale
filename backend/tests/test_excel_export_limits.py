"""Excel exports must reject content that openpyxl would silently truncate."""

import io
import uuid
from datetime import UTC, datetime

import pytest
from openpyxl import load_workbook

from app import excel
from app.models import (
    Module,
    Project,
    ProjectLayout,
    StringEntry,
    Tag,
    Translation,
    TranslationStatus,
)
from tests.helpers import json_upload, make_project, publish_strings


@pytest.mark.parametrize("stage", ["draft", "public"])
@pytest.mark.parametrize("layout", ["flat", "modular"])
def test_excel_cell_limit_roundtrip(client, stage, layout):
    project = make_project(client, layout=layout)
    root = f"/api/projects/{project['id']}"
    text = "x" * 32767
    payload = {"vi": {"long": text}, "en": {"long": text}}
    if layout == "modular":
        payload = {"modules": {"auth": payload}}
    response = client.post(f"{root}/import", files=json_upload(payload))
    assert response.status_code == 200, response.text
    entry = client.get(f"{root}/strings").json()["items"][0]
    response = client.patch(f"{root}/strings/{entry['id']}", json={"description": text})
    assert response.status_code == 200, response.text
    if stage == "public":
        publish_strings(client, project["id"], [entry["id"]])
    before = client.get(f"{root}/strings").json()

    exported = client.get(f"{root}/export", params={"format": "xlsx", "stage": stage})
    assert exported.status_code == 200, exported.text
    wb = load_workbook(io.BytesIO(exported.content))
    try:
        row = list(wb["auth" if layout == "modular" else "strings"].values)[1]
        assert row[1] == row[3] == row[4] == text
    finally:
        wb.close()
    imported = client.post(f"{root}/import", files={"file": ("unchanged.xlsx", exported.content)})
    assert imported.status_code == 200, imported.text
    assert imported.json()["updated"] == 0
    assert client.get(f"{root}/strings").json() == before


@pytest.mark.parametrize("stage", ["draft", "public"])
@pytest.mark.parametrize("layout", ["flat", "modular"])
@pytest.mark.parametrize(
    ("field", "length"),
    [
        ("vi", 32768),
        ("vi", 40000),
        ("en", 32768),
        ("en", 40000),
        ("description", 32768),
        ("description", 40000),
        ("tags", 33152),
    ],
)
def test_excel_export_rejects_oversized_cells_without_changing_content(
    client, session_factory, stage, layout, field, length
):
    project = make_project(client, layout=layout)
    pid = uuid.UUID(project["id"])
    root = f"/api/projects/{pid}"
    text = "x" * length
    source = text if field == "vi" else "Source"
    target = text if field == "en" else "Target"
    with session_factory() as db:
        module = Module(project_id=pid, slug="auth", name="Auth") if layout == "modular" else None
        # Direct inserts also cover legacy text longer than today's write limit.
        entry = StringEntry(
            project_id=pid,
            key="long",
            key_policy_allows_duplicates=False,
            source_text=source,
            description=text if field == "description" else "Description",
            module=module,
            published_module=module,
            status=TranslationStatus.public,
            published_key="long",
            published_source_text=source,
            translations=[Translation(locale="en", value=target, published_value=target)],
            tags=[Tag(project_id=pid, name=f"{i:03}" + "t" * 125) for i in range(257)]
            if field == "tags"
            else [],
        )
        if field == "tags":
            assert len(",".join(tag.name for tag in entry.tags)) == length
        db.add(entry)
        db.commit()
    before = client.get(f"{root}/strings").json()

    response = client.get(f"{root}/export", params={"format": "xlsx", "stage": stage})
    assert response.status_code == 400, response.headers
    assert "Content-Disposition" not in response.headers
    detail = response.json()["detail"]
    for expected in (
        "auth" if layout == "modular" else "strings",
        "long",
        field,
        str(length),
        "32767",
    ):
        assert expected in detail
    assert client.get(f"{root}/strings").json() == before
    exported = client.get(f"{root}/export", params={"format": "json", "stage": stage})
    assert exported.status_code == 200, exported.text
    values = exported.json()
    if layout == "modular":
        values = values["modules"]["auth"]
    assert values["vi"]["long"] == source
    assert values["en"]["long"] == target


def test_excel_length_check_uses_selected_locale_and_published_snapshot():
    project = Project(
        slug="snapshot",
        base_language="vi",
        target_languages=["en", "ja"],
        layout=ProjectLayout.flat,
    )
    entry = StringEntry(
        key="working",
        source_text="x" * 40000,
        status=TranslationStatus.public,
        published_key="published",
        published_source_text="Published source",
        translations=[
            Translation(locale="en", value="x" * 40000, published_value="Published target")
        ],
    )
    raw = excel.build_workbook(project, [entry], "public")
    wb = load_workbook(io.BytesIO(raw))
    try:
        row = list(wb["strings"].values)[1]
        assert row[0] == "published"
        assert row[3:5] == ("Published source", "Published target")
    finally:
        wb.close()
    assert excel.build_workbook(project, [entry], "draft", locale="ja")
    entry.source_text = "Short working source"
    entry.published_source_text = "x" * 40000
    with pytest.raises(ValueError, match="published.*vi.*40000"):
        excel.build_workbook(project, [entry], "public")


@pytest.mark.parametrize(
    ("stage", "state"),
    [("draft", "deleted"), ("public", "deleted"), ("draft", "pending_delete"), ("public", "draft")],
)
def test_excel_length_check_ignores_omitted_rows(stage, state):
    project = Project(
        slug="omitted", base_language="vi", target_languages=[], layout=ProjectLayout.flat
    )
    entry = StringEntry(
        key="omitted",
        source_text="x" * 40000,
        published_source_text="x" * 40000,
        status=TranslationStatus.draft if state == "draft" else TranslationStatus.public,
        pending_delete=state == "pending_delete",
        deleted_at=datetime.now(UTC) if state == "deleted" else None,
    )
    raw = excel.build_workbook(project, [entry], stage)
    wb = load_workbook(io.BytesIO(raw))
    try:
        assert list(wb["strings"].values) == [("key", "description", "tags", "vi")]
    finally:
        wb.close()


@pytest.mark.parametrize("template", [False, True])
def test_excel_length_check_covers_hidden_metadata(template):
    project = Project(
        slug="x" * 32768, base_language="vi", target_languages=[], layout=ProjectLayout.flat
    )
    with pytest.raises(ValueError, match="_meta.*project_slug.*value.*32768"):
        if template:
            excel.build_template_workbook(project)
        else:
            excel.build_workbook(project, [], "draft")


@pytest.mark.parametrize("field", ["key", "vi"])
def test_excel_template_rejects_oversized_cells(client, monkeypatch, field):
    project = make_project(client, layout="flat")
    text = "x" * 32768
    monkeypatch.setattr(
        excel,
        "DEMO_IMPORT_ROWS",
        {"auth": [(text, "Source") if field == "key" else ("long", text)]},
    )
    response = client.get(
        f"/api/projects/{project['id']}/import-template", params={"format": "xlsx"}
    )
    assert response.status_code == 400, response.headers
    assert "Content-Disposition" not in response.headers
    assert field in response.json()["detail"]
    assert "32768" in response.json()["detail"]
