"""Reject user input that would overflow a varchar, before Postgres returns 500."""

from __future__ import annotations

import io
import uuid

import pytest
from alembic.config import Config
from fastapi import HTTPException
from openpyxl import Workbook
from sqlalchemy import create_engine, text

from alembic import command
from app.config import settings
from app.excel import FLAT_SHEET
from app.helpers import validate_locale_code
from app.limits import IMPORT_MAX_BYTES, TEXT_MAX_LENGTH
from app.main import app
from app.models import Project, StringEntry, Translation
from app.services.activities import _apply_published_snapshot, _apply_working_copy_only
from tests.helpers import json_upload, make_project

LONG_KEY = "k" * 513
MAX_KEY = "k" * 512
KEY_PREVIEW = "k" * 40
LONG_SLUG = "a" * 129
LONG_LOCALE = "a" * 17
REGEX_MAX_LOCALE = "abc-12345678"


def _set_target_languages(project_id: str, locales: list[str]) -> None:
    from app.database import get_db

    generator = app.dependency_overrides[get_db]()
    db = next(generator)
    try:
        row = db.get(Project, uuid.UUID(project_id))
        assert row is not None
        row.target_languages = locales
        db.commit()
    finally:
        generator.close()


def test_validate_locale_code_accepts_regex_maximum():
    assert validate_locale_code(REGEX_MAX_LOCALE) == REGEX_MAX_LOCALE
    assert len(REGEX_MAX_LOCALE) == 12


def test_project_locale_over_column_limit_is_400(client):
    response = client.post(
        "/api/projects",
        json={
            "name": "Too long locale",
            "base_language": LONG_LOCALE,
            "target_languages": ["en"],
            "layout": "flat",
        },
    )
    assert response.status_code == 400, response.text
    assert "locale" in response.text
    assert "16" in response.text


def test_project_accepts_locale_longer_than_the_old_varchar_10(client):
    response = client.post(
        "/api/projects",
        json={
            "name": "Wide locale",
            "base_language": REGEX_MAX_LOCALE,
            "target_languages": ["en-abcdefgh"],
            "layout": "flat",
        },
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["base_language"] == REGEX_MAX_LOCALE
    assert "en-abcdefgh" in body["target_languages"]


def test_string_update_rejects_key_over_512(client):
    project = make_project(client, layout="flat")
    created = client.post(
        f"/api/projects/{project['id']}/strings",
        json={"key": "hello", "source_text": "Xin chào"},
    )
    assert created.status_code == 201, created.text
    updated = client.patch(
        f"/api/projects/{project['id']}/strings/{created.json()['id']}",
        json={"key": LONG_KEY},
    )
    assert updated.status_code == 422, updated.text
    assert "key" in updated.text
    assert "512" in updated.text


def test_import_and_push_reject_key_over_512(client):
    project = make_project(client, layout="flat")
    pid = project["id"]
    uploaded = client.post(
        f"/api/projects/{pid}/import",
        files=json_upload({LONG_KEY: "Xin chào"}),
    )
    preview = f"{KEY_PREVIEW!r}… (513 chars)"
    assert uploaded.status_code == 400, uploaded.text
    assert preview in uploaded.text
    assert "512" in uploaded.text

    pushed = client.post(
        f"/api/projects/{pid}/strings/import",
        json={"strings": {LONG_KEY: "Xin chào"}, "base_language": "vi"},
    )
    assert pushed.status_code == 400, pushed.text
    assert preview in pushed.text
    assert "512" in pushed.text
    assert client.get(f"/api/projects/{pid}/strings").json()["total"] == 0


def test_import_stores_512_character_key_and_activity_summary(client):
    project = make_project(client, layout="flat")
    pid = project["id"]
    uploaded = client.post(
        f"/api/projects/{pid}/import",
        files=json_upload({MAX_KEY: "Xin chào"}),
    )
    assert uploaded.status_code == 200, uploaded.text
    stored = client.get(f"/api/projects/{pid}/strings").json()["items"]
    assert stored[0]["key"] == MAX_KEY
    activities = client.get(f"/api/projects/{pid}/activities").json()["items"]
    summaries = [item["summary"] for item in activities]
    assert any(MAX_KEY in summary for summary in summaries)


def test_module_slug_update_rejects_over_128(client):
    project = make_project(client)
    created = client.post(
        f"/api/projects/{project['id']}/modules",
        json={"slug": "auth", "name": "Auth"},
    )
    assert created.status_code == 201, created.text
    updated = client.patch(
        f"/api/projects/{project['id']}/modules/{created.json()['id']}",
        json={"slug": LONG_SLUG},
    )
    assert updated.status_code == 422, updated.text
    assert "slug" in updated.text
    assert "128" in updated.text


def test_module_name_update_rejects_over_255(client):
    project = make_project(client)
    created = client.post(
        f"/api/projects/{project['id']}/modules",
        json={"slug": "auth", "name": "Auth"},
    )
    updated = client.patch(
        f"/api/projects/{project['id']}/modules/{created.json()['id']}",
        json={"name": "n" * 256},
    )
    assert updated.status_code == 422, updated.text
    assert "name" in updated.text
    assert "255" in updated.text


def test_push_rejects_module_slug_over_128(client):
    project = make_project(client)
    pushed = client.post(
        f"/api/projects/{project['id']}/strings/import",
        json={
            "base_language": "vi",
            "modules": {LONG_SLUG: {"vi": {"hello": "Xin chào"}}},
        },
    )
    assert pushed.status_code == 400, pushed.text
    assert f"module slug {'a' * 40!r}… (129 chars)" in pushed.text
    assert "128" in pushed.text


def test_tag_update_rejects_name_and_color_over_limit(client):
    project = make_project(client, layout="flat")
    created = client.post(
        f"/api/projects/{project['id']}/tags",
        json={"name": "ui", "color": "#112233"},
    )
    assert created.status_code == 201, created.text
    tag_id = created.json()["id"]
    root = f"/api/projects/{project['id']}/tags/{tag_id}"

    renamed = client.patch(root, json={"name": "t" * 129})
    assert renamed.status_code == 422, renamed.text
    assert "name" in renamed.text
    assert "128" in renamed.text

    recolored = client.patch(root, json={"color": "c" * 33})
    assert recolored.status_code == 422, recolored.text
    assert "color" in recolored.text
    assert "32" in recolored.text


def _xlsx(rows: list[list[str]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = FLAT_SHEET
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_excel_import_rejects_long_key_and_tag_name(client):
    project = make_project(client, layout="flat")
    pid = project["id"]
    content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    long_key = client.post(
        f"/api/projects/{pid}/import",
        files={"file": ("keys.xlsx", _xlsx([["key", "vi"], [LONG_KEY, "Xin chào"]]), content_type)},
    )
    assert long_key.status_code == 400, long_key.text
    assert f"Sheet {FLAT_SHEET!r} row 2" in long_key.text
    assert f"key {KEY_PREVIEW!r}… (513 chars)" in long_key.text
    assert "512" in long_key.text

    long_tag = client.post(
        f"/api/projects/{pid}/import",
        files={
            "file": (
                "tags.xlsx",
                _xlsx([["key", "tags", "vi"], ["hello", "t" * 129, "Xin chào"]]),
                content_type,
            )
        },
    )
    assert long_tag.status_code == 400, long_tag.text
    assert f"tag name {'t' * 40!r}… (129 chars)" in long_tag.text
    assert "128" in long_tag.text


def test_snapshot_apply_rejects_keys_over_512():
    entry = type("Entry", (), {"key": "ok", "published_key": None})()
    with pytest.raises(HTTPException) as published:
        _apply_published_snapshot(None, entry, {"published_key": LONG_KEY})
    assert published.value.status_code == 400
    assert "published_key" in published.value.detail
    assert "512" in published.value.detail

    working = type("Entry", (), {})()
    db = type("DB", (), {"info": {}})()
    with pytest.raises(HTTPException) as current:
        _apply_working_copy_only(
            db,
            working,
            {
                "key": LONG_KEY,
                "source_text": "x",
                "description": None,
                "module_id": None,
                "tag_ids": [],
                "translations": {},
            },
        )
    assert current.value.status_code == 400
    assert "key" in current.value.detail
    assert "512" in current.value.detail


def test_legacy_underscore_locale_still_creates_imports_and_pushes(client):
    rejected = client.post(
        "/api/projects",
        json={
            "name": "Legacy rejected at the door",
            "base_language": "vi",
            "target_languages": ["en_US"],
            "layout": "flat",
        },
    )
    assert rejected.status_code == 400, rejected.text

    project = make_project(client, layout="flat")
    pid = project["id"]
    _set_target_languages(pid, ["en_US"])

    created = client.post(
        f"/api/projects/{pid}/strings",
        json={"key": "hello", "source_text": "Xin chào"},
    )
    assert created.status_code == 201, created.text
    assert "en_US" in {item["locale"] for item in created.json()["translations"]}

    uploaded = client.post(
        f"/api/projects/{pid}/import",
        files=json_upload({"imported": "Nhập"}),
    )
    assert uploaded.status_code == 200, uploaded.text

    pushed = client.post(
        f"/api/projects/{pid}/strings/import",
        json={"strings": {"pushed": "Đẩy"}, "base_language": "vi"},
    )
    assert pushed.status_code == 200, pushed.text

    stored = {
        item["key"]: item for item in client.get(f"/api/projects/{pid}/strings").json()["items"]
    }
    for key in ("hello", "imported", "pushed"):
        locales = {item["locale"] for item in stored[key]["translations"]}
        assert "en_US" in locales

    _set_target_languages(pid, ["a" * 17])
    overflow = client.post(
        f"/api/projects/{pid}/strings",
        json={"key": "overflow", "source_text": "Không"},
    )
    assert overflow.status_code == 400, overflow.text
    assert "locale" in overflow.text
    assert "16" in overflow.text


def test_widen_downgrade_names_rows_that_do_not_fit(throwaway_database, monkeypatch):
    url = throwaway_database
    monkeypatch.setattr(settings, "database_url", url)
    cfg = Config("alembic.ini")
    command.upgrade(cfg, "head")
    engine = create_engine(url)
    project_id = uuid.uuid4().hex
    with engine.begin() as connection:
        connection.execute(
            text(
                "INSERT INTO projects (id, name, slug, base_language, target_languages, layout) "
                "VALUES (:id, 'Wide', 'wide', :base, '[\"en\"]', 'flat')"
            ),
            {"id": project_id, "base": "abcdefghijk"},
        )
        connection.execute(
            text(
                "INSERT INTO activities ("
                "id, project_id, actor_type, actor_label, action, entity_type, entity_id, "
                "event_type, summary, is_revertible"
                ") VALUES ("
                ":id, :project_id, 'system', 'system', 'update', 'string', 'entity', "
                "'string.updated', :summary, false"
                ")"
            ),
            {
                "id": uuid.uuid4().hex,
                "project_id": project_id,
                "summary": "s" * 600,
            },
        )
    with pytest.raises(RuntimeError, match="Cannot downgrade p5b83e4f5678") as caught:
        command.downgrade(cfg, "o4f61c2d3456")
    message = str(caught.value)
    assert "1 row(s) in activities.summary exceed 512 characters" in message
    assert "1 row(s) in projects.base_language exceed 10 characters" in message


LONG_TEXT = "x" * (TEXT_MAX_LENGTH + 1)
MAX_TEXT = "y" * TEXT_MAX_LENGTH


def test_string_text_over_limit_is_422(client):
    project = make_project(client, layout="flat")
    root = f"/api/projects/{project['id']}/strings"
    created = client.post(root, json={"key": "hello", "source_text": "Hi"})
    assert created.status_code == 201, created.text
    sid = created.json()["id"]
    cases = [
        ("post", root, {"key": "too-source", "source_text": LONG_TEXT}),
        ("post", root, {"key": "too-desc", "source_text": "Hi", "description": LONG_TEXT}),
        ("post", root, {"key": "too-tr", "source_text": "Hi", "translations": {"en": LONG_TEXT}}),
        ("patch", f"{root}/{sid}", {"source_text": LONG_TEXT}),
        ("patch", f"{root}/{sid}", {"description": LONG_TEXT}),
        ("patch", f"{root}/{sid}", {"translations": {"en": LONG_TEXT}}),
        ("put", f"{root}/{sid}/translations/en", {"value": LONG_TEXT}),
    ]
    for method, url, body in cases:
        response = getattr(client, method)(url, json=body)
        assert response.status_code == 422, (method, body.keys(), response.text[:400])
        assert str(TEXT_MAX_LENGTH) in response.text
    stored = client.get(f"{root}/{sid}").json()
    assert stored["source_text"] == "Hi"
    assert stored["description"] is None
    assert all(item["value"] != LONG_TEXT for item in stored["translations"])
    assert client.get(root).json()["total"] == 1


def test_overlong_expected_baseline_can_be_replaced(client, session_factory):
    project = make_project(client, layout="flat")
    root = f"/api/projects/{project['id']}/strings"
    created = client.post(
        root, json={"key": "hello", "source_text": "Hi", "translations": {"en": "Hello"}}
    )
    assert created.status_code == 201, created.text
    sid = uuid.UUID(created.json()["id"])
    with session_factory() as db:
        row = db.get(StringEntry, sid)
        assert row is not None
        row.source_text = LONG_TEXT
        row.description = LONG_TEXT
        by_locale = {item.locale: item for item in row.translations}
        if "en" in by_locale:
            by_locale["en"].value = LONG_TEXT
        else:
            row.translations.append(Translation(locale="en", value=LONG_TEXT))
        db.commit()

    url = f"{root}/{sid}"
    replaced = client.patch(
        url,
        json={
            "source_text": "Short",
            "description": "Short",
            "translations": {"en": "Short"},
            "expected": {
                "source_text": LONG_TEXT,
                "description": LONG_TEXT,
                "translations": {"en": LONG_TEXT},
            },
        },
    )
    assert replaced.status_code == 200, replaced.text[:400]
    body = replaced.json()
    assert body["source_text"] == "Short"
    assert body["description"] == "Short"
    assert {item["locale"]: item["value"] for item in body["translations"]}["en"] == "Short"

    rejected = client.patch(
        url, json={"source_text": LONG_TEXT, "expected": {"source_text": "Short"}}
    )
    assert rejected.status_code == 422, rejected.text[:400]
    assert client.get(url).json()["source_text"] == "Short"

    with session_factory() as db:
        row = db.get(StringEntry, sid)
        assert row is not None
        by_locale = {item.locale: item for item in row.translations}
        by_locale["en"].value = LONG_TEXT
        db.commit()
    replaced_translation = client.put(
        f"{url}/translations/en", json={"value": "Short", "expected": LONG_TEXT}
    )
    assert replaced_translation.status_code == 200, replaced_translation.text[:400]
    too_long = client.put(
        f"{url}/translations/en", json={"value": LONG_TEXT, "expected": "Short"}
    )
    assert too_long.status_code == 422, too_long.text[:400]


def test_string_text_at_limit_is_accepted(client):
    project = make_project(client, layout="flat")
    root = f"/api/projects/{project['id']}/strings"
    created = client.post(
        root,
        json={
            "key": "max",
            "source_text": MAX_TEXT,
            "description": MAX_TEXT,
            "translations": {"en": MAX_TEXT},
        },
    )
    assert created.status_code == 201, created.text[:400]
    body = created.json()
    assert len(body["source_text"]) == TEXT_MAX_LENGTH
    assert len(body["description"]) == TEXT_MAX_LENGTH
    assert len(body["translations"][0]["value"]) == TEXT_MAX_LENGTH
    updated = client.patch(
        f"{root}/{body['id']}",
        json={"source_text": MAX_TEXT, "description": MAX_TEXT, "translations": {"en": MAX_TEXT}},
    )
    assert updated.status_code == 200, updated.text[:400]
    put = client.put(f"{root}/{body['id']}/translations/en", json={"value": MAX_TEXT})
    assert put.status_code == 200, put.text[:400]
    assert len(put.json()["translations"][0]["value"]) == TEXT_MAX_LENGTH


def test_oversized_upload_is_400_before_workbook(client, monkeypatch):
    project = make_project(client, layout="flat")

    def fail_parse(*_args, **_kwargs):
        raise AssertionError("parse_import_file called")

    def fail_workbook(*_args, **_kwargs):
        raise AssertionError("workbook opened")

    monkeypatch.setattr("app.routers.sync.parse_import_file", fail_parse)
    monkeypatch.setattr("app.excel.load_workbook", fail_workbook)
    response = client.post(
        f"/api/projects/{project['id']}/import",
        files={
            "file": (
                "big.xlsx",
                b"x" * (IMPORT_MAX_BYTES + 1),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 400, response.text
    assert str(IMPORT_MAX_BYTES) in response.text
    assert client.get(f"/api/projects/{project['id']}/strings").json()["total"] == 0


def test_import_rejects_overlong_translation_without_writes(client, monkeypatch):
    project = make_project(client, layout="flat")
    root = f"/api/projects/{project['id']}"
    uploaded = client.post(
        f"{root}/import",
        files=json_upload({"en": {"ok": "Hello", "bad": LONG_TEXT}}),
    )
    assert uploaded.status_code == 400, uploaded.text
    assert str(TEXT_MAX_LENGTH) in uploaded.text
    assert client.get(f"{root}/strings").json()["total"] == 0

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["key", "vi", "en"])
    sheet.append(["ok", "Hello", "English"])
    sheet.append(["bad", "Hello", "English"])
    sheet.cell(3, 3)._value = LONG_TEXT
    monkeypatch.setattr("app.excel.load_workbook", lambda *_args, **_kwargs: workbook)
    excel = client.post(
        f"{root}/import",
        files={"file": ("strings.xlsx", b"placeholder", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert excel.status_code == 400, excel.text
    assert str(TEXT_MAX_LENGTH) in excel.text
    assert client.get(f"{root}/strings").json()["total"] == 0

    described = Workbook()
    described_sheet = described.active
    described_sheet.append(["key", "vi", "description"])
    described_sheet.append(["ok", "Hello", "Context"])
    described_sheet.cell(2, 3)._value = LONG_TEXT

    def fail_rows(*_args, **_kwargs):
        raise AssertionError("import_rows called")

    monkeypatch.setattr("app.excel.load_workbook", lambda *_args, **_kwargs: described)
    monkeypatch.setattr("app.excel.import_rows", fail_rows)
    description = client.post(
        f"{root}/import",
        files={"file": ("described.xlsx", b"placeholder", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert description.status_code == 400, description.text
    assert str(TEXT_MAX_LENGTH) in description.text
    assert client.get(f"{root}/strings").json()["total"] == 0
