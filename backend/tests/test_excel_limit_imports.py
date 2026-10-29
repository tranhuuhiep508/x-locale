"""Every Excel entry point rejects resource excess before any content writes."""

import io
import re
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from openpyxl import Workbook

from app import excel_limits
from app.limits import TEXT_MAX_LENGTH
from tests.helpers import make_project
from tests.test_excel import _replace_inline_cell
from tests.test_excel_limits import archive_bytes, shared_string_table
from tests.test_import_bulk import content_writes, sql_calls
from tests.test_import_concurrency import catalog_state, seed_import_project
from tests.test_import_jobs import poll


def _use_shared_strings(parts, cells, values, *, rich=False):
    name = "xl/worksheets/sheet1.xml"
    for ref, text, index in cells:
        needle = f'<c r="{ref}" t="inlineStr"><is><t>{text}</t></is></c>'.encode()
        assert needle in parts[name]
        replacement = f'<c r="{ref}" t="s"><v>{index}</v></c>'.encode()
        parts[name] = parts[name].replace(needle, replacement)
    parts["xl/sharedStrings.xml"] = shared_string_table(values, rich=rich).encode()
    parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(
        b"</Types>",
        b'<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>',
    )


@pytest.mark.parametrize("dry_run", [False, True])
@pytest.mark.parametrize("as_job", [False, True])
@pytest.mark.parametrize(
    "failure",
    [
        "text",
        "dimensions",
        "expanded",
        "sheet-type",
        "shared-formula",
        "shared-string",
        "renamed-sheet",
        "unqualified-id",
        "duplicate-name",
    ],
)
def test_resource_failure_never_opens_workbook_or_writes_content(
    client, session_factory, database_engine, monkeypatch, dry_run, as_job, failure
):
    project = seed_import_project(client, session_factory)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "new-module"
    sheet.append(["key", "vi", "en", "tags", "description"])
    sheet.append(["valid", "Source", "English", "new-tag", "Context"])
    sheet.append(["invalid", "MARKER", "English", "new-tag", "Context"])
    if failure in {"sheet-type", "renamed-sheet", "unqualified-id", "duplicate-name"}:
        workbook.create_sheet("alias")
    raw = io.BytesIO()
    workbook.save(raw)
    with ZipFile(raw) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    name = "xl/worksheets/sheet1.xml"
    if failure == "text":
        parts[name] = parts[name].replace(b"MARKER", b"x" * (TEXT_MAX_LENGTH + 1))
        message = "cell/shared-string characters"
    elif failure == "dimensions":
        parts[name] = re.sub(b"<dimension[^>]*/>", b'<dimension ref="A1:XFD1048576"/>', parts[name])
        message = "columns"
    elif failure == "expanded":
        monkeypatch.setattr(excel_limits, "EXCEL_MAX_MEMBER_BYTES", 4096)
        parts["oversized.bin"] = b"x" * 8192
        message = "expanded member bytes"
    elif failure == "duplicate-name":
        parts["xl/workbook.xml"] = parts["xl/workbook.xml"].replace(
            b'name="alias"', b'name="new-module"'
        )
        message = "distinct names"
    elif failure in {"sheet-type", "renamed-sheet", "unqualified-id"}:
        rels = ElementTree.fromstring(parts["xl/_rels/workbook.xml.rels"])
        for rel in rels:
            if rel.attrib["Type"].endswith("/worksheet"):
                rel.attrib["Type"] = rel.attrib["Type"].replace("/worksheet", "/styles")
                rel.attrib["Target"] = "/xl/worksheets/sheet1.xml"
        parts["xl/_rels/workbook.xml.rels"] = ElementTree.tostring(rels)
        del parts["xl/worksheets/sheet2.xml"]
        message = "internal worksheet relationships"
        if failure != "sheet-type":
            namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
            rel_id = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
            root = ElementTree.fromstring(parts["xl/workbook.xml"])
            for declaration in root.find(f"{{{namespace}}}sheets"):
                if failure == "renamed-sheet":
                    declaration.tag = f"{{{namespace}}}renamed"
                else:
                    declaration.set("id", declaration.attrib.pop(rel_id))
            parts["xl/workbook.xml"] = ElementTree.tostring(root)
            message = "Invalid Excel worksheet declaration"
    elif failure == "shared-formula":
        monkeypatch.setattr(excel_limits, "EXCEL_MAX_DECODED_CHARACTERS", 200)
        for row, text in [(2, "Source"), (3, "MARKER")]:
            needle = f'<c r="B{row}" t="inlineStr"><is><t>{text}</t></is></c>'.encode()
            formula = '"' + "x" * 100 + '"' if row == 2 else ""
            replacement = f'<c r="B{row}"><f t="shared" si="0">{formula}</f><v>0</v></c>'.encode()
            assert needle in parts[name]
            parts[name] = parts[name].replace(needle, replacement)
        message = "decoded characters"
    else:
        monkeypatch.setattr(excel_limits, "EXCEL_MAX_DECODED_CHARACTERS", 300)
        _use_shared_strings(parts, [("E2", "Context", 0), ("E3", "Context", 0)], [" " + "x" * 99])
        message = "decoded characters"

    def fail_load(*_args, **_kwargs):
        raise AssertionError("openpyxl ran before resource validation")

    monkeypatch.setattr("app.excel.load_workbook", fail_load)
    before = catalog_state(session_factory, project)
    with sql_calls(database_engine) as calls:
        response = client.post(
            f"/api/projects/{project['id']}/import",
            params={"dry_run": dry_run, "status": "public"},
            headers={"Prefer": "respond-async"} if as_job else {},
            files={"file": ("bounded.xlsx", archive_bytes(parts))},
        )
        if as_job:
            assert response.status_code == 202, response.text
            job = poll(client, project, "xlsx", response.json()["job_id"])
            assert job["status"] == "failed", job
            assert message in str(job["error"])
        else:
            assert response.status_code == 400, response.text
            assert message in response.json()["detail"]
    assert not content_writes(calls)
    assert catalog_state(session_factory, project) == before


@pytest.mark.parametrize("dry_run", [False, True])
@pytest.mark.parametrize("as_job", [False, True])
def test_bounded_shared_formulas_still_import_as_text(client, dry_run, as_job):
    project = make_project(client, layout="flat")
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["key", "vi"])
    sheet.append(["first", "FIRST"])
    sheet.append(["second", "SECOND"])
    raw = io.BytesIO()
    workbook.save(raw)
    content = _replace_inline_cell(
        raw.getvalue(), "B2", "FIRST", '<f t="shared" si="0">A2+1</f><v>0</v>'
    )
    content = _replace_inline_cell(content, "B3", "SECOND", '<f t="shared" si="0"/><v>0</v>')
    response = client.post(
        f"/api/projects/{project['id']}/import",
        params={"dry_run": dry_run},
        headers={"Prefer": "respond-async"} if as_job else {},
        files={"file": ("shared.xlsx", content)},
    )
    if as_job:
        assert response.status_code == 202, response.text
        job = poll(client, project, "xlsx", response.json()["job_id"])
        assert job["status"] == "completed", job
        result = job["result"]
    else:
        assert response.status_code == 200, response.text
        result = response.json()
    assert result["created"] == (0 if dry_run else 2)
    assert result["diff"]["create_count"] == 2
    assert {item["key"]: item["source_text"] for item in result["diff"]["create"]} == {
        "first": "=A2+1",
        "second": "=A3+1",
    }
    items = client.get(f"/api/projects/{project['id']}/strings").json()["items"]
    assert {item["key"]: item["source_text"] for item in items} == (
        {} if dry_run else {"first": "=A2+1", "second": "=A3+1"}
    )


@pytest.mark.parametrize("dry_run", [False, True])
@pytest.mark.parametrize("as_job", [False, True])
@pytest.mark.parametrize("rich", [False, True])
@pytest.mark.parametrize("table_first", [False, True])
def test_bounded_shared_descriptions_still_import(client, dry_run, as_job, rich, table_first):
    project = make_project(client, layout="flat")
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["key", "vi", "description"])
    sheet.append(["first", "First source", "DESCRIPTION"])
    sheet.append(["second", "Second source", "DESCRIPTION"])
    raw = io.BytesIO()
    workbook.save(raw)
    with ZipFile(raw) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    _use_shared_strings(
        parts, [("C2", "DESCRIPTION", 0), ("C3", "DESCRIPTION", 0)], [" shared text "], rich=rich
    )
    if table_first:
        parts = dict(reversed(list(parts.items())))
    response = client.post(
        f"/api/projects/{project['id']}/import",
        params={"dry_run": dry_run},
        headers={"Prefer": "respond-async"} if as_job else {},
        files={"file": ("shared.xlsx", archive_bytes(parts))},
    )
    if as_job:
        assert response.status_code == 202, response.text
        job = poll(client, project, "xlsx", response.json()["job_id"])
        assert job["status"] == "completed", job
        result = job["result"]
    else:
        assert response.status_code == 200, response.text
        result = response.json()
    assert result["created"] == (0 if dry_run else 2)
    assert result["diff"]["create_count"] == 2
    items = client.get(f"/api/projects/{project['id']}/strings").json()["items"]
    assert {item["key"]: item["description"] for item in items} == (
        {} if dry_run else {"first": "shared text", "second": "shared text"}
    )


@pytest.mark.parametrize("dry_run", [False, True])
@pytest.mark.parametrize("as_job", [False, True])
@pytest.mark.parametrize("layout", ["flat", "modular"])
@pytest.mark.filterwarnings(
    "ignore:Unknown extension is not supported and will be removed:UserWarning"
)
def test_row_extension_metadata_preserves_imported_strings(client, dry_run, as_job, layout):
    project = make_project(client, layout=layout)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "common"
    sheet.append(["key", "vi", "en"])
    sheet.append(["hello", "Xin chào", "Hello"])
    raw = io.BytesIO()
    workbook.save(raw)
    with ZipFile(raw) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["xl/worksheets/sheet1.xml"] = parts["xl/worksheets/sheet1.xml"].replace(
        b"</row>",
        b'<extLst><ext uri="urn:test"><metadata xmlns="urn:test">'
        b'<row r="custom-row-id"/><c r="custom-cell-id"/>'
        b'<dimension ref="custom-range"/></metadata></ext></extLst></row>',
    )
    response = client.post(
        f"/api/projects/{project['id']}/import",
        params={"dry_run": dry_run},
        headers={"Prefer": "respond-async"} if as_job else {},
        files={"file": ("row-extensions.xlsx", archive_bytes(parts))},
    )
    if as_job:
        assert response.status_code == 202, response.text
        job = poll(client, project, "xlsx", response.json()["job_id"])
        assert job["status"] == "completed", job
        result = job["result"]
    else:
        assert response.status_code == 200, response.text
        result = response.json()
    assert result["created"] == (0 if dry_run else 1)
    assert result["diff"]["create_count"] == 1
    assert result["diff"]["create"][0]["source_text"] == "Xin chào"
    items = client.get(f"/api/projects/{project['id']}/strings").json()["items"]
    if dry_run:
        assert items == []
    else:
        assert len(items) == 1
        assert (items[0]["key"], items[0]["source_text"]) == ("hello", "Xin chào")
        assert {
            translation["locale"]: translation["value"] for translation in items[0]["translations"]
        } == {"en": "Hello"}
