"""Base-language content belongs in source_text, never target translations."""

import io
import uuid
from unittest.mock import Mock

import pytest
from fastapi import HTTPException
from openpyxl import Workbook, load_workbook

from app.models import Activity, Job, JobStatus, Project, StringEntry, Translation
from app.schemas import BatchFilter
from app.services import translate
from app.services.strings import filtered_string_query
from tests.helpers import json_upload, make_project, publish_strings

ERROR = "Base language 'vi' must be edited through source_text."


@pytest.fixture
def catalog(client, session_factory):
    project = make_project(client)
    root = f"/api/projects/{project['id']}"
    row = client.post(
        f"{root}/strings", json={"key": "save", "source_text": "Save"}
    ).json()
    # Model a project created before target-list normalization, including retained data.
    with session_factory() as db:
        db.get(Project, uuid.UUID(project["id"])).target_languages = ["vi", "en"]
        db.add(Translation(string_id=uuid.UUID(row["id"]), locale="vi", value="Legacy"))
        db.commit()
    return project, root, row


def content_state(client, root, session_factory):
    with session_factory() as db:
        counts = [db.query(model).count() for model in (StringEntry, Translation, Activity)]
    return client.get(f"{root}/strings").json(), counts


@pytest.mark.parametrize("method", ["create", "patch", "put"])
@pytest.mark.parametrize("value", ["", "Legacy", "Lưu"])
def test_base_translation_rejection_is_atomic(client, session_factory, catalog, method, value):
    project, root, row = catalog
    publish_strings(client, project["id"], [row["id"]])
    before = content_state(client, root, session_factory)
    if method == "put":
        response = client.put(
            f"{root}/strings/{row['id']}/translations/vi", json={"value": value}
        )
    else:
        body = {
            "key": "new-key", "source_text": "Changed source", "status": "public",
            "translations": {"en": "Changed target", "vi": value},
        }
        response = (
            client.post(f"{root}/strings", json=body)
            if method == "create"
            else client.patch(f"{root}/strings/{row['id']}", json=body)
        )
    assert response.status_code == 400, response.text
    assert response.json()["detail"] == ERROR
    assert content_state(client, root, session_factory) == before


def test_legacy_targets_do_not_preallocate_base_rows_and_targets_still_work(client, catalog):
    _, root, _ = catalog
    created = client.post(
        f"{root}/strings",
        json={"key": "new", "source_text": "Source", "translations": {"en": "Created"}},
    )
    assert created.status_code == 201, created.text
    row = created.json()
    assert [(t["locale"], t["value"]) for t in row["translations"]] == [("en", "Created")]
    for method in ("patch", "put"):
        response = (
            client.patch(f"{root}/strings/{row['id']}", json={"translations": {"en": method}})
            if method == "patch"
            else client.put(f"{root}/strings/{row['id']}/translations/en", json={"value": method})
        )
        assert response.status_code == 200, response.text
        exported = client.get(f"{root}/export", params={"layout": "flat"}).json()
        assert exported["en"]["new"] == method
    unknown = client.put(f"{root}/strings/{row['id']}/translations/fr", json={"value": "Non"})
    assert unknown.status_code == 400
    assert unknown.json()["detail"] == "Locale 'fr' is not configured"


def test_project_targets_exclude_base_on_create_and_locale_updates(client):
    project = make_project(client, targets=["en", "vi", "en", "ja"])
    assert project["target_languages"] == ["en", "ja"]
    root = f"/api/projects/{project['id']}"
    changed = client.patch(root, json={"base_language": "en"})
    assert changed.status_code == 200, changed.text
    assert changed.json()["target_languages"] == ["ja"]
    changed = client.patch(root, json={"target_languages": ["ja", "en", "vi", "ja"]})
    assert changed.status_code == 200, changed.text
    assert changed.json()["target_languages"] == ["ja", "vi"]


@pytest.mark.parametrize("targets", [["vi", "en"], ["vi"]])
def test_missing_any_ignores_legacy_base_targets(client, session_factory, targets):
    project = make_project(client)
    project_id = uuid.UUID(project["id"])
    root = f"/api/projects/{project['id']}"
    with session_factory() as db:
        db.get(Project, project_id).target_languages = targets
        db.commit()
    for key, translations in [("complete", {"en": "Save"} if "en" in targets else {}), ("missing", {})]:
        response = client.post(
            f"{root}/strings",
            json={"key": key, "source_text": "Lưu", "translations": translations},
        )
        assert response.status_code == 201, response.text
    missing_id = response.json()["id"]

    def assert_missing(expected):
        response = client.get(f"{root}/strings", params={"missing_any": True})
        assert response.status_code == 200, response.text
        assert response.json()["total"] == len(expected)
        assert sorted(row["key"] for row in response.json()["items"]) == expected
        # Batch actions and AI scopes must select the same rows as the catalog.
        with session_factory() as db:
            query = filtered_string_query(db, db.get(Project, project_id), BatchFilter(missing_any=True))
            assert sorted(row.key for row in query.all()) == expected

    assert_missing(["missing"] if "en" in targets else [])
    if "en" in targets:
        saved = client.put(f"{root}/strings/{missing_id}/translations/en", json={"value": "Save"})
        assert saved.status_code == 200, saved.text
        assert_missing([])


@pytest.mark.parametrize("kind", ["cli", "json", "xlsx"])
def test_source_imports_work_with_legacy_targets_without_creating_base_rows(client, catalog, kind):
    _, root, old_row = catalog
    values = {"save": "Lưu", "imported": "Mới"}
    if kind == "cli":
        response = client.post(f"{root}/strings/import", json={"strings": values, "base_language": "vi"})
    elif kind == "json":
        response = client.post(f"{root}/import", files=json_upload({"vi": values}))
    else:
        workbook = Workbook()
        workbook.active.title = "_unassigned"
        workbook.active.append(["key", "description", "tags", "vi", "en"])
        for key, value in values.items():
            workbook.active.append([key, None, None, value, ""])
        raw = io.BytesIO()
        workbook.save(raw)
        workbook.close()
        response = client.post(f"{root}/import", files={"file": ("source.xlsx", raw.getvalue())})
    assert response.status_code == 200, response.text
    rows = {row["key"]: row for row in client.get(f"{root}/strings").json()["items"]}
    assert rows["save"]["source_text"] == "Lưu"
    assert rows["imported"]["source_text"] == "Mới"
    assert all(t["locale"] != "vi" for t in rows["imported"]["translations"])
    legacy = client.get(f"{root}/strings/{old_row['id']}").json()
    assert next(t["value"] for t in legacy["translations"] if t["locale"] == "vi") == "Legacy"


@pytest.mark.parametrize("layout,format", [("flat", "json"), ("modular", "json"), ("flat", "xlsx")])
def test_source_corrections_respect_publication(client, layout, format):
    project = make_project(client, layout=layout)
    root = f"/api/projects/{project['id']}"
    row = client.post(f"{root}/strings", json={"key": "save", "source_text": "Save"}).json()
    publish_strings(client, project["id"], [row["id"]])

    def source(stage):
        response = client.get(
            f"{root}/export", params={"layout": layout, "format": format, "stage": stage, "locale": "vi"}
        )
        assert response.status_code == 200, response.text
        if format == "xlsx":
            workbook = load_workbook(io.BytesIO(response.content))
            sheet = workbook["strings"]
            result = sheet.cell(2, 4).value
            workbook.close()
            return result
        payload = response.json()
        return (payload["unassigned"] if layout == "modular" else payload)["vi"]["save"]

    changed = client.patch(f"{root}/strings/{row['id']}", json={"source_text": "Lưu"})
    assert changed.status_code == 200, changed.text
    assert source("draft") == "Lưu"
    assert source("public") == "Save"
    publish_strings(client, project["id"], [row["id"]])
    assert source("public") == "Lưu"


@pytest.mark.parametrize("endpoint", ["", "/preview", "/missing", "/proposals"])
def test_ai_rejects_base_before_generation_or_job_creation(
    client, session_factory, catalog, monkeypatch, endpoint
):
    _, root, _ = catalog
    generate = Mock(side_effect=AssertionError("AI must not run"))
    monkeypatch.setattr(translate, "translate_batch", generate)
    before = content_state(client, root, session_factory)
    response = client.post(
        f"{root}/translate{endpoint}", json={"source_text": "Save", "locales": ["en", "vi"]}
    )
    assert response.status_code == 400, response.text
    assert response.json()["detail"] == ERROR
    generate.assert_not_called()
    assert content_state(client, root, session_factory) == before
    with session_factory() as db:
        assert db.query(Job).count() == 0


@pytest.mark.parametrize("endpoint", ["", "/preview", "/missing", "/proposals"])
@pytest.mark.parametrize("locales", [None, []])
def test_ai_defaults_filter_legacy_base_target(client, catalog, monkeypatch, endpoint, locales):
    _, root, row = catalog

    def generate(source, items, on_progress=None):
        assert source == "vi"
        assert all(item.locales == ("en",) for item in items)
        return {item.id: {"en": "Save target"} for item in items}

    monkeypatch.setattr(translate, "translate_batch", generate)
    body = {"source_text": "Save"}
    if locales is not None:
        body["locales"] = locales
    response = client.post(f"{root}/translate{endpoint}", json=body)
    assert response.status_code == 200, response.text
    if endpoint == "/preview":
        assert response.json()["translations"] == {"en": "Save target"}
    else:
        assert response.json()["locales"] == ["en"]
    saved = client.get(f"{root}/strings/{row['id']}").json()
    assert next(t["value"] for t in saved["translations"] if t["locale"] == "vi") == "Legacy"


@pytest.mark.parametrize("value", ["", "Legacy", "Lưu"])
def test_proposal_apply_rejects_entire_batch_before_description_writes(
    client, session_factory, catalog, value
):
    _, root, row = catalog
    before = content_state(client, root, session_factory)
    common = {"string_id": row["id"], "expected_source_text": "Save"}
    response = client.post(f"{root}/translate/apply", json={"items": [
        {**common, "translations": {"en": "Valid"}, "description": "Changed", "expected_description": None},
        {**common, "translations": {"vi": value}},
    ]})
    assert response.status_code == 400, response.text
    assert response.json()["detail"] == ERROR
    assert content_state(client, root, session_factory) == before
    # Preserve the existing silent skip for other unconfigured proposal locales.
    response = client.post(f"{root}/translate/apply", json={"items": [
        {**common, "translations": {"fr": "Ignored", "en": "Valid"}},
    ]})
    assert response.status_code == 200, response.text
    assert response.json()["locales"] == ["en"]


@pytest.mark.parametrize("kind", ["translate", "translate_proposals"])
def test_queued_base_locale_job_fails_without_writing_content(
    client, session_factory, catalog, monkeypatch, kind
):
    project, root, row = catalog
    generate = Mock(side_effect=AssertionError("AI must not run"))
    monkeypatch.setattr(translate, "translate_batch", generate)
    before = content_state(client, root, session_factory)
    with session_factory() as db:
        job = Job(project_id=uuid.UUID(project["id"]), kind=kind, status=JobStatus.pending, payload={})
        db.add(job)
        db.commit()
        job_id = job.id
    args = (uuid.UUID(project["id"]), [uuid.UUID(row["id"])], ["en", "vi"], False)
    with pytest.raises(HTTPException, match="source_text"):
        if kind == "translate":
            translate.run_translate_job(*args, uuid.uuid4(), job_id)
        else:
            translate.run_propose_job(*args, job_id)
    generate.assert_not_called()
    assert content_state(client, root, session_factory) == before
    with session_factory() as db:
        job = db.get(Job, job_id)
        assert job.status == JobStatus.failed
        assert ERROR in job.error
        assert job.completed_at is not None
