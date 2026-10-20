"""Whole-payload validation, bounded SQL, and consolidated Excel undo."""

import io
import uuid
from contextlib import contextmanager

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import event, func, select

from app import excel
from app.models import Activity, Module, StringEntry, StringTag, Tag, Translation
from tests.helpers import json_upload, make_project, publish_strings


@contextmanager
def sql_calls(engine):
    calls = []

    def record(_conn, _cursor, statement, _parameters, _context, many):
        calls.append((statement.lower(), many))

    event.listen(engine, "before_cursor_execute", record)
    try:
        yield calls
    finally:
        event.remove(engine, "before_cursor_execute", record)


def content_writes(calls):
    tables = ("strings", "translations", "string_tags", "modules", "tags", "activities")
    return [
        sql
        for sql, _ in calls
        if any(
            sql.startswith(f"{verb} {table} ")
            for verb in ("insert into", "update", "delete from")
            for table in tables
        )
    ]


def counts(session_factory):
    with session_factory() as db:
        return [
            db.scalar(select(func.count()).select_from(model))
            for model in (StringEntry, Translation, Module, Tag, StringTag, Activity)
        ]


@pytest.mark.parametrize("kind", ["json", "cli"])
@pytest.mark.parametrize("field", ["key", "value", "locale", "module", "long-key"])
def test_late_json_nul_rejects_before_any_content_write(
    client, database_engine, session_factory, kind, field
):
    project = make_project(client)
    bad = {"vi": {"bad": "Bad"}}
    slug = "last"
    if field == "key":
        bad = {"vi": {"k" * 511 + "\x00": "Bad"}}
    elif field == "long-key":
        bad = {"vi": {"k" * 513: "Bad"}}
    elif field == "value":
        bad = {"vi": {"bad": "Value\x00"}}
    elif field == "locale":
        bad = {"vi\x00": {"bad": "Bad"}}
    else:
        slug = "last\x00"
    data = {"modules": {"first": {"vi": {"good": "Good"}}, slug: bad}}
    before = counts(session_factory)
    with sql_calls(database_engine) as calls:
        root = f"/api/projects/{project['id']}"
        response = (
            client.post(f"{root}/strings/import", json=data)
            if kind == "cli"
            else client.post(f"{root}/import", files=json_upload(data))
        )
    assert response.status_code == 400, response.text
    assert (
        "at most 512" if field == "long-key" else "contains a NUL (\\u0000) character"
    ) in response.json()["detail"]
    assert not content_writes(calls)
    assert counts(session_factory) == before


@pytest.mark.parametrize("field", ["key", "value", "description", "tag", "locale", "module"])
@pytest.mark.parametrize("dry_run", [False, True])
def test_decoded_excel_nul_rejects_before_writes(
    client, database_engine, session_factory, monkeypatch, field, dry_run
):
    project = make_project(client)
    workbook = Workbook()
    first = workbook.active
    first.title = "first"
    first.append(["key", "vi", "en", "description", "tags"])
    first.append(["good", "Good", "English", "Context", "new-tag"])
    last = workbook.create_sheet("last")
    last.append(["key", "vi", "en", "description", "tags"])
    last.append(["bad", "Bad", "English", "Context", "new-tag"])
    if field == "module":
        last.title = "last\x00"
    elif field == "locale":
        last.cell(1, 3)._value = "en\x00"
    else:
        column = {"key": 1, "value": 2, "description": 4, "tag": 5}[field]
        last.cell(2, column)._value = "Bad\x00"
    # XML cannot encode literal NUL; exercise validation at the decoded-cell boundary.
    monkeypatch.setattr(excel, "load_workbook", lambda *_args, **_kwargs: workbook)
    before = counts(session_factory)
    with sql_calls(database_engine) as calls:
        response = client.post(
            f"/api/projects/{project['id']}/import",
            params={"dry_run": dry_run},
            files={"file": ("nul.xlsx", b"decoded workbook")},
        )
    assert response.status_code == 400, response.text
    assert "contains a NUL (\\u0000) character" in response.json()["detail"]
    assert "Sheet" in response.json()["detail"]
    if field not in {"module", "locale"}:
        assert "row 2" in response.json()["detail"]
    assert not content_writes(calls)
    assert counts(session_factory) == before


@pytest.mark.parametrize("size", [100, 1100])
def test_import_sql_grows_by_chunks_and_no_change_writes_nothing(client, database_engine, size):
    project = make_project(client, targets=["en", "fr"])
    payload = {
        locale: {f"key.{i}": f"{locale} {i}" for i in range(size)} for locale in ("vi", "en", "fr")
    }
    path = f"/api/projects/{project['id']}/import"
    with sql_calls(database_engine) as calls:
        imported = client.post(path, files=json_upload(payload))
    assert imported.status_code == 200, imported.text
    assert imported.json()["created"] == size
    assert len(calls) <= 25
    assert not any(many for _, many in calls)
    with sql_calls(database_engine) as calls:
        unchanged = client.post(path, files=json_upload(payload))
    assert unchanged.status_code == 200
    assert unchanged.json()["created"] == unchanged.json()["updated"] == 0
    assert len(calls) <= 15
    assert not content_writes(calls)


@pytest.mark.parametrize("layout", ["flat", "modular"])
@pytest.mark.parametrize(
    "new,source_column,source,target,expected_source",
    [
        (False, False, None, "Hello", "Xin chào"),
        (False, False, None, "Hi", "Xin chào"),
        (True, False, None, "Hello", "new"),
        (False, True, None, "Hello", ""),
        (True, True, None, "Hello", ""),
        (False, True, "Chào bạn", "Hello", "Chào bạn"),
        (True, True, "Chào bạn", "Hello", "Chào bạn"),
    ],
    ids=[
        "unchanged",
        "target-edit",
        "new-target-only",
        "blank",
        "new-blank",
        "source",
        "new-source",
    ],
)
def test_excel_source_presence_preview_and_apply(
    client,
    session_factory,
    database_engine,
    layout,
    new,
    source_column,
    source,
    target,
    expected_source,
):
    project = make_project(client, layout=layout)
    root = f"/api/projects/{project['id']}"
    module_id = None
    if layout == "modular":
        module_id = client.post(
            f"{root}/modules", json={"slug": "common", "name": "Common"}
        ).json()["id"]
    created = client.post(
        f"{root}/strings",
        json={
            "key": "greeting",
            "source_text": "Xin chào",
            "module_id": module_id,
            "translations": {"en": "Hello"},
        },
    )
    assert created.status_code == 201, created.text
    publish_strings(client, project["id"], [created.json()["id"]])
    exported = client.get(
        f"{root}/export", params={"format": "xlsx", "stage": "draft", "locale": "en"}
    )
    assert exported.status_code == 200, exported.text
    workbook = load_workbook(io.BytesIO(exported.content))
    sheet = workbook["common" if layout == "modular" else excel.FLAT_SHEET]
    headers = [cell.value for cell in sheet[1]]
    assert "vi" not in headers
    key = "new" if new else "greeting"
    sheet.cell(2, headers.index("key") + 1, key)
    sheet.cell(2, headers.index("en") + 1, target)
    if source_column:
        column = len(headers) + 1
        sheet.cell(1, column, "vi")
        sheet.cell(2, column).value = source
    raw = io.BytesIO()
    workbook.save(raw)
    workbook.close()
    files = {"file": ("source-presence.xlsx", raw.getvalue())}

    before_counts = counts(session_factory)
    with sql_calls(database_engine) as calls:
        preview = client.post(f"{root}/import", params={"dry_run": True}, files=files)
    assert preview.status_code == 200, preview.text
    assert not content_writes(calls)
    assert counts(session_factory) == before_counts
    diff = preview.json()["diff"]
    assert diff["create_count"] == int(new)
    # Excel intentionally counts source/description edits, not translation-only edits.
    assert diff["update_count"] == int(not new and expected_source != "Xin chào")
    items = diff["create"] + diff["update"]
    if items:
        assert [item["source_text"] for item in items] == [expected_source]

    unchanged = not new and not source_column and target == "Hello"
    with sql_calls(database_engine) as calls:
        applied = client.post(f"{root}/import", files=files)
    assert applied.status_code == 200, applied.text
    assert applied.json()["diff"] == diff
    assert applied.json()["created"] == diff["create_count"]
    assert applied.json()["updated"] == diff["update_count"]
    if unchanged:
        assert not content_writes(calls)
    with session_factory() as db:
        entry = db.scalar(
            select(StringEntry).where(
                StringEntry.project_id == uuid.UUID(project["id"]), StringEntry.key == key
            )
        )
        assert entry.source_text == expected_source
        assert entry.module_id == (uuid.UUID(module_id) if module_id else None)
        assert entry.translations[0].value == target
        assert entry.published_source_text == (None if new else "Xin chào")
        assert entry.translations[0].published_value == (None if new else "Hello")
        activity_count = db.scalar(
            select(func.count())
            .select_from(Activity)
            .where(Activity.batch_id == uuid.UUID(applied.json()["batch_id"]))
        )
        assert activity_count == (0 if unchanged else 1)


@pytest.mark.parametrize("layout", ["flat", "modular"])
@pytest.mark.parametrize("status", ["draft", "public"])
def test_unchanged_tagged_excel_preserves_metadata_and_writes_nothing(
    client, session_factory, database_engine, layout, status
):
    project = make_project(client, layout=layout)
    root = f"/api/projects/{project['id']}"
    module_id = None
    if layout == "modular":
        module_id = client.post(
            f"{root}/modules", json={"slug": "stable", "name": "Stable"}
        ).json()["id"]
    tag = client.post(f"{root}/tags", json={"name": "stable"}).json()
    created = client.post(
        f"{root}/strings",
        json={
            "key": "same",
            "source_text": "Same",
            "description": "Context",
            "module_id": module_id,
            "tag_ids": [tag["id"]],
            "translations": {"en": "English"},
        },
    ).json()
    if status == "public":
        publish_strings(client, project["id"], [created["id"]])
    key = client.post(f"{root}/api-keys", json={"name": "Spreadsheet importer"}).json()["key"]
    exported = client.get(f"{root}/export", params={"format": "xlsx", "stage": status})
    assert exported.status_code == 200, exported.text
    with session_factory() as db:
        before = dict(
            db.execute(
                select(StringEntry.__table__).where(StringEntry.id == uuid.UUID(created["id"]))
            )
            .mappings()
            .one()
        )
    with sql_calls(database_engine) as calls:
        imported = client.post(
            f"{root}/import",
            params={"status": status},
            headers={"X-API-Key": key},
            files={"file": ("same.xlsx", exported.content)},
        )
    assert imported.status_code == 200, imported.text
    assert imported.json()["created"] == imported.json()["updated"] == 0
    assert not content_writes(calls)
    with session_factory() as db:
        after = dict(
            db.execute(
                select(StringEntry.__table__).where(StringEntry.id == uuid.UUID(created["id"]))
            )
            .mappings()
            .one()
        )
        assert after == before
        assert (
            db.scalar(
                select(func.count())
                .select_from(Activity)
                .where(Activity.batch_id == uuid.UUID(imported.json()["batch_id"]))
            )
            == 0
        )


def test_repeated_excel_rows_create_one_final_event_and_undo_exactly(client, session_factory):
    project = make_project(client, layout="flat")
    root = f"/api/projects/{project['id']}"
    created = client.post(
        f"{root}/strings",
        json={"key": "existing", "source_text": "Original", "translations": {"en": "Original EN"}},
    ).json()
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["key", "vi", "en"])
    for row in (
        ("existing", "Intermediate", "Middle EN"),
        ("existing", "Final", "Final EN"),
        ("new", "First", "First EN"),
        ("new", "Last", "Last EN"),
    ):
        sheet.append(row)
    raw = io.BytesIO()
    workbook.save(raw)
    imported = client.post(f"{root}/import", files={"file": ("repeated.xlsx", raw.getvalue())})
    assert imported.status_code == 200, imported.text
    result = imported.json()
    assert (result["created"], result["updated"], result["total"]) == (1, 3, 4)
    with session_factory() as db:
        audit = db.scalars(
            select(Activity).where(Activity.batch_id == uuid.UUID(result["batch_id"]))
        ).all()
        assert len(audit) == 2
        prior = next(row for row in audit if row.before)
        assert prior.before["source_text"] == "Original"
        assert prior.after["source_text"] == "Final"
        assert prior.before["translations"]["en"] == "Original EN"
        assert prior.after["translations"]["en"] == "Final EN"
    undone = client.post(f"{root}/activities/batch/{result['batch_id']}/revert")
    assert undone.status_code == 200, undone.text
    rows = client.get(f"{root}/strings").json()["items"]
    assert len(rows) == 1 and rows[0]["id"] == created["id"]
    assert rows[0]["source_text"] == "Original"
    assert next(t for t in rows[0]["translations"] if t["locale"] == "en")["value"] == "Original EN"


@pytest.mark.parametrize("kind", ["json", "xlsx"])
@pytest.mark.parametrize("status", ["draft", "public"])
def test_import_publishes_final_state_only_when_requested(
    client, session_factory, database_engine, kind, status
):
    project = make_project(client, layout="flat")
    root = f"/api/projects/{project['id']}"
    existing = client.post(
        f"{root}/strings",
        json={"key": "existing", "source_text": "Original", "translations": {"en": "Original EN"}},
    ).json()
    publish_strings(client, project["id"], [existing["id"]])
    if kind == "json":
        files = json_upload(
            {
                "vi": {"existing": "Final", "new": "Final"},
                "en": {"existing": "Final EN", "new": "Final EN"},
            }
        )
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["key", "vi", "en"])
        for key in ("existing", "new"):
            sheet.append([key, "First", "First EN"])
            sheet.append([key, "Final", "Final EN"])
        raw = io.BytesIO()
        workbook.save(raw)
        files = {"file": ("publish.xlsx", raw.getvalue())}
    imported = client.post(f"{root}/import", params={"status": status}, files=files)
    assert imported.status_code == 200, imported.text
    result = imported.json()
    with session_factory() as db:
        entries = db.scalars(
            select(StringEntry).where(StringEntry.project_id == uuid.UUID(project["id"]))
        ).all()
        for entry in entries:
            assert entry.source_text == "Final"
            translation = next(t for t in entry.translations if t.locale == "en")
            assert translation.value == "Final EN"
            if status == "public":
                assert entry.status.value == "public"
                assert entry.published_source_text == "Final"
                assert translation.published_value == "Final EN"
            elif entry.key == "existing":
                assert entry.published_source_text == "Original"
                assert translation.published_value == "Original EN"
            else:
                assert entry.status.value == "draft"
                assert entry.published_source_text is None
                assert translation.published_value is None
        activities = db.scalars(
            select(Activity).where(Activity.batch_id == uuid.UUID(result["batch_id"]))
        ).all()
        assert len(activities) == 2
        prior = next(activity for activity in activities if activity.before is not None)
        assert prior.before["source_text"] == prior.before["published_source_text"] == "Original"
        assert prior.after["source_text"] == "Final"
        assert prior.after["published_source_text"] == (
            "Final" if status == "public" else "Original"
        )
    # An unchanged final snapshot must not produce writes or timestamp-only activities.
    if kind == "json":
        with sql_calls(database_engine) as calls:
            repeated = client.post(f"{root}/import", params={"status": status}, files=files)
        assert repeated.status_code == 200
        assert not content_writes(calls)
    undone = client.post(f"{root}/activities/batch/{result['batch_id']}/revert")
    assert undone.status_code == 200, undone.text
    with session_factory() as db:
        entries = {
            entry.key: entry
            for entry in db.scalars(
                select(StringEntry).where(StringEntry.project_id == uuid.UUID(project["id"]))
            )
        }
        entry = entries["existing"]
        assert entry.source_text == entry.published_source_text == "Original"
        assert entry.translations[0].value == entry.translations[0].published_value == "Original EN"
        if status == "public":
            assert entries["new"].pending_delete
            assert entries["new"].published_source_text == "Final"
        else:
            assert entries["new"].deleted_at is not None


@pytest.mark.parametrize("kind", ["json", "xlsx"])
def test_explicit_public_import_publishes_existing_unchanged_draft(
    client, session_factory, database_engine, kind
):
    project = make_project(client, layout="flat")
    root = f"/api/projects/{project['id']}"
    client.post(
        f"{root}/strings",
        json={"key": "same", "source_text": "Same", "translations": {"en": "English"}},
    )
    if kind == "json":
        files = json_upload({"vi": {"same": "Same"}, "en": {"same": "English"}})
    else:
        workbook = Workbook()
        workbook.active.append(["key", "vi", "en"])
        workbook.active.append(["same", "Same", "English"])
        raw = io.BytesIO()
        workbook.save(raw)
        files = {"file": ("same.xlsx", raw.getvalue())}
    preview = client.post(
        f"{root}/import", params={"status": "public", "dry_run": True}, files=files
    )
    assert preview.status_code == 200
    with session_factory() as db:
        assert db.scalar(select(StringEntry)).status.value == "draft"
    imported = client.post(f"{root}/import", params={"status": "public"}, files=files)
    assert imported.status_code == 200, imported.text
    with session_factory() as db:
        entry = db.scalar(select(StringEntry))
        assert entry.status.value == "public"
        assert entry.published_source_text == "Same"
        assert entry.translations[0].published_value == "English"
        audit = db.scalars(
            select(Activity).where(Activity.batch_id == uuid.UUID(imported.json()["batch_id"]))
        ).all()
        assert len(audit) == 1 and audit[0].event_type == "string.published"
    with sql_calls(database_engine) as calls:
        assert (
            client.post(f"{root}/import", params={"status": "public"}, files=files).status_code
            == 200
        )
    assert not content_writes(calls)


def test_public_json_import_restores_pending_delete_and_undo_requeues_it(
    client, session_factory, database_engine
):
    project = make_project(client, layout="flat")
    root = f"/api/projects/{project['id']}"
    created = client.post(
        f"{root}/strings",
        json={"key": "same", "source_text": "Same", "translations": {"en": "English"}},
    ).json()
    publish_strings(client, project["id"], [created["id"]])
    assert client.delete(f"{root}/strings/{created['id']}").status_code == 204
    files = json_upload({"vi": {"same": "Same"}, "en": {"same": "English"}})

    preview = client.post(
        f"{root}/import", params={"status": "public", "dry_run": True}, files=files
    )
    assert preview.status_code == 200
    with session_factory() as db:
        assert db.get(StringEntry, uuid.UUID(created["id"])).pending_delete

    imported = client.post(f"{root}/import", params={"status": "public"}, files=files)
    assert imported.status_code == 200, imported.text
    assert imported.json()["updated"] == 1
    batch_id = imported.json()["batch_id"]
    with session_factory() as db:
        entry = db.get(StringEntry, uuid.UUID(created["id"]))
        assert not entry.pending_delete
        assert entry.source_text == entry.published_source_text == "Same"
        assert entry.translations[0].value == entry.translations[0].published_value == "English"
        audit = db.scalars(select(Activity).where(Activity.batch_id == uuid.UUID(batch_id))).all()
        assert len(audit) == 1
        assert audit[0].event_type == "string.restored"
        assert audit[0].before["pending_delete"] is True
        assert audit[0].after["pending_delete"] is False
    draft = client.get(f"{root}/export", params={"layout": "flat", "stage": "draft"})
    assert draft.status_code == 200
    assert draft.json()["vi"]["same"] == "Same"

    with sql_calls(database_engine) as calls:
        repeated = client.post(f"{root}/import", params={"status": "public"}, files=files)
    assert repeated.status_code == 200
    assert repeated.json()["updated"] == 0
    assert not content_writes(calls)

    undone = client.post(f"{root}/activities/batch/{batch_id}/revert")
    assert undone.status_code == 200, undone.text
    with session_factory() as db:
        entry = db.get(StringEntry, uuid.UUID(created["id"]))
        assert entry.pending_delete
        assert entry.published_source_text == "Same"


@pytest.mark.parametrize("description", [None, "   ", "  Replacement  ", "absent"])
@pytest.mark.parametrize("changed", [False, True])
def test_excel_blank_description_preserves_text_independently_of_other_changes(
    client, session_factory, description, changed
):
    project = make_project(client, layout="flat")
    root = f"/api/projects/{project['id']}"
    client.post(
        f"{root}/strings",
        json={
            "key": "existing",
            "source_text": "Original",
            "description": "Keep me",
            "translations": {"en": "English"},
        },
    )
    workbook = Workbook()
    headers = ["key", "vi", "en"]
    if description != "absent":
        headers.append("description")
    workbook.active.append(headers)
    for key in ("existing", "new"):
        row = [key, "Changed" if changed else "Original", "Changed EN" if changed else "English"]
        if description != "absent":
            row.append(description)
        workbook.active.append(row)
    raw = io.BytesIO()
    workbook.save(raw)
    imported = client.post(f"{root}/import", files={"file": ("description.xlsx", raw.getvalue())})
    assert imported.status_code == 200, imported.text
    with session_factory() as db:
        entries = {entry.key: entry for entry in db.scalars(select(StringEntry))}
        supplied = description == "  Replacement  "
        assert entries["existing"].description == ("Replacement" if supplied else "Keep me")
        assert entries["new"].description == ("Replacement" if supplied else None)


@pytest.mark.parametrize("dry_run", [False, True])
@pytest.mark.parametrize(
    "bad", [{"title": "Title"}, {"vi": {"title": 123}}, {"vi": []}, [], "Title"]
)
def test_invalid_later_module_returns_400_before_writes(
    client, database_engine, session_factory, dry_run, bad
):
    project = make_project(client)
    data = {"modules": {"first": {"vi": {"good": "Good"}}, "home": bad}}
    before = counts(session_factory)
    with sql_calls(database_engine) as calls:
        response = client.post(
            f"/api/projects/{project['id']}/import",
            files=json_upload(data),
            params={"dry_run": dry_run},
        )
    assert response.status_code == 400, response.text
    assert (
        response.json()["detail"]
        == "Module 'home' must map locale codes to objects containing string keys and values"
    )
    assert not content_writes(calls)
    assert counts(session_factory) == before


def test_invalid_unassigned_map_is_rejected(client):
    project = make_project(client)
    response = client.post(
        f"/api/projects/{project['id']}/import",
        files=json_upload(
            {"modules": {"home": {"vi": {"good": "Good"}}}, "unassigned": {"bad": "Bad"}}
        ),
    )
    assert response.status_code == 400
    assert "Unassigned strings must map locale codes" in response.json()["detail"]
