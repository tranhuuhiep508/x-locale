"""Nightly 100k-string benchmark against PostgreSQL and four API workers.

Run with POSTGRES_TEST_URL (an admin database URL) and an output directory.
The database is temporary and is dropped even when an assertion fails.
"""

from __future__ import annotations

import argparse
import io
import json
import os
import socket
import subprocess
import sys
import threading
import time
import uuid
from concurrent.futures import ThreadPoolExecutor
from contextlib import contextmanager
from pathlib import Path

import httpx
from openpyxl import Workbook
from sqlalchemy import create_engine, insert
from sqlalchemy.engine import make_url

from app.auth import hash_api_key
from app.models import ApiKey, Project, StringEntry, Translation

ROWS = 100_000
TARGETS = ["en", "fr", "de", "es", "ja", "ko", "zh", "th", "id", "pt"]
MAX_WORKER_RSS_BYTES = 2 * 1024**3
BACKEND_ROOT = Path(__file__).resolve().parents[1]


def _port() -> int:
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        return int(sock.getsockname()[1])


@contextmanager
def database(admin_url: str):
    base = make_url(admin_url)
    name = f"xlocale_bench_{uuid.uuid4().hex}"
    admin = create_engine(base.set(database="postgres"), isolation_level="AUTOCOMMIT")
    with admin.connect() as connection:
        connection.exec_driver_sql(f'CREATE DATABASE "{name}"')
    url = base.set(database=name).render_as_string(hide_password=False)
    try:
        yield url
    finally:
        with admin.connect() as connection:
            connection.exec_driver_sql(f'DROP DATABASE "{name}" WITH (FORCE)')
        admin.dispose()


def _project(engine, slug: str) -> tuple[str, str]:
    project_id = uuid.uuid4()
    key = f"bench-{uuid.uuid4().hex}"
    with engine.begin() as connection:
        connection.execute(insert(Project.__table__), {
            "id": project_id, "name": slug, "slug": slug,
            "base_language": "vi", "target_languages": TARGETS, "layout": "flat",
        })
        connection.execute(insert(ApiKey.__table__), {
            "id": uuid.uuid4(), "project_id": project_id, "name": "benchmark",
            "key_prefix": key[:16], "key_hash": hash_api_key(key),
        })
    return str(project_id), key


def seed(engine, project_id: str) -> None:
    for start in range(0, ROWS, 1000):
        string_rows = []
        translation_rows = []
        for number in range(start, min(start + 1000, ROWS)):
            string_id = uuid.uuid4()
            key = f"bench.{number:06d}"
            string_rows.append({
                "id": string_id, "project_id": uuid.UUID(project_id), "key": key,
                "source_text": f"Nguồn {number}", "status": "public",
                "published_key": key, "published_source_text": f"Nguồn {number}",
                "pending_delete": False,
            })
            for locale in TARGETS:
                value = f"{locale} translation {number}"
                translation_rows.append({
                    "id": uuid.uuid4(), "string_id": string_id, "locale": locale,
                    "value": value, "published_value": value,
                })
        with engine.begin() as connection:
            connection.execute(insert(StringEntry.__table__), string_rows)
            connection.execute(insert(Translation.__table__), translation_rows)


def _worker_pids(parent: int) -> list[int]:
    children = Path(f"/proc/{parent}/task/{parent}/children")
    if not children.exists():
        return []
    return [int(value) for value in children.read_text().split()]


def _rss(pid: int) -> int:
    status = Path(f"/proc/{pid}/status")
    if not status.exists():
        return 0
    for line in status.read_text().splitlines():
        if line.startswith("VmRSS:"):
            return int(line.split()[1]) * 1024
    return 0


class MemorySampler:
    def __init__(self, server_pid: int):
        self.server_pid = server_pid
        self.peak = 0
        self.stop = threading.Event()
        self.thread = threading.Thread(target=self._run, daemon=True)

    def _run(self) -> None:
        while not self.stop.wait(0.1):
            workers = _worker_pids(self.server_pid)
            self.peak = max(self.peak, *(_rss(pid) for pid in workers), 0)

    def __enter__(self):
        self.thread.start()
        return self

    def __exit__(self, *_):
        self.stop.set()
        self.thread.join()


def _wait_for_api(base: str, process: subprocess.Popen) -> None:
    deadline = time.monotonic() + 90
    while time.monotonic() < deadline:
        if process.poll() is not None:
            raise RuntimeError("API workers exited before becoming healthy")
        try:
            if httpx.get(f"{base}/health", timeout=2).status_code == 200:
                return
        except httpx.HTTPError:
            pass
        time.sleep(0.25)
    raise TimeoutError("API workers did not become healthy")


def catalog_run(base: str, project_id: str, key: str) -> float:
    headers = {"X-API-Key": key}
    paths = [
        f"/api/projects/{project_id}/strings?page=1&page_size=50",
        f"/api/projects/{project_id}/strings?q=bench.0005&page_size=50",
        f"/api/projects/{project_id}/strings?status=public&page_size=50",
        f"/api/projects/{project_id}/activities?page=1&page_size=50",
    ]

    def request(number: int) -> float:
        start = time.perf_counter()
        with httpx.Client(base_url=base, headers=headers, timeout=30) as client:
            response = client.get(paths[number % len(paths)])
            response.raise_for_status()
        return time.perf_counter() - start

    with ThreadPoolExecutor(max_workers=20) as pool:
        samples = sorted(pool.map(request, range(200)))
    return samples[int(len(samples) * 0.95) - 1]


def json_file() -> bytes:
    payload = {locale: {f"import.{n:06d}": f"{locale} value {n}" for n in range(ROWS)} for locale in ["vi", *TARGETS]}
    return json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode()


def excel_file() -> bytes:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("strings")
    sheet.append(["key", "description", "tags", "vi", *TARGETS])
    for number in range(ROWS):
        sheet.append([f"import.{number:06d}", "", "", f"Nguồn {number}", *(f"{locale} value {number}" for locale in TARGETS)])
    data = io.BytesIO()
    workbook.save(data)
    return data.getvalue()


def timed_request(base: str, key: str, path: str, *, files=None) -> float:
    start = time.perf_counter()
    with httpx.Client(base_url=base, headers={"X-API-Key": key}, timeout=360) as client:
        response = client.post(path, files=files) if files else client.get(path)
        response.raise_for_status()
    return time.perf_counter() - start


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    admin_url = os.environ["POSTGRES_TEST_URL"]
    args.output.mkdir(parents=True, exist_ok=True)
    with database(admin_url) as url:
        env = {**os.environ, "DATABASE_URL": url, "AUTH_DEV_BYPASS": "true", "X_LOCALE_SECRET": "benchmark-secret-at-least-32-characters"}
        subprocess.run([sys.executable, "-m", "alembic", "upgrade", "head"], cwd=BACKEND_ROOT, env=env, check=True)
        engine = create_engine(url)
        project_id, key = _project(engine, "benchmark-catalog")
        seed(engine, project_id)
        port = _port()
        base = f"http://127.0.0.1:{port}"
        server = subprocess.Popen([sys.executable, "-m", "uvicorn", "app.main:app", "--host", "127.0.0.1", "--port", str(port), "--workers", "4"], cwd=BACKEND_ROOT, env=env)
        results: dict[str, list[dict[str, float]]] = {kind: [] for kind in ("catalog", "export", "json_import", "excel_import")}
        failures: list[str] = []
        try:
            _wait_for_api(base, server)
            # Warm connections, query plans, and application imports before measurement.
            timed_request(base, key, f"/api/projects/{project_id}/strings?page_size=1")
            timed_request(base, key, f"/api/projects/{project_id}/export?stage=public&locale=en")
            warm_id, warm_key = _project(engine, "benchmark-warmup")
            timed_request(base, warm_key, f"/api/projects/{warm_id}/import", files={"file": ("vi.json", b'{"warm.json":"Warm"}')})
            warm_book = Workbook(write_only=True)
            warm_sheet = warm_book.create_sheet("strings")
            warm_sheet.append(["key", "vi", "en"])
            warm_sheet.append(["warm.excel", "Warm", "Warm"])
            warm_data = io.BytesIO()
            warm_book.save(warm_data)
            timed_request(base, warm_key, f"/api/projects/{warm_id}/import", files={"file": ("warm.xlsx", warm_data.getvalue())})
            payloads = {"json_import": json_file(), "excel_import": excel_file()}
            for run in range(3):
                with MemorySampler(server.pid) as memory:
                    p95 = catalog_run(base, project_id, key)
                results["catalog"].append({"run": run + 1, "p95_seconds": p95, "peak_worker_rss_bytes": memory.peak})
                with MemorySampler(server.pid) as memory:
                    duration = timed_request(base, key, f"/api/projects/{project_id}/export?stage=public&locale=en")
                results["export"].append({"run": run + 1, "seconds": duration, "peak_worker_rss_bytes": memory.peak})
                for kind, extension in (("json_import", "json"), ("excel_import", "xlsx")):
                    import_id, import_key = _project(engine, f"benchmark-{extension}-{run}")
                    with MemorySampler(server.pid) as memory:
                        duration = timed_request(base, import_key, f"/api/projects/{import_id}/import", files={"file": (f"import.{extension}", payloads[kind])})
                    results[kind].append({"run": run + 1, "seconds": duration, "peak_worker_rss_bytes": memory.peak})
            for kind, runs in results.items():
                for run in runs:
                    duration = run.get("p95_seconds", run.get("seconds", 0))
                    budget = 1 if kind == "catalog" else 15 if kind == "export" else 300
                    if duration >= budget:
                        failures.append(f"{kind} run {run['run']}: {duration:.2f}s >= {budget}s")
                    if run["peak_worker_rss_bytes"] > MAX_WORKER_RSS_BYTES:
                        failures.append(f"{kind} run {run['run']}: API worker exceeded 2 GiB RSS")
        finally:
            report = {"rows": ROWS, "target_locales": len(TARGETS), "api_workers": 4, "results": results, "failures": failures}
            (args.output / "benchmark.json").write_text(json.dumps(report, indent=2))
            print(json.dumps(report, indent=2))
            server.terminate()
            server.wait(timeout=20)
            engine.dispose()
        if failures:
            raise SystemExit(1)


if __name__ == "__main__":
    main()
